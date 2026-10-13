"""API tests with DB and Google mocked (no network needed)."""
import importlib
import os
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from make_sample_template import create_template

RAW_COMMENT = "  Cliente pide enlace\x0b dedicado.\nLlamar antes de visitar.  "
BODY = {"ticket_id": "T-1001", "client_name": "Juan Pérez / Ferretería", "client_code": 4521,
        "sales_exec": "Ana López", "contact_name": "Roxana Lopez", "contact_phone": "5880 9994",
        "raw_comment": RAW_COMMENT, "latitude": 12.1364, "longitude": -86.2514,
        "service_type": "internet", "medium": "fibra", "bandwidth": "100 Mbps"}
HEADERS = {"X-API-Key": "secret"}


@pytest.fixture
def client(tmp_path, monkeypatch):
    template = tmp_path / "t.xlsx"
    create_template(template)
    env = {"DB_HOST": "x", "DB_PORT": "5432", "DB_NAME": "x", "DB_USER": "x", "DB_PASSWORD": "x",
           "GOOGLE_API_KEY": "x", "SERVICE_API_KEY": "secret",
           "TEMPLATE_PATH": str(template), "OUTPUT_DIR": str(tmp_path / "out"),
           "LOG_FILE": str(tmp_path / "log.txt"),
           "RED_FIBRA_DB": str(tmp_path / "missing_red_fibra.db"),   # origin stays pending
           "VISIO_DB_NAME": "",                                      # Visio PostgreSQL disabled
           "RULES_PATH": os.path.join(os.path.dirname(__file__), "..", "business_rules.yaml")}
    for k, v in env.items():
        monkeypatch.setenv(k, v)
    import config
    config.get_settings.cache_clear()
    import main
    importlib.reload(main)
    monkeypatch.setattr(main.db, "check_connection", lambda: None)
    seg = main.db.Segment(id=7, tipo="aereo", lat=12.13, lon=-86.25, distance_m=120.0)
    route = main.routing.Route(distance_m=1850.0, duration_s=300.0,
                               points=[(12.13, -86.25), (12.14, -86.25), (12.14, -86.24)])
    monkeypatch.setattr(main.db, "find_nearest_connection_point", lambda *a, **k: seg)
    # 1850 m by route for 120 m in straight line: no other KMZ mufa to compare (no network call)
    monkeypatch.setattr(main.db, "find_connection_candidates", lambda *a, **k: [])
    monkeypatch.setattr(main.db, "find_mufas_near_route", lambda *a, **k: [])
    monkeypatch.setattr(main.db, "find_point_coords", lambda ids: {})
    monkeypatch.setattr(main.db, "find_upstream_mufa_chain",
                        lambda mufa_id, nombre, *a, **k: [main.db.ChainMufa(mufa_id, nombre)])
    monkeypatch.setattr(
        main.db,
        "find_existing_poles_near_route",
        lambda *a, **k: [main.db.ExistingPole(id=8, lat=12.135, lon=-86.25, route_position_m=50.0)],
    )
    monkeypatch.setattr(main.routing, "get_route", lambda *a, **k: route)
    # LibreOffice recalculation has its own test; API tests skip it for speed.
    monkeypatch.setattr(main.recalc, "find_soffice", lambda *a, **k: None)
    # No tile downloads in API tests: the map is unavailable unless a test provides one.
    monkeypatch.setattr(main.mapa_ruta, "generar_mapa_ruta", lambda *a, **k: None)
    return TestClient(main.app)


def test_health(client):
    assert client.get("/health").json() == {"status": "ok"}


def test_requires_api_key(client):
    assert client.post("/presupuesto", json=BODY).status_code == 401


def test_coordinates_outside_nicaragua_is_422(client):
    r = client.post("/presupuesto", json={**BODY, "latitude": 40.4, "longitude": -3.7}, headers=HEADERS)
    assert r.status_code == 422


def test_missing_field_is_422(client):
    body = {k: v for k, v in BODY.items() if k != "medium"}
    assert client.post("/presupuesto", json=body, headers=HEADERS).status_code == 422


@pytest.mark.parametrize("field", ["client_name", "client_code", "sales_exec", "contact_name",
                                   "contact_phone", "raw_comment"])
def test_missing_new_field_is_422(client, field):
    body = {k: v for k, v in BODY.items() if k != field}
    assert client.post("/presupuesto", json=body, headers=HEADERS).status_code == 422


def test_nonexistent_budget_excel_returns_404(client):
    response = client.get("/presupuesto/missing123/excel", headers=HEADERS)
    assert response.status_code == 404
    assert response.json()["error"] == "not_found"


def test_success_generates_excel(client):
    r = client.post("/presupuesto", json=BODY, headers=HEADERS)
    assert r.status_code == 200
    data = r.json()
    assert data["budget_id"].isalnum()
    assert data["distance_km"] == 1.85
    assert data["exceeds_5km"] is False
    assert data["pending_validation_items"]
    assert data["pending_validation_poles"] == [{
        "id": 8,
        "latitude": 12.135,
        "longitude": -86.25,
        "distance_along_route_m": 50.0,
    }]
    assert os.path.basename(data["excel_path"]) == f"Juan_Pérez_Ferretería_{data['budget_id']}.xlsx"
    assert data["budget"]["quantities"]["Instalacion de Fibra Optica Aerea ADSS-FLAT/FTTH (6H - 12H)"] == 0
    assert data["budget"]["quantities"]["Instalacion de Fibra Optica Area ADSS (24H - 48H)"] == 2305
    assert os.path.exists(data["excel_path"])

    workbook = load_workbook(data["excel_path"], data_only=False)
    budget_sheet = workbook["Presupuesto"]
    assert budget_sheet["G12"].value == 0
    assert budget_sheet["G13"].value == 2305
    assert budget_sheet["G14"].value == 0
    assert budget_sheet["G20"].value == 1
    # Fiber material = aerial installation only (duct ignored): Flat <- G12, ADSS 24H <- G13.
    assert budget_sheet["G34"].value == "=G12"
    assert budget_sheet["G41"].value == "=G13"
    assert budget_sheet["G35"].data_type == "f"
    assert budget_sheet["H12"].value == "=G12*1"
    assert workbook["COSTOS"]["B1"].value == "=1+1"
    assert workbook["Ruta proyectada"]["B3"].value == "T-1001"

    route_sheet = workbook["Ruta proyectada"]
    assert route_sheet["A5"].value == "Origen de Tráfico (Nodo): Pendiente de validar"
    assert route_sheet["A6"].value == "Elaborado por: Fiber_Budget"
    assert route_sheet["A7"].value == "Ejecutivo de ventas: Ana López"
    assert route_sheet["A8"].value == "Código cliente: 4521"
    assert route_sheet["A9"].value == "Contactos: Roxana Lopez"
    assert route_sheet["A10"].value == "Teléfono contacto: 5880 9994"
    assert budget_sheet["E8"].value == (
        '=CONCATENATE("Fecha: ",TEXT(DAY(TODAY()-A1),"00")&"/"&TEXT(MONTH(TODAY()-A1),"00")'
        '&"/"&YEAR(TODAY()-A1))')
    # Client name goes where the title formula reads it (G81); comment into the box below the label.
    assert budget_sheet["G81"].value == "Juan Pérez / Ferretería"
    assert budget_sheet["H8"].value == "Fiber_Budget"          # beside "Elaborado por:" (F8:G8)
    assert budget_sheet["F4"].value == '=CONCATENATE("PRESUPUESTO DE OBRA"," - ",G81)'
    assert budget_sheet["C63"].value == "Comentarios:   "
    assert budget_sheet["C64"].value == (
        "Texto de validacion fijo.\n\n\n--- Comentario original del ticket ---\n"
        + RAW_COMMENT.replace("\x0b", ""))
    assert {"G81:I81", "C64:I73"} <= {str(r) for r in budget_sheet.merged_cells.ranges}
    # 1,850 m route -> ADSS: no S/FTTH clamps; green preformed = the template's 2 x Clevis (row 35), kept.
    hardware_rows = {34: "=G12", 38: 0, 39: 0, 40: "=G35*2", 41: "=G13"}
    for row in range(34, 52):
        assert budget_sheet[f"G{row}"].value == hardware_rows.get(row, "=G12+G27")
    # POSTE 8 MTRS (row 45) has a formula in this fixture: kept, and reported.
    assert any("Presupuesto!G45" in warning for warning in data["warnings"])
    assert data["pending_validation_fields"] == ["Origen de Tráfico (Nodo)"]
    assert data["values_cached"] is False
    with zipfile.ZipFile(data["excel_path"]) as package:
        assert b'fullCalcOnLoad="1"' in package.read("xl/workbook.xml")

    download = client.get(f"/presupuesto/{data['budget_id']}/excel", headers=HEADERS)
    assert download.status_code == 200
    assert download.content[:2] == b"PK"
    assert data["pdf_url"] == f"/presupuesto/{data['budget_id']}/pdf"


def test_pdf_endpoint_errors(client, monkeypatch):
    import main

    assert client.get("/presupuesto/missing123/pdf", headers=HEADERS).status_code == 404
    data = client.post("/presupuesto", json=BODY, headers=HEADERS).json()
    pdf_url = data["pdf_url"]
    assert client.get(pdf_url).status_code == 401
    # LibreOffice missing (the fixture's find_soffice returns None).
    missing = client.get(pdf_url, headers=HEADERS)
    assert (missing.status_code, missing.json()["detail"]) == (503, "LibreOffice no disponible")
    # Conversion failure.
    monkeypatch.setattr(main.recalc, "find_soffice", lambda *a, **k: "soffice.exe")

    def broken(*a, **k):
        raise main.pdf_export.PdfExportError("LibreOffice failed (TimeoutExpired)")

    monkeypatch.setattr(main.pdf_export, "xlsx_a_pdf", broken)
    failed = client.get(pdf_url, headers=HEADERS)
    assert failed.status_code == 500 and "No se pudo generar el PDF" in failed.json()["detail"]


def test_pdf_endpoint_returns_the_converted_file(client, monkeypatch, tmp_path):
    import main

    data = client.post("/presupuesto", json=BODY, headers=HEADERS).json()
    fake_pdf = tmp_path / "budget.pdf"
    fake_pdf.write_bytes(b"%PDF-1.7 fake")
    monkeypatch.setattr(main.recalc, "find_soffice", lambda *a, **k: "soffice.exe")
    monkeypatch.setattr(main.pdf_export, "xlsx_a_pdf", lambda xlsx, soffice: fake_pdf)
    response = client.get(data["pdf_url"], headers=HEADERS)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


def test_repeated_ticket_returns_existing_budget(client, monkeypatch):
    import main

    original_get_route = main.routing.get_route
    route_calls = 0

    def count_route_calls(*args, **kwargs):
        nonlocal route_calls
        route_calls += 1
        return original_get_route(*args, **kwargs)

    monkeypatch.setattr(main.routing, "get_route", count_route_calls)
    first = client.post("/presupuesto", json=BODY, headers=HEADERS)
    second = client.post("/presupuesto", json=BODY, headers=HEADERS)

    assert first.status_code == second.status_code == 200
    assert first.json()["budget_id"] == second.json()["budget_id"]
    assert route_calls == 1
    assert main._budget_index_path().is_file()


def test_exceeds_5km_uses_configured_threshold(client, monkeypatch):
    import main

    original_load_rules = main.load_rules

    def rules_with_lower_threshold(path):
        rules = original_load_rules(path)
        rules["budget_alerts"]["distance_threshold_km"] = 1.0
        return rules

    monkeypatch.setattr(main, "load_rules", rules_with_lower_threshold)
    response = client.post("/presupuesto", json=BODY, headers=HEADERS)

    assert response.status_code == 200
    assert response.json()["distance_km"] == 1.85
    assert response.json()["exceeds_5km"] is True


def test_default_gpon_kit_is_written_to_materials(client):
    data = client.post("/presupuesto", json=BODY, headers=HEADERS).json()
    assert data["equipment"]["technology"] == "gpon"
    assert data["equipment"]["technology_source"] == "bandwidth"
    assert data["equipment"]["missing_in_template"] == []
    sheet = load_workbook(data["excel_path"])["Presupuesto"]
    written = {sheet.cell(row, 5).value: sheet.cell(row, 7).value for row in range(56, 63)}
    assert written == {"ODF 2 Hilos": 1, "Huawei HG8310M Gpon": 1,
                       "PATCH CORD SC-SC SM (SINGLE MODE) UPC SIMPLEX 3MTS": 1, "PATCH CORD CAT6 7FT": 1,
                       "HAP AC³ WITH ROUTER/RBD53IG-5HACD2HND": 1, "SFS-7020-WA-SC-DD": 0, "FMC-1001S": 0}
    assert sheet["G32"].value == 0   # Visita Tecnica row in the fixture


def test_metro_with_rack_odf_reports_rows_missing_from_template(client):
    body = {**BODY, "ticket_id": "T-METRO", "technology": "METRO", "odf_rack_ports": "24"}
    response = client.post("/presupuesto", json=body, headers=HEADERS)
    assert response.status_code == 200
    data = response.json()
    assert data["equipment"]["technology_source"] == "request"
    assert data["equipment"]["missing_in_template"] == ["SFS-7020-WB-SC-DD", "Odf Rack 24 Acopladores"]
    assert {"item": "SFS-7020-WB-SC-DD", "quantity": 1} in data["pending_validation_items"]
    sheet = load_workbook(data["excel_path"])["Presupuesto"]
    written = {sheet.cell(row, 5).value: sheet.cell(row, 7).value for row in range(56, 63)}
    assert written["FMC-1001S"] == 2 and written["SFS-7020-WA-SC-DD"] == 1
    assert written["ODF 2 Hilos"] == 1 and written["Huawei HG8310M Gpon"] == 0


@pytest.mark.parametrize("extra", [{"technology": None, "odf_rack_ports": None},
                                   {"technology": "", "odf_rack_ports": ""}, {}])
def test_optional_equipment_fields_accept_missing_null_or_empty(client, extra):
    response = client.post("/presupuesto", json={**BODY, **extra}, headers=HEADERS)
    assert response.status_code == 200


@pytest.mark.parametrize("extra", [{"technology": "fibra"}, {"odf_rack_ports": 48}])
def test_invalid_equipment_values_are_422(client, extra):
    assert client.post("/presupuesto", json={**BODY, **extra}, headers=HEADERS).status_code == 422


def test_reserve_connection_quantities_and_origin_name(client, monkeypatch):
    import main

    reserve = main.db.Segment(id=55, tipo="reserva", lat=12.131, lon=-86.251, distance_m=60.0,
                              is_reserve=True, nombre="RVA KAREN")
    monkeypatch.setattr(main.db, "find_nearest_connection_point", lambda *a, **k: reserve)
    response = client.post("/presupuesto", json=BODY, headers=HEADERS)
    assert response.status_code == 200
    data = response.json()
    quantities = data["quantities"]
    assert quantities["Instalacion de mufas"] == 1
    assert quantities["Intervencion de Mufa Existente"] == 0
    assert quantities["Sangria de Cable o Buffer"] == 2
    assert quantities["Caracterizacion de la fibra optica  - Pruebas de OTDR y VFL"] == 2
    assert data["connection_point"]["name"] == "RVA KAREN"

    # "Origen de trafico (MF)" is the mufa the node's service comes from, not the connection
    # point: unknown here, so it stays empty (the connection point has its own field).
    workbook = load_workbook(data["excel_path"], data_only=False)
    assert workbook["Ruta proyectada"]["B1"].value in (None, "")
    assert workbook["Ruta proyectada"]["B2"].value == f"{reserve.lat:.6f}, {reserve.lon:.6f}"


def _visio_node(**overrides):
    node = {"id": 2423, "nombre": "MF LOYOLA #2 MUFA FOSC", "tipo": "MUFA", "propietario": None,
            "lat": 12.153661, "lon": -86.264043, "fuente_coord": "visio", "zona": "Loyola",
            "notas": None, "distancia_m": 75.5, "enlaces": 0, "clientes": 2}
    return {**node, **overrides}


def test_visio_node_is_connection_point_when_kmz_has_nothing_close(client, monkeypatch):
    import main

    far = main.db.Segment(id=55, tipo="mufa", lat=12.20, lon=-86.30, distance_m=1800.0,
                          is_mufa=True, nombre="MF LEJANA")
    busqueda = main.red_fibra_pg.BusquedaVisio(fuente="red_fibra", radio_m=2000.0, nodo=_visio_node(),
                                               fecha_red="2026-10-05 15:10:24")
    captured = {}

    def fake_route(origin, destination):
        captured["origin"] = origin
        return main.routing.Route(distance_m=120.0, duration_s=30.0, points=[origin, destination])

    def chain_must_not_run(*a, **k):
        raise AssertionError("a Visio node is not in the KMZ: no PostGIS chain")

    monkeypatch.setattr(main.db, "find_nearest_connection_point", lambda *a, **k: far)
    monkeypatch.setattr(main.db, "find_upstream_mufa_chain", chain_must_not_run)
    monkeypatch.setattr(main.red_fibra_pg, "buscar_origen", lambda *a, **k: busqueda)
    monkeypatch.setattr(main.routing, "get_route", fake_route)
    data = client.post("/presupuesto", json=BODY, headers=HEADERS).json()

    assert captured["origin"] == (12.153661, -86.264043)
    origen_block = data["origen"]
    assert (origen_block["fuente"], origen_block["nombre"], origen_block["tipo"]) == (
        "red_fibra", "MF LOYOLA #2 MUFA FOSC", "MUFA")
    assert origen_block["distancia_ruta_m"] == 120.0 and origen_block["distancia_lineal_m"] == 75.5
    assert origen_block["fecha_red"] == "2026-10-05 15:10:24"
    assert data["distance_km"] == 0.12                      # n8n's 2 km rule uses this route
    assert data["quantities"]["Intervencion de Mufa Existente"] == 1
    route_sheet = load_workbook(data["excel_path"])["Ruta proyectada"]
    assert route_sheet["B2"].value == "12.153661, -86.264043"


def test_kmz_point_is_kept_and_flagged_when_missing_from_visio(client, monkeypatch):
    import main

    mufa = main.db.Segment(id=99, tipo="mufa", lat=12.131, lon=-86.251, distance_m=40.0,
                           is_mufa=True, nombre="MF KMZ")
    busqueda = main.red_fibra_pg.BusquedaVisio(fuente="red_fibra", radio_m=2000.0,
                                               nodo=_visio_node(distancia_m=300.0))
    monkeypatch.setattr(main.db, "find_nearest_connection_point", lambda *a, **k: mufa)
    monkeypatch.setattr(main.red_fibra_pg, "buscar_origen", lambda *a, **k: busqueda)
    monkeypatch.setattr(main.origen, "find_traffic_origin", lambda *a, **k: main.origen.TrafficOrigin(
        reason="x", connection_match={"name": "MF OTRA", "distance_m": 126.0, "exact": False}))
    data = client.post("/presupuesto", json=BODY, headers=HEADERS).json()

    origen_block = data["origen"]
    assert (origen_block["fuente"], origen_block["nombre"]) == ("metodo_anterior", "MF KMZ")
    assert origen_block["verificacion_kmz"] == {"coincide": False, "mufa_visio": "MF OTRA", "distancia_m": 126.0}
    assert origen_block["requiere_revision"] is True
    assert "Punto de conexión" in data["pending_validation_fields"]
    route_sheet = load_workbook(data["excel_path"])["Ruta proyectada"]
    assert route_sheet["B2"].value == "12.131000, -86.251000 (Coordenada estimada / validar en campo)"


def test_visio_database_down_keeps_kmz_and_reports_sin_conexion(client, monkeypatch):
    import main

    monkeypatch.setattr(main.red_fibra_pg, "enabled", lambda: True)
    monkeypatch.setattr(main.red_fibra_pg, "connect",
                        lambda: (_ for _ in ()).throw(main.psycopg.OperationalError("password authentication failed")))
    response = client.post("/presupuesto", json=BODY, headers=HEADERS)
    assert response.status_code == 200
    origen_block = response.json()["origen"]
    assert origen_block["fuente"] == "sin_conexion_bd"
    assert origen_block["encontrado"] is False
    assert origen_block["nombre"] is None or origen_block["fuente_coord"] == "kmz"


def test_no_network_point_near_client_is_reported_in_excel(client, monkeypatch):
    import main

    far = main.db.Segment(id=55, tipo="mufa", lat=11.05, lon=-87.4, distance_m=9000.0,
                          is_mufa=True, nombre="MF LEJANA")
    busqueda = main.red_fibra_pg.BusquedaVisio(fuente="red_fibra", radio_m=2000.0)
    monkeypatch.setattr(main.db, "find_nearest_connection_point", lambda *a, **k: far)
    monkeypatch.setattr(main.red_fibra_pg, "buscar_origen", lambda *a, **k: busqueda)
    data = client.post("/presupuesto", json=BODY, headers=HEADERS).json()

    assert data["status"] == "ok"                           # the budget is still produced
    assert (data["origen"]["encontrado"], data["origen"]["mensaje"]) == (False, "sin nodo de red cercano")
    route_sheet = load_workbook(data["excel_path"])["Ruta proyectada"]
    assert route_sheet["B2"].value.endswith(
        "(No se encontró nodo de red a menos de 2000 m; validar con Planta Externa)")


def test_route_map_goes_into_route_sheet_box_and_formulas_stay(client, monkeypatch, tmp_path):
    import main
    from PIL import Image

    captured = {}

    def fake_map(points, origen, cliente, alternativas, aproximada, salida, **kwargs):
        captured.update(points=points, origen=origen, cliente=cliente, postes=kwargs["postes"], salida=salida)
        salida.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (1000, 700), (10, 120, 10)).save(salida)
        return salida

    monkeypatch.setattr(main.mapa_ruta, "generar_mapa_ruta", fake_map)
    data = client.post("/presupuesto", json={**BODY, "ticket_id": "T/77 X"}, headers=HEADERS).json()

    assert captured["salida"].name == "T_77_X.png" and captured["salida"].parent.name == "mapas"
    assert captured["origen"]["lat"] == 12.13 and captured["cliente"]["nombre"] == BODY["client_name"]
    assert captured["postes"] == [(12.135, -86.25)]
    assert data["ruta"]["mapa_generado"] is True and data["ruta"]["puntos"] == 3
    assert data["ruta"]["google_maps_url"] == (
        "https://www.google.com/maps/dir/?api=1&origin=12.130000,-86.250000"
        "&destination=12.136400,-86.251400&travelmode=walking")

    generated = load_workbook(data["excel_path"])
    route_sheet = generated["Ruta proyectada"]
    assert len(route_sheet._images) == 1
    anchor = route_sheet._images[0].anchor
    # Pinned to the corners of C18:H52: from C18 to the start of I53 (0-based col 8, row 52),
    # so Excel and LibreOffice (PDF) both fill the box whatever their column widths.
    assert (anchor._from.col, anchor._from.row, anchor.to.col, anchor.to.row) == (2, 17, 8, 52)
    assert (anchor.to.colOff, anchor.to.rowOff) == (0, 0)
    assert anchor.editAs == "twoCell"     # LibreOffice (PDF) stretches it between the corners too
    assert route_sheet["C18"].value is None
    # The budget sheet keeps every template formula.
    template = load_workbook(main.settings.template_path)["Presupuesto"]
    budget = generated["Presupuesto"]
    for row in template.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("=") and cell.coordinate not in (
                    "E8", "G34", "G38", "G39", "G40", "G41"):   # date rewrite and named material rules
                assert budget[cell.coordinate].value == cell.value, cell.coordinate


def test_map_failure_still_produces_excel_with_notice(client):
    data = client.post("/presupuesto", json=BODY, headers=HEADERS).json()
    assert data["status"] == "ok"
    assert (data["ruta"]["mapa_generado"], data["ruta"]["mapa_archivo"]) == (False, None)
    route_sheet = load_workbook(data["excel_path"])["Ruta proyectada"]
    assert route_sheet["C18"].value == "Mapa no disponible"
    assert route_sheet._images == []


def test_origen_endpoint_returns_block_without_budget(client, monkeypatch):
    import main

    monkeypatch.setattr(main.db, "find_nearest_connection_point",
                        lambda *a, **k: (_ for _ in ()).throw(main.db.DatabaseError("down")))
    busqueda = main.red_fibra_pg.BusquedaVisio(fuente="red_fibra", radio_m=500.0, nodo=_visio_node())
    monkeypatch.setattr(main.red_fibra_pg, "buscar_origen", lambda *a, **k: busqueda)
    assert client.get("/origen?lat=12.153328&lon=-86.263438").status_code == 401
    response = client.get("/origen?lat=12.153328&lon=-86.263438&radio=500", headers=HEADERS)
    assert response.status_code == 200
    block = response.json()["origen"]
    assert (block["nombre"], block["fuente"], block["distancia_ruta_m"]) == (
        "MF LOYOLA #2 MUFA FOSC", "red_fibra", None)
    assert client.get("/origen?lat=40.4&lon=-3.7", headers=HEADERS).status_code == 422


def test_upstream_chain_sets_interventions_backbone_and_labels(client, monkeypatch):
    import main

    mufa = main.db.Segment(id=18274, tipo="mufa", lat=12.199261, lon=-86.364555, distance_m=38.3,
                           is_mufa=True, nombre="MUFA REBECA")
    chain = [main.db.ChainMufa(18274, "MUFA REBECA"), main.db.ChainMufa(18319, "MF KAREN BERMUDEZ"),
             main.db.ChainMufa(18320, "MF CACIQUE")]
    route_mufa = main.db.RouteMufa(id=18274, tipo="mufa", lat=12.199261, lon=-86.364555,
                                   distance_m=3.6, is_backbone=False)
    monkeypatch.setattr(main.db, "find_nearest_connection_point", lambda *a, **k: mufa)
    monkeypatch.setattr(main.db, "find_upstream_mufa_chain", lambda *a, **k: chain)
    monkeypatch.setattr(main.db, "find_mufas_near_route", lambda *a, **k: [route_mufa])
    data = client.post("/presupuesto", json=BODY, headers=HEADERS).json()

    quantities = data["quantities"]
    assert quantities["Intervencion de Mufa Existente"] == 3
    assert quantities["Empalme de fibra optica cables de Red Backbone (core/troncal)"] == 3
    assert quantities["Rotulacion de Cable de fibra en mufas y ODF"] == 3
    assert [m["name"] for m in data["budget"]["intervened_mufas"]] == [
        "MUFA REBECA", "MF KAREN BERMUDEZ", "MF CACIQUE"]


def test_traffic_origin_fills_route_sheet_and_reports_saturation(client, monkeypatch):
    import main

    mufa = main.db.Segment(id=18274, tipo="mufa", lat=12.199261, lon=-86.364555, distance_m=38.3,
                           is_mufa=True, nombre="MUFA REBECA")
    chain = [main.db.ChainMufa(18274, "MUFA REBECA"), main.db.ChainMufa(18319, "MF KAREN BERMUDEZ"),
             main.db.ChainMufa(18320, "MF SIN COORDENADA")]
    captured = {}

    def fake_origin(db_path, points, rules):
        captured["points"] = points
        return main.origen.TrafficOrigin(node="Nodo La Gruta", confidence="baja", label="OLT GRUTA",
                                         path=["MF REBECA", "MF X"], saturated_mufas=["MF X"])

    monkeypatch.setattr(main.db, "find_nearest_connection_point", lambda *a, **k: mufa)
    monkeypatch.setattr(main.db, "find_upstream_mufa_chain", lambda *a, **k: chain)
    monkeypatch.setattr(main.db, "find_point_coords", lambda ids: {18319: (12.199, -86.365)})
    monkeypatch.setattr(main.origen, "find_traffic_origin", fake_origin)
    data = client.post("/presupuesto", json=BODY, headers=HEADERS).json()

    # Connection point first (its own coordinates), then chain mufas that have coordinates.
    assert [(p.nombre, p.lat, p.lon) for p in captured["points"]] == [
        ("MUFA REBECA", 12.199261, -86.364555), ("MF KAREN BERMUDEZ", 12.199, -86.365)]
    assert data["traffic_origin"]["node"] == "Nodo La Gruta"
    assert data["pending_validation_fields"] == ["Origen de Tráfico (Nodo)"]   # baja -> validate
    assert any("SATURADA" in w and "MF X" in w for w in data["warnings"])
    route_sheet = load_workbook(data["excel_path"])["Ruta proyectada"]
    assert route_sheet["A5"].value == "Origen de Tráfico (Nodo): Nodo La Gruta (validar)"


def test_traced_diagram_path_sets_intervened_mufas(client, monkeypatch):
    import main

    mufa = main.db.Segment(id=1521, tipo="mufa", lat=12.119842, lon=-86.263267, distance_m=20.0,
                           is_mufa=True, nombre="MF EDIFICIO DELTA")
    route_mufas = [main.db.RouteMufa(id=1521, tipo="mufa", lat=12.119842, lon=-86.263267,
                                     distance_m=0.0, is_backbone=True),
                   main.db.RouteMufa(id=555, tipo="mufa", lat=12.1201, lon=-86.2626,
                                     distance_m=4.0, is_backbone=False)]
    path = ["MUFA FOSC EDIFICIO DELTA", "MUFA FRENTE A EDIFICIO PELLAS", "MUFA CASAVISION", "MUFA EDIFICIO CAR"]
    monkeypatch.setattr(main.db, "find_nearest_connection_point", lambda *a, **k: mufa)
    monkeypatch.setattr(main.db, "find_mufas_near_route", lambda *a, **k: route_mufas)
    monkeypatch.setattr(main.origen, "find_traffic_origin", lambda *a, **k: main.origen.TrafficOrigin(
        node="Nodo Villa Fontana", confidence="media", path=path, path_traced=True))
    data = client.post("/presupuesto", json=BODY, headers=HEADERS).json()

    # 4 mufas on the diagram path (connection mufa = PostGIS 1521, deduped) + 1 other route mufa.
    quantities = data["quantities"]
    assert data["budget"]["mufa_count_source"] == "diagramas"
    assert quantities["Intervencion de Mufa Existente"] == 5
    assert quantities["Empalme de fibra optica cables de Red Backbone (core/troncal)"] == 5
    assert quantities["Rotulacion de Cable de fibra en mufas y ODF"] == 5
    assert [m["name"] for m in data["budget"]["intervened_mufas"]] == [
        "MF EDIFICIO DELTA", *path[1:], None]


def test_untraced_diagram_path_keeps_postgis_chain(client, monkeypatch):
    import main

    mufa = main.db.Segment(id=18274, tipo="mufa", lat=12.199261, lon=-86.364555, distance_m=38.3,
                           is_mufa=True, nombre="MUFA REBECA")
    chain = [main.db.ChainMufa(18274, "MUFA REBECA"), main.db.ChainMufa(18319, "MF KAREN BERMUDEZ"),
             main.db.ChainMufa(18320, "MF CACIQUE")]
    monkeypatch.setattr(main.db, "find_nearest_connection_point", lambda *a, **k: mufa)
    monkeypatch.setattr(main.db, "find_upstream_mufa_chain", lambda *a, **k: chain)
    monkeypatch.setattr(main.origen, "find_traffic_origin", lambda *a, **k: main.origen.TrafficOrigin(
        node="Nodo La Gruta", confidence="baja", path=["MF REBECA", "MF SARATOGA"], path_traced=False))
    data = client.post("/presupuesto", json=BODY, headers=HEADERS).json()
    assert data["budget"]["mufa_count_source"] == "postgis"
    assert data["quantities"]["Intervencion de Mufa Existente"] == 3


def test_traffic_origin_failure_never_blocks_the_budget(client, monkeypatch):
    import main

    def broken(*a, **k):
        raise main.sqlite3.OperationalError("database is locked")

    monkeypatch.setattr(main.origen, "find_traffic_origin", broken)
    response = client.post("/presupuesto", json=BODY, headers=HEADERS)
    assert response.status_code == 200
    assert response.json()["traffic_origin"]["node"] is None
    assert any("Traffic-origin lookup failed" in w for w in response.json()["warnings"])


def test_mufa_connection_is_used_for_routing_and_route_sheet(client, monkeypatch):
    import main

    mufa = main.db.Segment(id=99, tipo="mufa existente", lat=12.131, lon=-86.251,
                           distance_m=40.0, is_mufa=True)
    captured = {}
    monkeypatch.setattr(main.db, "find_nearest_connection_point", lambda *a, **k: mufa)
    monkeypatch.setattr(main.db, "find_mufas_near_route", lambda *a, **k: [])

    def fake_route(origin, destination):
        captured["origin"] = origin
        return main.routing.Route(distance_m=189.0, duration_s=60.0,
                                  points=[origin, destination])

    monkeypatch.setattr(main.routing, "get_route", fake_route)
    response = client.post("/presupuesto", json=BODY, headers=HEADERS)
    assert response.status_code == 200
    data = response.json()
    assert data["budget"]["quantities"]["Intervencion de Mufa Existente"] == 1
    assert data["budget"]["quantities"]["Instalacion de mufas"] == 0
    assert data["budget"]["quantities"]["Sangria de Cable o Buffer"] == 0
    assert data["budget"]["quantities"]["Caracterizacion de la fibra optica  - Pruebas de OTDR y VFL"] == 2
    assert captured["origin"] == (mufa.lat, mufa.lon)

    workbook = load_workbook(data["excel_path"], data_only=False)
    assert workbook["Ruta proyectada"]["B2"].value == f"{mufa.lat:.6f}, {mufa.lon:.6f}"
    # 189 m -> Flat: FTTH clamps follow the S clamps (row 38); no green preformed.
    budget_sheet = workbook["Presupuesto"]
    assert (budget_sheet["G39"].value, budget_sheet["G40"].value) == ("=G38", 0)
