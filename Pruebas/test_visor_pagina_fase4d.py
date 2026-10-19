"""Visor v2, fase 4B (ruta de fusión editable y su Excel), en la página real (Edge sin ventana): el plan sale con la ruta, una mufa aérea no se
toca por su nombre, marcar / desmarcar y cambiar hilos reescribe los textos (Python real) y el presupuesto cuenta las mufas que se tocan, y
«Exportar al Excel» agrega la franja al libro de prueba (nunca al real) o avisa si está abierto."""
import asyncio
import json

import pytest
from openpyxl import load_workbook

from arnes_pagina import hay_edge, pagina_edge
from test_ruta_fusion import libro_minimo
from test_visor_pagina_fase1 import DATOS

pytestmark = pytest.mark.skipif(not hay_edge(), reason="necesita Edge y la librería websockets")

PLAN = {"potencia": "NODO MASAYA", "compartir_potencia": False, "avisos_rojos": [], "tocadas": 2, "lineas": [], "pasos": [
    {"id": "p1", "tipo": "mufa", "nombre": "MF UNO", "lat": 12.15, "lon": -86.26, "tocar": True, "motivo": None, "hilo_cliente": "H__", "hilo_potencia": "H4", "splitter": None, "texto_manual": None, "fuente_hilos": "excel"},
    {"id": "p2", "tipo": "mufa", "nombre": "MF AEREA DOS", "lat": 12.16, "lon": -86.27, "tocar": False, "motivo": "aérea", "hilo_cliente": "H__", "hilo_potencia": "H__", "splitter": None, "texto_manual": None},
    {"id": "p3", "tipo": "mufa", "nombre": "MF TRES", "lat": 12.17, "lon": -86.28, "tocar": True, "motivo": None, "hilo_cliente": "H7", "hilo_potencia": "H__", "splitter": None, "texto_manual": None}]}

STUBS = """
window.__datos = %(datos)s; window.__estimaciones = [];
window.__stubs.calcular = () => ({ok: true, datos: window.__datos, coordenada: '12.153000, -86.263000'});
window.__stubs.estimar_presupuesto = (lat, lon, metraje, marcas, ruta, extra, mufas) => {
  window.__estimaciones.push({mufas});
  return {ok: true, supuestos: {service_type: 'Residencial', bandwidth: '10 Mbps', technology: 'gpon', cable: 'Flat', metraje_final_m: metraje, umbral_m: 600, postes_marcados: 0},
          campos: {}, resultado: {calculado: true, archivo: 'x.xlsx', avisos: [], resumen: {totales: [{clave: 'total', etiqueta: 'TOTAL', valor: 1000}], secciones: []}}};
};
"""


async def _recorrido(tmp_path):
    import main
    libro = libro_minimo(tmp_path / "rutas.xlsx")
    anterior = main.settings.ruta_fusion_xlsx
    object.__setattr__(main.settings, "ruta_fusion_xlsx", libro)
    object.__setattr__(main.settings, "output_dir", tmp_path / "out")
    try:
        datos = {**DATOS, "plan_fusion": PLAN}
        async with pagina_edge(1366, 768) as p:
            await p.esperar("window.EditarFusion && window.visor && window.visor.mapa", 20)
            await p.evaluar(STUBS % {"datos": json.dumps(datos)})
            await p.evaluar("document.getElementById('coord').value = '12.153, -86.263'; document.getElementById('generar').click(); true")
            await p.esperar("window.__estimaciones.length === 1", 15)
            await p.evaluar("window.visor.mostrarPestana('fusion'); true")
            await p.esperar("document.querySelectorAll('.ef-paso').length === 3", 5)
            assert await p.evaluar("[...document.querySelectorAll('.ef-paso')].map(x => x.classList.contains('ef-no-tocar'))") == [False, True, False]
            assert "no se toca: aérea" in await p.evaluar("document.querySelectorAll('.ef-paso')[1].innerText")
            assert "2 mufa(s) a intervenir" in await p.evaluar("document.querySelector('.ef-bloque h2').innerText.toLowerCase()")
            assert (await p.evaluar("window.__estimaciones[0].mufas")) is None                          # not edited: the computed count
            # ---- tocar la aérea a mano: Python rescribe los textos y el presupuesto cuenta 3 mufas
            await p.evaluar("(() => { const c = document.querySelector('.ef-paso:nth-child(2) [data-ef=tocar]'); c.checked = true; c.dispatchEvent(new Event('change', {bubbles: true})); return true; })()")
            await p.esperar("document.querySelector('.ef-bloque h2').innerText.toLowerCase().includes('3 mufa(s)')", 10)
            await p.esperar("window.__estimaciones.length === 2", 10)
            assert (await p.evaluar("window.__estimaciones[1].mufas")) == 3 and "editada" in await p.evaluar("document.querySelector('.ef-bloque h2').innerText.toLowerCase()")
            textos = await p.evaluar("[...document.querySelectorAll('.ef-texto')].map(t => t.value)")
            assert textos[0].startswith("En MF UNO 12.150000°, -86.260000° ingresar la acometida del cliente. H1 del cliente fusionarlo con H4 hacia MF AEREA DOS")
            assert "hacia NODO MASAYA" in textos[2]
            # ---- cambiar un hilo reescribe el texto
            await p.evaluar("(() => { const i = document.querySelectorAll('[data-ef=hilo_potencia]')[2]; i.value = '12'; i.dispatchEvent(new Event('change', {bubbles: true})); return true; })()")
            await p.esperar("[...document.querySelectorAll('.ef-texto')].pop().value.includes('H12 hacia NODO MASAYA')", 10)
            # ---- exportar al libro (el de prueba)
            await p.evaluar("(() => { const c = document.querySelector('[data-ef=cliente]'); c.value = 'TICKET-5 Ana'; c.dispatchEvent(new Event('change', {bubbles: true})); const m = document.querySelector('[data-ef=mapa]'); m.checked = false; m.dispatchEvent(new Event('change', {bubbles: true})); return true; })()")
            await p.evaluar("document.querySelector('[data-ef=exportar]').click(); true")
            await p.esperar("document.querySelector('.ef-estado.ok')", 15)
            assert "filas 10" in await p.evaluar("document.querySelector('.ef-estado').innerText")
            ws = load_workbook(libro)["2025 - 2026"]
            valores = [c.value for fila in ws.iter_rows(min_row=10) for c in fila if isinstance(c.value, str)]
            assert "TICKET-5 Ana" in valores and any(v.endswith("H12 hacia NODO MASAYA") for v in valores) and valores[-1] == "RUTA SUGERIDA"
            assert len(list((tmp_path / "out" / "ruta_fusion_respaldos").glob("rutas.*.xlsx"))) == 1
            # ---- con el libro abierto en Excel no se escribe
            (tmp_path / "~$rutas.xlsx").write_bytes(b"x")
            await p.evaluar("document.querySelector('[data-ef=exportar]').click(); true")
            await p.esperar("document.querySelector('.ef-estado.error')", 15)
            assert "abierto en Excel" in await p.evaluar("document.querySelector('.ef-estado').innerText")
            assert await p.errores() == []
    finally:
        object.__setattr__(main.settings, "ruta_fusion_xlsx", anterior)


def test_ruta_de_fusion_editable_y_exportar(tmp_path):
    asyncio.run(_recorrido(tmp_path))
