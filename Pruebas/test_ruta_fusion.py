"""Ruta de fusión (visor v2, fase 4B): el plan (visor_rutas/ruta_fusion_plan.py), la franja que se agrega al libro de rutas de fusión a nivel XML
(ruta_fusion_xlsx.py) y «Exportar» (exportar_fusion.py). El libro de prueba se arma a mano con las mismas partes del real (hoja "2025 - 2026"
con filas y celdas combinadas, dibujo con una imagen, comentarios en hilo, hojas aparte): el real tiene clientes y no va al repositorio."""
import io
import re
import zipfile
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from test_visor_rutas import entorno  # noqa: F401  (fixture)

# --------------------------------------------------------------------------- un libro mínimo con la estructura del real

PNG = None


def _png(ancho=800, alto=400) -> bytes:
    from PIL import Image
    b = io.BytesIO()
    Image.new("RGB", (ancho, alto), (20, 90, 160)).save(b, "PNG")
    return b.getvalue()


def libro_minimo(destino: Path, con_dibujo: bool = True) -> Path:
    estilos = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
               '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts><fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>'
               '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders><cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
               '<cellXfs count="1600">' + '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>' * 1600 + '</cellXfs></styleSheet>')
    filas = ['<row r="1" spans="1:9" ht="15"><c r="A1" s="398"><v>46000</v></c></row>',
             '<row r="3" spans="1:9" ht="15.75"><c r="A3" s="202" t="s"><v>0</v></c><c r="B3" s="383" t="s"><v>1</v></c></row>',
             '<row r="4" spans="1:9" ht="15"><c r="A4" s="371" t="s"><v>2</v></c></row>', '<row r="5" spans="1:9" ht="15"><c r="A5" s="374"/></row>',
             '<row r="6" spans="1:9" ht="15"><c r="A6" s="380" t="s"><v>3</v></c></row>', '<row r="7" spans="1:9" ht="15"><c r="A7" s="386"/></row>',
             '<row r="8" spans="1:9" ht="15"/>', '<row r="9" spans="1:9" ht="15"/>']
    hoja2 = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
             'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><dimension ref="A1:XFC9"/>'
             '<sheetViews><sheetView workbookViewId="0"/></sheetViews><sheetFormatPr defaultRowHeight="14.45"/><cols><col min="1" max="1" width="17.8" customWidth="1"/></cols>'
             f'<sheetData>{"".join(filas)}</sheetData><mergeCells count="3"><mergeCell ref="A1:I2"/><mergeCell ref="B3:D3"/><mergeCell ref="A4:D5"/></mergeCells>'
             '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>' + ('<drawing r:id="rId1"/>' if con_dibujo else '') + '<legacyDrawing r:id="rId2"/></worksheet>')
    hoja1 = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><dimension ref="A1"/>'
             '<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>otra hoja</t></is></c></row></sheetData></worksheet>')
    dibujo = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
              'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><xdr:twoCellAnchor editAs="oneCell"><xdr:from><xdr:col>0</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>6</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>'
              '<xdr:to><xdr:col>4</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>8</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="2" name="Imagen 5"/><xdr:cNvPicPr/></xdr:nvPicPr>'
              '<xdr:blipFill><a:blip xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:embed="rId1"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
              '<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="5095875" cy="2619375"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:pic><xdr:clientData/></xdr:twoCellAnchor></xdr:wsDr>')
    partes = {
        "[Content_Types].xml": ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                               '<Default Extension="png" ContentType="image/png"/><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                               '<Default Extension="xml" ContentType="application/xml"/><Default Extension="vml" ContentType="application/vnd.openxmlformats-officedocument.vmlDrawing"/>'
                               '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
                               '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                               '<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
                               '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
                               '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
                               '<Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
                               '<Override PartName="/xl/comments1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.comments+xml"/></Types>'),
        "_rels/.rels": '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="2023 - 2024" sheetId="1" r:id="rId1"/>'
                            '<sheet name="2025 - 2026" sheetId="2" r:id="rId2"/></sheets></workbook>'),
        "xl/_rels/workbook.xml.rels": ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                                       '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
                                       '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet2.xml"/>'
                                       '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
                                       '<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>'),
        "xl/styles.xml": estilos,
        "xl/sharedStrings.xml": ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="4" uniqueCount="4">'
                                 '<si><t>Cliente</t></si><si><t>C1 Ana</t></si><si><t>En MF VIEJA</t></si><si><t>RUTA SUGERIDA</t></si></sst>'),
        "xl/worksheets/sheet1.xml": hoja1, "xl/worksheets/sheet2.xml": hoja2,
        "xl/worksheets/_rels/sheet2.xml.rels": ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                                                + ('<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/>' if con_dibujo else "")
                                                + '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/vmlDrawing" Target="../drawings/vmlDrawing1.vml"/>'
                                                '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="../comments1.xml"/></Relationships>'),
        "xl/drawings/vmlDrawing1.vml": '<xml xmlns:v="urn:schemas-microsoft-com:vml"></xml>',
        "xl/comments1.xml": ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><comments xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><authors><author>Ana</author></authors>'
                             '<commentList><comment ref="A3" authorId="0"><text><t>comentario que no se debe perder</t></text></comment></commentList></comments>'),
        "xl/threadedComments/threadedComment1.xml": "<threadedComments>hilo</threadedComments>",
    }
    if con_dibujo:
        partes["xl/drawings/drawing1.xml"] = dibujo
        partes["xl/drawings/_rels/drawing1.xml.rels"] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                                                         '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/></Relationships>')
        partes["xl/media/image1.png"] = _png(300, 150)
    with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as z:
        for nombre, datos in partes.items():
            z.writestr(nombre, datos)
    return destino


def _partes(ruta: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(ruta) as z:
        return {n: z.read(n) for n in z.namelist()}


LINEAS = ["En MF UNO 12.020434°, -86.151570° ingresar la acometida del cliente. H1 del cliente fusionarlo con H4 hacia MF DOS",
          "En MF DOS 12.021648°, -86.150574° H4 hacia MF UNO fusionarlo con H__ hacia MASAYA (hilos por identificar)"]


# --------------------------------------------------------------------------- el libro

def test_a_strip_is_added_at_the_end_without_touching_anything_else(tmp_path):
    from visor_rutas import ruta_fusion_xlsx as rx
    libro = libro_minimo(tmp_path / "rutas.xlsx")
    antes = _partes(libro)
    r = rx.agregar_franja(libro, fecha=date(2026, 10, 14), cliente=None, lineas=LINEAS, avisos=["Validar la mufa aérea con Planta Externa"], compartir=True,
                          respaldos=tmp_path / "resp")
    despues = _partes(libro)
    cambiadas = {n for n in antes if antes[n] != despues[n]}
    assert cambiadas == {"xl/worksheets/sheet2.xml"} and set(despues) == set(antes)          # comments, other sheet, styles… byte-identical
    assert r["fila_inicial"] == 10 and r["hoja"] == "2025 - 2026"                           # right after the last row (9): the two empty rows it already had are the gap
    assert Path(r["respaldo"]).is_file() and _partes(Path(r["respaldo"])) == antes          # the backup is the workbook as it was
    ws = load_workbook(libro)["2025 - 2026"]                                                 # a real reader opens it
    assert ws["A10"].value == rx.serie_fecha(date(2026, 10, 14))                              # the date as the workbook stores it (a serial number)
    textos = [c.value for fila in ws.iter_rows(min_row=10) for c in fila if isinstance(c.value, str)]
    assert textos[:4] == ["Cliente", "Diagramas", "Cuadrilla", LINEAS[0]] and "COMPARTIR POTENCIA DEL SP" in textos and textos[-1] == "RUTA SUGERIDA"
    assert ws["B13"].value is None                                                           # hand-made budget: Cliente stays empty
    assert "A10:I11" in {str(m) for m in ws.merged_cells.ranges} and "A10:I11" not in {str(m) for m in load_workbook(r["respaldo"])["2025 - 2026"].merged_cells.ranges}
    assert len(ws.merged_cells.ranges) == 3 + len([m for m in ws.merged_cells.ranges if m.min_row >= 10])
    assert r["fila_final"] == r["fila_ruta"] and ws.max_row >= r["fila_final"]
    # the old content is where it was
    assert ws["B3"].value == "C1 Ana" and ws["A4"].value == "En MF VIEJA" and ws["A6"].value == "RUTA SUGERIDA"


def test_the_client_the_date_and_a_second_strip(tmp_path):
    from visor_rutas import ruta_fusion_xlsx as rx
    libro = libro_minimo(tmp_path / "r.xlsx")
    uno = rx.agregar_franja(libro, fecha=date(2026, 10, 14), cliente="TICKET-7 Ana Pérez", lineas=LINEAS[:1], respaldos=tmp_path / "resp")
    dos = rx.agregar_franja(libro, fecha=date(2026, 10, 15), lineas=LINEAS, respaldos=tmp_path / "resp")
    assert dos["fila_inicial"] == uno["fila_final"] + 3                                     # two empty rows between strips
    ws = load_workbook(libro)["2025 - 2026"]
    assert ws.cell(uno["fila_inicial"] + 3, 2).value == "TICKET-7 Ana Pérez"
    assert ws.cell(dos["fila_inicial"], 1).value == rx.serie_fecha(date(2026, 10, 15))
    filas = [c.row for fila in ws.iter_rows() for c in fila[:1] if c.value is not None]
    assert filas == sorted(filas) and len(set(filas)) == len(filas)                         # no repeated or disordered rows


def test_long_texts_get_taller_boxes_and_special_characters_are_escaped(tmp_path):
    from visor_rutas import ruta_fusion_xlsx as rx
    assert [rx.filas_necesarias("x" * n) for n in (10, 156, 157, 312, 313, 600)] == [2, 2, 4, 4, 6, 8]
    libro = libro_minimo(tmp_path / "r.xlsx")
    largo = "En MF A & B <fusión> \"x\" " + "texto " * 60
    r = rx.agregar_franja(libro, lineas=[largo], respaldos=tmp_path / "resp")
    ws = load_workbook(libro)["2025 - 2026"]
    assert ws.cell(r["fila_inicial"] + 7, 1).value == largo
    caja = [m for m in ws.merged_cells.ranges if m.min_row == r["fila_inicial"] + 7][0]
    assert (caja.max_row - caja.min_row + 1, caja.min_col, caja.max_col) == (rx.filas_necesarias(largo), 1, 4)


def test_the_map_goes_under_ruta_sugerida_and_the_next_strip_clears_it(tmp_path):
    from visor_rutas import ruta_fusion_xlsx as rx
    libro = libro_minimo(tmp_path / "r.xlsx")
    antes = _partes(libro)
    r = rx.agregar_franja(libro, lineas=LINEAS, imagen=_png(900, 450), respaldos=tmp_path / "resp")
    despues = _partes(libro)
    assert "xl/media/image2.png" in despues and despues["xl/media/image1.png"] == antes["xl/media/image1.png"]
    dibujo = despues["xl/drawings/drawing1.xml"].decode()
    assert dibujo.count("<xdr:twoCellAnchor") == 2 and 'r:embed="rId2"' in dibujo and 'id="3"' in dibujo
    assert 'Target="../media/image2.png"' in despues["xl/drawings/_rels/drawing1.xml.rels"].decode()
    desde = int(re.findall(r"<xdr:from>.*?<xdr:row>(\d+)</xdr:row>", dibujo, re.S)[-1])
    assert desde == r["fila_ruta"]                                                           # right below RUTA SUGERIDA (0-based row = next row)
    assert r["fila_final"] > r["fila_ruta"] + 5                                              # rows are reserved for the picture
    otra = rx.agregar_franja(libro, lineas=LINEAS[:1], respaldos=tmp_path / "resp")
    assert otra["fila_inicial"] >= r["fila_final"] + 3
    assert load_workbook(libro)["2025 - 2026"].max_row >= otra["fila_final"]


def test_without_a_drawing_layer_the_strip_is_added_without_the_map(tmp_path):
    from visor_rutas import ruta_fusion_xlsx as rx
    libro = libro_minimo(tmp_path / "r.xlsx", con_dibujo=False)
    r = rx.agregar_franja(libro, lineas=LINEAS, imagen=_png(), respaldos=tmp_path / "resp")
    assert any("sin el mapa" in a for a in r["avisos"])


def test_an_open_or_invalid_workbook_is_never_written(tmp_path):
    from visor_rutas import ruta_fusion_xlsx as rx
    libro = libro_minimo(tmp_path / "r.xlsx")
    original = libro.read_bytes()
    (tmp_path / "~$r.xlsx").write_bytes(b"x")                                               # Excel's lock file
    with pytest.raises(rx.ExcelAbierto):
        rx.agregar_franja(libro, lineas=LINEAS, respaldos=tmp_path / "resp")
    (tmp_path / "~$r.xlsx").unlink()
    with pytest.raises(rx.LibroNoValido):
        rx.agregar_franja(libro, lineas=[], respaldos=tmp_path / "resp")
    with pytest.raises(rx.LibroNoValido):
        rx.agregar_franja(libro, lineas=LINEAS, hoja="No existe", respaldos=tmp_path / "resp")
    with pytest.raises(rx.LibroNoValido):
        rx.agregar_franja(tmp_path / "no_hay.xlsx", lineas=LINEAS)
    texto = tmp_path / "x.xlsx"
    texto.write_bytes(b"no es un zip")
    with pytest.raises(zipfile.BadZipFile):
        rx.agregar_franja(texto, lineas=LINEAS)
    assert libro.read_bytes() == original and not list(tmp_path.glob(".*.tmp")) and not (tmp_path / "resp").exists()


def test_only_the_last_backups_are_kept(tmp_path):
    from visor_rutas import ruta_fusion_xlsx as rx
    libro = libro_minimo(tmp_path / "r.xlsx")
    resp = tmp_path / "resp"
    for _ in range(4):
        rx.agregar_franja(libro, lineas=LINEAS[:1], respaldos=resp, conservar=2)
        import time
        time.sleep(1.1)                                                                      # the backup name has seconds
    assert len(list(resp.glob("r.*.xlsx"))) == 2


# --------------------------------------------------------------------------- el plan

def _parada(nombre, lat=12.0, lon=-86.0, anillo=None, **extra):
    return {"orden": 1, "nombre": nombre, "lat": lat, "lon": lon, "hilos_anillo": anillo or [], "fuente_hilos": "sin_dato", "emparejado_por_nombre": False, **extra}


def test_mufas_aereas_and_de_corte_are_not_touched_by_name(entorno):
    from visor_rutas import ruta_fusion_plan as plan
    assert plan.motivo_no_tocar("MF AÉREA CENTRO") == "aérea" and plan.motivo_no_tocar("MUFA DE CORTE #3") == "de corte"
    assert plan.motivo_no_tocar("MF CORTE COMASA #1") == "de corte" and plan.motivo_no_tocar("mf aerea") == "aérea"
    assert plan.motivo_no_tocar("MF FUERZA AEREA") is None                                  # a place, not an aerial mufa
    assert plan.motivo_no_tocar("MF CORTEZA") is None and plan.motivo_no_tocar("MF AEREAL") is None   # whole words only
    assert plan.motivo_no_tocar("MF SOTO GAGO CASA") is None


def test_hilos_come_from_the_ring_excel_and_otherwise_are_unknown(entorno):
    from visor_rutas import ruta_fusion_plan as plan
    assert [plan.hilo(v) for v in ("H12", "12", "hilo 7", "", None, "H0", "x")] == ["H12", "H12", "H7", "H__", "H__", "H__", "H__"]
    anillo = [{"de_mufa": "MF UNO", "hilo_entrante": "H4", "hacia_mufa": "MF TRES", "hilo_saliente": "H9"}]
    assert plan._hilos_de_anillo({"hilos_anillo": anillo}, "MF UNO", "MF TRES") == ("H4", "H9")
    assert plan._hilos_de_anillo({"hilos_anillo": anillo}, "MF TRES", "MF UNO") == ("H9", "H4")        # the other way round
    assert plan._hilos_de_anillo({"hilos_anillo": anillo}, "MF OTRA", "MF NADA") == ("H__", "H__")
    assert plan._hilos_de_anillo({}, "MF UNO", None) == ("H__", "H__")


def _datos_ruta():
    # OLT side first, like fusion.ruta_fusion: MF TRES (the node's mufa) ... MF UNO (the connection)
    return {"origen_trafico": {"nodo": "NODO MASAYA"}, "ruta_fusion": [
        _parada("MF TRES", 12.03, -86.03, orden=1),
        _parada("MF AEREA DOS", 12.02, -86.02, orden=2),
        _parada("MF CUATRO", 12.025, -86.025, orden=3, anillo=[{"de_mufa": "MF UNO", "hilo_entrante": "H4", "hacia_mufa": "MF TRES", "hilo_saliente": "H9"}]),
        _parada("MF UNO CORTE", 12.01, -86.01, orden=4, es_conexion=True)]}


def test_the_plan_goes_from_the_client_to_the_power_and_skips_the_untouchable(entorno):
    from visor_rutas import ruta_fusion_plan as plan
    p = plan.armar(_datos_ruta())
    assert [x["nombre"] for x in p["pasos"]] == ["MF UNO CORTE", "MF CUATRO", "MF AEREA DOS", "MF TRES"]     # client side first
    assert [x["tocar"] for x in p["pasos"]] == [True, True, False, True]                    # the connection mufa is always touched
    assert p["pasos"][2]["motivo"] == "aérea" and p["pasos"][0]["motivo"] is None and p["tocadas"] == 3
    assert any("MF UNO CORTE" in a and "de corte" in a for a in p["avisos_rojos"])          # but a warning says it is "de corte"
    uno, cuatro, tres = p["lineas"]
    assert uno == "En MF UNO CORTE 12.010000°, -86.010000° ingresar la acometida del cliente. H1 del cliente fusionarlo con H__ hacia MF CUATRO (hilos por identificar)"
    assert cuatro == "En MF CUATRO 12.025000°, -86.025000° H4 hacia MF UNO CORTE fusionarlo con H9 hacia MF TRES"      # from the ring Excel
    assert tres.startswith("En MF TRES 12.030000°, -86.030000° H__ hacia MF CUATRO fusionarlo con H__ hacia NODO MASAYA")
    assert plan.armar({"ruta_fusion": [], "origen_trafico": {}})["lineas"] == []


def test_editing_the_plan_rewrites_the_texts(entorno):
    from visor_rutas import ruta_fusion_plan as plan
    p = plan.armar(_datos_ruta())
    pasos = [dict(x) for x in p["pasos"]]
    pasos[2]["tocar"] = True                                                                # the user touches the aerial one after all
    pasos[1]["hilo_cliente"], pasos[1]["hilo_potencia"] = "H4", "12"
    pasos[3]["splitter"] = "1/8"
    pasos[0]["texto_manual"] = "Texto a mano  del técnico"
    nuevo = plan.redactar({**p, "pasos": pasos, "compartir_potencia": True, "avisos_rojos": [" Ojo   con la saturación ", ""]})
    assert nuevo["tocadas"] == 4 and nuevo["lineas"][0] == "Texto a mano del técnico" and nuevo["compartir_potencia"] is True
    assert nuevo["avisos_rojos"] == ["Ojo con la saturación"]
    assert nuevo["lineas"][1].startswith("En MF CUATRO") and "H4 hacia MF UNO CORTE fusionarlo con H12 hacia MF AEREA DOS" in nuevo["lineas"][1]
    assert nuevo["lineas"][3].endswith("fusionarlo con H__ del SPL 1/8 (hilos por identificar)")
    # a new mufa between two, and one removed from the touched list
    pasos.insert(1, {"id": "nueva", "tipo": "nueva", "nombre": "NUEVA", "lat": 12.0, "lon": -86.0, "tocar": True, "hilo_cliente": "H2", "hilo_potencia": "H3"})
    pasos[0]["tocar"] = False
    otra = plan.redactar({**p, "pasos": pasos})
    assert otra["lineas"][0].startswith("En MF NUEVA 12.000000°, -86.000000° instalar MF FOSC y ingresar la acometida del cliente")
    with pytest.raises(ValueError):
        plan.redactar({"pasos": [{"nombre": ""}]})
    with pytest.raises(ValueError):
        plan.redactar({"pasos": [{"nombre": "x"}] * 61})
    with pytest.raises(ValueError):
        plan.redactar("no")


def test_calcular_ruta_brings_the_plan(entorno):
    main, calculo, _ = entorno
    r = calculo.calcular_ruta(12.1364, -86.2514)
    assert set(r["plan_fusion"]) >= {"pasos", "lineas", "tocadas", "compartir_potencia", "avisos_rojos"}


# --------------------------------------------------------------------------- exportar

def test_export_adds_the_strip_to_the_configured_workbook(entorno, tmp_path, monkeypatch):
    main, calculo, _ = entorno
    from visor_rutas import exportar_fusion, ruta_fusion_plan
    from visor_rutas.ventana_web import Api
    libro = libro_minimo(tmp_path / "rutas.xlsx")
    monkeypatch.setattr(main.settings.__class__, "ruta_fusion_xlsx", libro, raising=False) if False else None
    object.__setattr__(main.settings, "ruta_fusion_xlsx", libro)
    plan = ruta_fusion_plan.armar(_datos_ruta())
    api = Api()
    assert api.redactar_fusion({**plan, "compartir_potencia": True})["plan"]["compartir_potencia"] is True
    r = api.exportar_fusion(plan, {"cliente": "TICKET-9", "fecha": "2026-10-14"})
    assert r["ok"] and r["pasos"] == 3 and r["fila_inicial"] == 10 and Path(r["respaldo"]).parent == main.settings.output_dir / "ruta_fusion_respaldos"
    ws = load_workbook(libro)["2025 - 2026"]
    assert ws.cell(13, 2).value == "TICKET-9" and ws.cell(r["fila_inicial"], 1).value == 46309
    # errors come back as text, nothing is written
    assert "fecha" in api.exportar_fusion(plan, {"fecha": "ayer"})["error"]
    vacio = ruta_fusion_plan.redactar({**plan, "pasos": [{**p, "tocar": False} for p in plan["pasos"]]})
    assert "No hay mufas que intervenir" in api.exportar_fusion(vacio, {})["error"]
    (tmp_path / "~$rutas.xlsx").write_bytes(b"x")
    abierto = api.exportar_fusion(plan, {})
    assert abierto["ok"] is False and abierto["abierto"] is True and "Excel" in abierto["error"]
    assert exportar_fusion  # imported
