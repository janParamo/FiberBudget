"""Create a small template-shaped workbook for local tests only.

The production service must receive the company's real workbook instead.
"""
from pathlib import Path

from openpyxl import Workbook


def create_template(path: Path) -> None:
    """Write a fixture with the real workbook's input/formula boundaries."""
    wb = Workbook()
    ws = wb.worksheets[0]   # a new workbook always has one sheet (wb.active is typed Optional)
    ws.title = "Presupuesto"
    labels = {
        12: "Instalacion de Fibra Optica Aerea ADSS-FLAT/FTTH (6H - 12H)",
        13: "Instalacion de Fibra Optica Area ADSS (24H - 48H)",
        14: "Intervencion de Mufa Existente",
        15: "Empalme de fibra optica cables de Red Backbone (core/troncal)",
        16: "Empalme de Fibra Optica Red de Acceso, Ultima Milla",
        17: "Instalacion de Caja Terminal / ODF",
        18: "Rotulacion de Cable de fibra en mufas y ODF",
        19: "Caracterizacion de la fibra optica  - Pruebas de OTDR y VFL",
        20: "Instalacion de herraje en poste existente",
        21: "Instalacion de poste metalico tipo tuvo 8 mt.",
        22: "Instalacion de poste metalico tipo tuvo 9 mt.",
        23: "Desplazamiento a [Ciudad] / [xxx] Km",
        24: "Instalacion de mufas",
        25: "Sangria de Cable o Buffer",
        26: "Levantamieto de hilos mufa",
        27: "Movimiento de Reserva existente",
        28: "Instalacion de Fibra en Ducto ADSS-FLAT/FTTH (6H - 12H) / ADSS (24H - 48H)",
        32: "Visita Técnica (Certificación de FO existente cuantificación de obra)",
    }
    ws["E11"] = "Descripción"
    ws["G11"] = "Cantidad"
    for row, label in labels.items():
        ws.cell(row=row, column=5, value=label)
        ws.cell(row=row, column=7, value=0)
        ws.cell(row=row, column=8, value=f"=G{row}*1")
    material_labels = {
        34: "FIBRA OPTICA 12 Hilos ADSS Flat SingleMode Span 100 m",
        35: "Clevis tipo D Chico",
        36: 'Cinta Band It 3/4" (50mts)',
        37: "Hebilla Band It 3/4",
        38: "Grapa Metalica en S",
        39: "Grapa p/ cable fibra FTTH",
        40: "Preformada Punto Verde ADSS 11.7 mm / 12.8 mm",
        41: "FIBRA OPTICA 24 Hilos ADSS SingleMode Span 150 m",
    }
    for row in range(34, 52):
        ws.cell(row=row, column=5, value=material_labels.get(row, f"Material {row}"))
        ws.cell(row=row, column=7, value=f"=G12+G27")
    ws["G41"] = 0   # the real workbook's ADSS 24H fiber row is a constant, not a formula
    ws["G40"] = "=G35*2"   # as in the real workbook: green preformed = 2 x Clevis tipo D Chico
    ws["E45"] = "POSTE 8 MTRS"
    ws["C33"] = "MATERIALES"
    # Customer-equipment rows hold fixed zeros, like the real workbook (no WB / rack ODF rows).
    equipment_labels = ["ODF 2 Hilos", "Huawei HG8310M Gpon",
                        "PATCH CORD SC-SC SM (SINGLE MODE) UPC SIMPLEX 3MTS", "PATCH CORD CAT6 7FT",
                        "HAP AC³ WITH ROUTER/RBD53IG-5HACD2HND", "SFS-7020-WA-SC-DD", "FMC-1001S"]
    for row, label in enumerate(equipment_labels, start=56):
        ws.cell(row=row, column=5, value=label)
        ws.cell(row=row, column=7, value=0)
    ws["E53"], ws["H53"] = "Subtotal", "=SUM(H34:H51)"
    ws["E54"], ws["H54"] = "Impuesto", "=H53*0.15"
    ws["E55"], ws["H55"] = "TOTAL", "=H53+H54"
    ws["F4"] = '=CONCATENATE("PRESUPUESTO DE OBRA"," - ",G81)'
    ws["E8"] = '=CONCATENATE("Fecha: ",TEXT(TODAY()-A1,"DD/MM/yyyy"))'
    ws["F8"] = "Elaborado por:"           # value goes in the merged H8:I8 box, as in the real workbook
    ws.merge_cells("F8:G8")
    ws.merge_cells("H8:I8")
    ws["C63"] = "Comentarios:   "
    ws.merge_cells("C63:I63")
    ws["C64"] = "Texto de validacion fijo."
    ws.merge_cells("C64:I73")
    ws.merge_cells("G81:I81")

    route = wb.create_sheet("Ruta proyectada")
    route["A1"] = "Origen de trafico (MF)"
    route["A2"] = "Punto de conexion"
    route["A3"] = "Numero de ticket"
    route["A4"] = "Coordenadas del cliente"
    route["A5"] = "Origen de Tráfico (Nodo): "
    route["A6"] = "Elaborado por:"
    route["A7"] = "Ejecutivo de ventas:         "
    route["A8"] = "Código cliente: "
    route["A9"] = "Contactos:"
    route["A10"] = "Teléfono contacto:"
    route["C17"] = "ESQUEMA DE RUTA PROYECTADA"   # the route map goes in the box below
    route.merge_cells("C17:H17")
    route.merge_cells("C18:H52")
    for sheet_name in ("COSTOS", "MATERIALES", "Variación de costos"):
        catalog = wb.create_sheet(sheet_name)
        catalog["A1"] = f"sentinel-{sheet_name}"
        catalog["B1"] = "=1+1"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


if __name__ == "__main__":
    out = Path(__file__).resolve().parent.parent / "templates" / "plantilla_presupuesto.xlsx"
    create_template(out)
    print(f"Sample template written to {out}")
