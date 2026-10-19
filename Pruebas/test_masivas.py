"""Coordenadas masivas (visor v2, fase 6): leer las coordenadas de Excel, CSV, Word y texto (visor_rutas/masivas.py), marcar las inválidas, invertidas,
fuera de Nicaragua y duplicadas, y el lote que calcula un presupuesto por punto en segundo plano (con Google y LibreOffice simulados)."""
import io
import threading
import time
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from visor_rutas import masivas

LAT, LON = 12.1364, -86.2514


def xlsx(filas) -> bytes:
    libro = Workbook()
    for f in filas:
        libro.active.append(f)
    b = io.BytesIO()
    libro.save(b)
    return b.getvalue()


def docx(parrafos=(), tablas=()) -> bytes:
    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    p = lambda t: f"<w:p><w:r><w:t xml:space=\"preserve\">{t}</w:t></w:r></w:p>"
    cuerpo = "".join(p(t) for t in parrafos)
    for tabla in tablas:
        cuerpo += "<w:tbl>" + "".join("<w:tr>" + "".join(f"<w:tc>{p(c)}</w:tc>" for c in fila) + "</w:tr>" for fila in tabla) + "</w:tbl>"
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("word/document.xml", f'<?xml version="1.0" encoding="UTF-8"?><w:document {w}><w:body>{cuerpo}</w:body></w:document>')
    return b.getvalue()


def estados(r):
    return [(p["n"], p["estado"]) for p in r["puntos"]]


def test_excel_with_headers_finds_the_columns_names_and_codes():
    crudo = xlsx([["Cliente", "Código", "Latitud", "Longitud", "Notas"],
                  ["Casa Pérez", "C12345", 12.1364, -86.2514, "x"],
                  ["Tienda Ana", None, "12.15", "-86.26", ""],
                  ["Mal escrito", "C1", "abc", "-86.2", ""],
                  ["Al revés", None, -86.25, 12.13, ""],
                  ["Lejos", None, 40.0, 10.0, ""],
                  ["Repetida", None, 12.1364, -86.2514, ""],
                  [None, None, None, None, None]])
    r = masivas.extraer("clientes.xlsx", crudo)
    assert r["formato"] == "Excel" and estados(r) == [(1, "ok"), (2, "ok"), (3, "invalida"), (4, "invertida"), (5, "fuera"), (6, "duplicada")]
    p1, p2, p3, p4, p5, p6 = r["puntos"]
    assert (p1["nombre"], p1["codigo"], p1["lat"], p1["lon"], p1["incluir"]) == ("Casa Pérez", "C12345", LAT, LON, True)
    assert (p2["lat"], p2["lon"]) == (12.15, -86.26) and p2["codigo"] == ""
    assert p3["motivo"] and p3["incluir"] is False and p3["lat"] is None
    assert p4["sugerencia"] == [12.13, -86.25] and p4["incluir"] is False and "invertidas" in p4["motivo"]
    assert p5["incluir"] is False and "fuera de Nicaragua" in p5["motivo"]
    assert p6["motivo"] == "Igual a la fila 1" and p6["incluir"] is False


def test_excel_with_one_coordinate_column_and_without_headers():
    uno = masivas.extraer("a.xlsx", xlsx([["Sitio", "Coordenadas"], ["Torre 1", "12.1364, -86.2514"], ["Torre 2", "13°5'38.09\"N 86°0'7.42\"O"]]))
    assert estados(uno) == [(1, "ok"), (2, "ok")] and uno["puntos"][0]["nombre"] == "Torre 1" and abs(uno["puntos"][1]["lat"] - 13.0939) < 1e-3
    sin = masivas.extraer("b.xlsx", xlsx([["Casa Uno", "C777", 12.1364, -86.2514], ["Casa Dos", None, 12.2, -86.3]]))
    assert estados(sin) == [(1, "ok"), (2, "ok")] and sin["puntos"][0]["nombre"] == "Casa Uno" and sin["puntos"][0]["codigo"] == "C777"
    con_titulo_y_pares = masivas.extraer("c.xlsx", xlsx([["Reporte de clientes"], [], ["Nombre", "Lat", "Lon"], ["X", 12.1, -86.2]]))
    assert estados(con_titulo_y_pares) == [(1, "ok")]


def test_csv_with_semicolons_comma_decimals_and_latin1():
    crudo = "Nombre;Latitud;Longitud\nSan José;12,1364;-86,2514\nÁngel;12.2;-86.3\n".encode("latin-1")
    r = masivas.extraer("datos.csv", crudo)
    assert r["formato"] == "CSV" and estados(r) == [(1, "ok"), (2, "ok")] and r["puntos"][0]["nombre"] == "San José" and r["puntos"][0]["lat"] == LAT
    coma = masivas.extraer("c.csv", b"lat,lon\n12.1364,-86.2514\n12.2,-86.3\n")
    assert estados(coma) == [(1, "ok"), (2, "ok")]


def test_text_and_word_find_coordinates_in_any_format_with_the_name_around():
    texto = ("Visita 1: Casa Pérez 12.1364, -86.2514\n"
             "nada que ver aquí 2025\n"
             "MF 12 NORTE sin coordenada\n"
             "Tienda 13°5'38.09\"N 86°0'7.42\"O\n"
             "12.2 -86.3\n"
             "Repetida 12.1364 -86.2514\n")
    r = masivas.extraer("notas.txt", texto.encode("utf-8"))
    assert r["formato"] == "Texto" and estados(r) == [(1, "ok"), (2, "ok"), (3, "ok"), (4, "duplicada")]
    assert r["puntos"][0]["nombre"] == "Visita 1: Casa Pérez" and r["puntos"][1]["nombre"] == "Tienda" and r["puntos"][2]["nombre"] == ""
    w = masivas.extraer("notas.docx", docx(["Ruta nueva", "Casa Blanca 12.1364, -86.2514", "Otra  12.2, -86.3"], tablas=[[["Taller", "12.3", "-86.4"], ["Sin dato", "", ""]]]))
    assert w["formato"] == "Word" and [(p["nombre"], p["estado"]) for p in w["puntos"]] == [("Casa Blanca", "ok"), ("Otra", "ok"), ("Taller", "ok")]
    assert w["puntos"][2]["lat"] == 12.3


def test_files_that_cannot_be_read_say_why():
    for nombre, crudo, texto in (("a.xls", b"x", ".xls viejos"), ("a.pdf", b"x", "Formato no admitido"), ("a.txt", b"", "vacío"), ("a.txt", b"hola\nadios", "ninguna coordenada"),
                                 ("a.xlsx", b"no es excel", "no se pudo leer"), ("a.docx", b"no es word", "no se pudo leer"), ("a.xlsx", xlsx([["a", "b"], ["x", "y"]]), "ninguna coordenada")):
        with pytest.raises(ValueError, match=texto):
            masivas.extraer(nombre, crudo)
    with pytest.raises(ValueError, match="pesar"):
        masivas.extraer("a.txt", b"0" * (masivas.MAX_ARCHIVO + 1))


def test_the_edited_points_are_checked_again():
    base = masivas.extraer("a.txt", b"Uno 12.1364, -86.2514\nDos 12.2, -86.3\n")["puntos"]
    editados = [dict(base[0], texto="12.1364 -86.2514"), dict(base[1], texto="12.1364, -86.2514"), dict(base[1], n=3, texto="mala", incluir=True), dict(base[1], n=4, texto="12.3, -86.4", incluir=False)]
    r = masivas.revisar(editados)
    assert [(p["n"], p["estado"], p["incluir"]) for p in r] == [(1, "ok", True), (2, "duplicada", False), (3, "invalida", False), (4, "ok", False)]   # «incluir» chosen by the user is kept
    with pytest.raises(ValueError):
        masivas.revisar([])


# --------------------------------------------------------------------------- el lote

class CacheFalsa:
    """Stands for CacheRutas: counts the Google calls and knows which ones are cached."""
    def __init__(self):
        self._rutas, self.llamadas = {}, 0

    @staticmethod
    def _clave(o, d):
        return (o, d)

    def __call__(self, origin, destination):
        self.llamadas += 1
        self._rutas[(origin, destination)] = 1
        return object()


def _punto(n, nombre="P"):
    return {"n": n, "nombre": f"{nombre}{n}", "codigo": "", "texto": "12.1, -86.2", "lat": 12.1 + n / 1000, "lon": -86.2, "estado": "ok", "incluir": True, "motivo": "", "sugerencia": None}


def _esperar(lote, hasta="terminado", seg=10):
    t0 = time.time()
    while lote.estado == "corriendo" and time.time() - t0 < seg:
        time.sleep(0.02)
    assert lote.estado == hasta, (lote.estado, lote.mensaje)


def test_the_batch_runs_in_the_background_keeps_going_after_an_error_and_writes_the_summary(tmp_path, monkeypatch):
    def falso(p, carpeta, cache, bloqueo, ahora=None):
        if p["n"] == 2:
            raise RuntimeError("Google no respondió")
        archivo = carpeta / f"{p['n']:02d}_x.xlsx"
        archivo.write_bytes(b"x")
        return {"n": p["n"], "nombre": p["nombre"], "codigo": "", "lat": p["lat"], "lon": p["lon"], "ok": True, "conexion": "MF X", "distancia_m": 100.5, "metraje_final_m": 140.0,
                "tipo_cable": "Flat", "postes_existentes": 1, "postes_nuevos": 2, "mufas": 1, "total": 5000.0, "nombre_archivo": archivo.name, "avisos": ["Uno", "Dos"], "archivo": str(archivo)}
    monkeypatch.setattr(masivas, "procesar_punto", falso)
    g = masivas.GestorLotes()
    puntos = [_punto(1), _punto(2), _punto(3), dict(_punto(4), incluir=False), dict(_punto(5), estado="invalida")]
    lote = g.iniciar("clientes.xlsx", puntos, tmp_path, por_minuto=30, maximo=10, bloqueo=threading.Lock(), cache_factory=CacheFalsa)
    assert lote.hilo is not None and lote.carpeta.parent == tmp_path and lote.carpeta.name.startswith("masivo_") and lote.carpeta.name.endswith("_clientes")
    _esperar(lote)
    estado = g.obtener(lote.id).publico()
    assert (estado["estado"], estado["hechos"], estado["total"]) == ("terminado", 3, 3)                       # only the valid, included ones
    assert [r["ok"] for r in estado["resultados"]] == [True, False, True] and "No se pudo calcular (RuntimeError)" in estado["resultados"][1]["error"]
    assert g.obtener(lote.id).publico(desde=2)["resultados"][0]["n"] == 3                                  # the page asks only for what is new
    hoja = load_workbook(lote.carpeta / "Resumen.xlsx")["Resumen"]
    filas = [[c.value for c in f] for f in hoja.iter_rows(min_row=4)]
    assert [f[0] for f in filas] == [1, 2, 3] and filas[0][5] == "MF X" and filas[0][12] == 5000.0 and filas[0][14] == "Uno; Dos" and "Google no respondió" in filas[1][14] and filas[1][12] is None
    nombre, datos = g.comprimir(lote.id)
    assert nombre == lote.carpeta.name + ".zip" and sorted(zipfile.ZipFile(io.BytesIO(datos)).namelist()) == ["01_x.xlsx", "03_x.xlsx", "Resumen.xlsx"]
    with pytest.raises(KeyError):
        g.obtener("nada")


def test_one_batch_at_a_time_a_limit_of_points_and_cancel(tmp_path, monkeypatch):
    soltar = threading.Event()

    def lento(p, carpeta, cache, bloqueo, ahora=None):
        soltar.wait(5)
        return {"n": p["n"], "nombre": p["nombre"], "codigo": "", "lat": p["lat"], "lon": p["lon"], "ok": True, "avisos": []}
    monkeypatch.setattr(masivas, "procesar_punto", lento)
    g = masivas.GestorLotes()
    with pytest.raises(ValueError, match="hasta 2"):
        g.iniciar("a.txt", [_punto(1), _punto(2), _punto(3)], tmp_path, 30, 2, threading.Lock(), CacheFalsa)
    with pytest.raises(ValueError, match="No hay coordenadas válidas"):
        g.iniciar("a.txt", [dict(_punto(1), incluir=False)], tmp_path, 30, 2, threading.Lock(), CacheFalsa)
    lote = g.iniciar("a.txt", [_punto(1), _punto(2), _punto(3)], tmp_path, 30, 10, threading.Lock(), CacheFalsa)
    with pytest.raises(ValueError, match="Ya hay un lote en curso"):
        g.iniciar("b.txt", [_punto(1)], tmp_path, 30, 10, threading.Lock(), CacheFalsa)
    with pytest.raises(ValueError, match="todavía está calculando"):
        g.comprimir(lote.id)
    lote.cancelar.set()
    soltar.set()
    _esperar(lote, "cancelado")
    assert lote.publico()["hechos"] == 1 and (lote.carpeta / "Resumen.xlsx").is_file()                    # the one in progress finished; the rest were not done
    g.iniciar("c.txt", [_punto(1)], tmp_path, 30, 10, threading.Lock(), CacheFalsa)                       # and a new batch can start


def test_google_routes_are_asked_at_most_a_number_per_minute_and_the_wait_is_visible(tmp_path):
    class Lote:
        esperando_s = 0
        cancelar = threading.Event()
    cache = CacheFalsa()
    limitada = masivas.CacheLimitada(cache, por_minuto=2, lote=Lote)
    limitada((1, 1), (2, 2))
    limitada((1, 1), (2, 2))                                   # cached: it does not count
    limitada((3, 3), (4, 4))
    assert cache.llamadas == 3 and len(limitada._llamadas) == 2
    # the third distinct one has to wait: cancel releases it
    espera = threading.Thread(target=lambda: Lote.cancelar.set() if time.sleep(0.3) is None else None)
    espera.start()
    with pytest.raises(InterruptedError):
        limitada((5, 5), (6, 6))
    espera.join()
    assert cache.llamadas == 3
