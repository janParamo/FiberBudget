"""Budgets from the desktop viewer (visor_rutas/presupuesto_visor.py and excel_html.py).

The viewer must produce the SAME budget as POST /presupuesto (same function, rules and template),
keep its files in output/visor (never budget_index.json), and show values, never formulas.
"""
import importlib
import os
import time
from datetime import datetime
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from make_sample_template import create_template

BODY = {"ticket_id": "T-100", "client_name": "Cliente Uno", "client_code": "C1", "sales_exec": "Ana",
        "contact_name": "Roxana", "contact_phone": "5880", "raw_comment": "comentario",
        "latitude": 12.1364, "longitude": -86.2514, "service_type": "Empresarial", "medium": "FO",
        "bandwidth": "50 Mbps"}
HEADERS = {"X-API-Key": "secret"}


@pytest.fixture
def entorno(tmp_path, monkeypatch):
    template = tmp_path / "t.xlsx"
    create_template(template)
    env = {"DB_HOST": "x", "DB_PORT": "5432", "DB_NAME": "x", "DB_USER": "x", "DB_PASSWORD": "x",
           "GOOGLE_API_KEY": "x", "SERVICE_API_KEY": "secret",
           "TEMPLATE_PATH": str(template), "OUTPUT_DIR": str(tmp_path / "out"),
           "LOG_FILE": str(tmp_path / "log.txt"), "RED_FIBRA_DB": str(tmp_path / "missing.db"),
           "VISIO_DB_NAME": "", "RULES_PATH": os.path.join(os.path.dirname(__file__), "..", "business_rules.yaml")}
    for k, v in env.items():
        monkeypatch.setenv(k, v)
    import config
    config.get_settings.cache_clear()
    import main
    importlib.reload(main)
    seg = main.db.Segment(id=7, tipo="mufa", lat=12.13, lon=-86.25, distance_m=120.0, is_mufa=True, nombre="MF PRUEBA")
    llamadas = {"google": 0, "mapas": []}

    def get_route(origin, destination):
        llamadas["google"] += 1
        return main.routing.Route(distance_m=1850.0, duration_s=300.0,
                                  points=[(12.13, -86.25), (12.14, -86.25), (12.1364, -86.2514)])

    def mapa(*a, **k):
        llamadas["mapas"].append(Path(a[5]))
        return None

    monkeypatch.setattr(main.db, "check_connection", lambda: None)
    monkeypatch.setattr(main.db, "find_nearest_connection_point", lambda *a, **k: seg)
    monkeypatch.setattr(main.db, "find_connection_candidates", lambda *a, **k: [])
    monkeypatch.setattr(main.db, "find_mufas_near_route", lambda *a, **k: [])
    monkeypatch.setattr(main.db, "find_point_coords", lambda ids: {})
    monkeypatch.setattr(main.db, "find_upstream_mufa_chain",
                        lambda mufa_id, nombre, *a, **k: [main.db.ChainMufa(mufa_id, nombre)])
    monkeypatch.setattr(main.db, "find_existing_poles_near_route", lambda *a, **k: [
        main.db.ExistingPole(id=8, lat=12.135, lon=-86.25, route_position_m=50.0)])
    monkeypatch.setattr(main.routing, "get_route", get_route)
    monkeypatch.setattr(main.recalc, "find_soffice", lambda *a, **k: None)
    monkeypatch.setattr(main.mapa_ruta, "generar_mapa_ruta", mapa)
    from visor_rutas import presupuesto_visor
    return main, presupuesto_visor, tmp_path / "out", llamadas


def _campos(**cambios):
    campos = {"ticket_id": "T-100", "coordenada": "12.1364, -86.2514", "client_name": "Cliente Uno",
              "client_code": "C1", "sales_exec": "Ana", "contact_name": "Roxana", "contact_phone": "5880",
              "service_type": "Empresarial", "medium": "FO", "bandwidth": "50 Mbps", "technology": "auto",
              "odf_rack_ports": "ninguno", "raw_comment": "comentario"}
    return {**campos, **cambios}


def _archivos(carpeta: Path) -> set[str]:
    return {str(p.relative_to(carpeta)) for p in carpeta.rglob("*") if p.is_file()}


# --------------------------------------------------------------------------- misma generación que la API

def test_visor_uses_the_api_generation_same_quantities_and_own_folder(entorno):
    main, pv, salida, llamadas = entorno
    api = TestClient(main.app).post("/presupuesto", json=BODY, headers=HEADERS).json()
    antes = _archivos(salida)
    indice = (salida / "budget_index.json").read_text(encoding="utf-8")

    visor = pv.generar(pv.construir_solicitud(_campos()), ahora=datetime(2026, 10, 7, 15, 30, 5))

    assert visor["quantities"] == api["quantities"]
    assert visor["equipment"] == api["equipment"]
    assert visor["budget"]["mufa_count"] == api["budget"]["mufa_count"]
    assert visor["distance_km"] == api["distance_km"]
    # Only output/visor changed: same index, same API files.
    assert (salida / "budget_index.json").read_text(encoding="utf-8") == indice
    nuevos = _archivos(salida) - antes
    assert nuevos == {os.path.join("visor", "Cliente_Uno_12.136400_-86.251400_20261007-153005.xlsx")}
    assert Path(visor["excel_path"]).parent == salida / "visor"
    assert llamadas["mapas"][-1].parent == salida / "visor" / "mapas"     # the map too
    assert llamadas["mapas"][0].parent == salida / "mapas"                # the API's, as before


def test_form_validation_is_ticketrequest(entorno):
    _, pv, _, _ = entorno
    solicitud = pv.construir_solicitud(_campos(ticket_id="", contact_name="", technology="metro",
                                               odf_rack_ports="24"), ahora=datetime(2026, 1, 2, 3, 4, 5))
    assert solicitud.ticket_id == "VISOR-20260102-030405"
    assert solicitud.contact_name == "Cliente Uno"
    assert (solicitud.technology, solicitud.odf_rack_ports) == ("metro", 24)
    assert (solicitud.latitude, solicitud.longitude) == (12.1364, -86.2514)

    with pytest.raises(pv.ErrorFormulario) as error:
        pv.construir_solicitud(_campos(coordenada="-86.25, 12.13", client_name="", bandwidth=""))
    assert set(error.value.errores) == {"coordenada", "client_name", "bandwidth"}
    assert "invertidas" in error.value.errores["coordenada"]


def test_no_connection_point_is_a_clear_error(entorno, monkeypatch):
    main, pv, _, _ = entorno
    monkeypatch.setattr(main.db, "find_nearest_connection_point",
                        lambda *a, **k: (_ for _ in ()).throw(main.db.NoSegmentFound("x")))
    with pytest.raises(pv.PresupuestoError, match="No eligible mufa or reserva found"):
        pv.generar(pv.construir_solicitud(_campos()))


# --------------------------------------------------------------------------- estimación

def test_estimate_defaults_from_visor_estimacion(entorno):
    main, pv, _, _ = entorno
    rules = main.load_rules(main.settings.rules_path)
    assert rules["visor_estimacion"]["bandwidth"] == "10 Mbps"
    hasta = pv.supuestos_estimacion(600, rules)
    sobre = pv.supuestos_estimacion(600.5, rules)
    assert (hasta["service_type"], hasta["cable"], hasta["technology"], hasta["bandwidth"]) == (
        "Residencial", "Flat", "gpon", "10 Mbps")
    assert (sobre["service_type"], sobre["cable"]) == ("Empresarial", "ADSS")
    propio = {**rules, "visor_estimacion": {**rules["visor_estimacion"], "bandwidth": "200 Mbps",
                                            "servicio_hasta_umbral": "Corporativo"}}
    otro = pv.supuestos_estimacion(100, propio)
    assert (otro["service_type"], otro["bandwidth"], otro["technology"]) == ("Corporativo", "200 Mbps", "metro")


def test_estimate_reuses_the_viewer_route_without_calling_google_again(entorno):
    main, pv, salida, llamadas = entorno
    from visor_rutas import calculo
    cache = pv.CacheRutas()
    with calculo.usando_rutas(cache):
        ruta = calculo.calcular_ruta(12.1364, -86.2514)
    assert llamadas["google"] == 1
    estimacion = pv.estimar(12.1364, -86.2514, ruta["ruta"]["metraje_final_m"], cache)
    assert llamadas["google"] == 1 and cache.llamadas_google == 1          # same coordinate: cached
    s = estimacion["supuestos"]
    assert (s["service_type"], s["bandwidth"], s["technology"]) == ("Empresarial", "10 Mbps", "gpon")
    assert Path(estimacion["respuesta"]["excel_path"]).name.startswith("Estimacion_12.136400_-86.251400_")
    assert estimacion["campos"]["client_name"] == "Estimación visor"
    assert not (salida / "budget_index.json").exists()                    # the viewer never writes it


# --------------------------------------------------------------------------- archivos

def test_cleanup_only_touches_old_files_in_output_visor(entorno):
    main, pv, salida, _ = entorno
    viejo = time.time() - (main.settings.output_retention_days + 1) * 86400
    archivos = {"visor_viejo": salida / "visor" / "a.xlsx", "mapa_viejo": salida / "visor" / "mapas" / "a.png",
                "visor_nuevo": salida / "visor" / "b.xlsx", "api_viejo": salida / "api.xlsx",
                "indice": salida / "budget_index.json"}
    for nombre, ruta in archivos.items():
        ruta.parent.mkdir(parents=True, exist_ok=True)
        ruta.write_text("x")
        if nombre != "visor_nuevo":
            os.utime(ruta, (viejo, viejo))
    assert pv.limpiar_archivos_viejos() == 2
    assert {n for n, r in archivos.items() if r.exists()} == {"visor_nuevo", "api_viejo", "indice"}


def test_page_can_only_ask_for_files_inside_output_visor(entorno):
    _, pv, salida, _ = entorno
    (salida / "visor").mkdir(parents=True)
    dentro, fuera = salida / "visor" / "x.xlsx", salida / "x.xlsx"
    dentro.write_text("x")
    fuera.write_text("x")
    assert pv.archivo_del_visor(str(dentro)) == dentro.resolve()
    for ruta in (fuera, salida / "visor" / ".." / "x.xlsx", salida / "visor" / "nada.xlsx"):
        with pytest.raises(FileNotFoundError):
            pv.archivo_del_visor(str(ruta))


# --------------------------------------------------------------------------- Excel -> HTML y resumen

MONEDA = '_-"C$"* #,##0.00_-;\\-"C$"* #,##0.00_-;_-"C$"* "-"??_-;_-@'


def _libro_como_plantilla(ruta: Path) -> Path:
    """Real template layout: MANO DE OBRA / MATERIALES in C, quantities G, totals I, labels H."""
    wb = Workbook()
    ws = wb.worksheets[0]
    ws.title = "Presupuesto"
    ws["C11"], ws["I11"] = "MANO DE OBRA", 1200.0
    ws["E12"], ws["F12"], ws["G12"], ws["H12"], ws["I12"] = "Instalacion ADSS", "ML", 100, 10.0, 1000.0
    ws["E13"], ws["G13"], ws["I13"] = "Intervencion de Mufa Existente", 0, 0.0
    ws["E14"], ws["G14"], ws["H14"], ws["I14"] = "Empalme", 2, 100.0, "=H14*G14"
    ws["C34"], ws["I34"] = "MATERIALES", "=SUM(I35:I36)"
    ws["E35"], ws["G35"], ws["I35"] = "FIBRA 24 Hilos", "=G12", "=G35*3"
    ws["E36"], ws["G36"], ws["I36"] = "ODF 2 Hilos", 1, 50.0
    ws["C57"], ws["I57"] = "ALCALDIA", 0.0
    ws["H58"], ws["I58"] = "SUBTOTAL", "=I11+I34"
    ws["G60"], ws["I60"] = "Costo Total el 7,29 % ", 10.0
    ws["E61"], ws["H61"], ws["I61"] = "% Descuento al cliente sobre el total:", "TOTAL", "=I58+I60"
    for celda in ("I11", "I12", "I13", "I14", "I34", "I35", "I36", "I57", "I58", "I60", "I61"):
        ws[celda].number_format = MONEDA
    ws["B2"] = "Título combinado"
    ws.merge_cells("B2:D3")
    ws.column_dimensions["J"].hidden = True
    ws["J5"] = "columna oculta"
    wb.create_sheet("Ruta proyectada")["A1"] = "Origen de trafico (MF): MF X"
    wb.save(ruta)
    return ruta


def test_sheet_html_shows_values_never_formulas(tmp_path):
    from visor_rutas import excel_html
    xlsx = _libro_como_plantilla(tmp_path / "p.xlsx")
    html = excel_html.hoja_a_html(xlsx, "Presupuesto", calculado=False)
    assert "=H14" not in html and "=SUM" not in html and "=G12" not in html
    assert "C$ 1,000.00" in html and "C$ -" in html                     # currency and accounting zero
    assert 'class="sin-calculo"' in html                                # formula without LibreOffice: "—"
    assert 'colspan="3"' in html and 'rowspan="2"' in html              # merged B2:D3
    assert "columna oculta" not in html                                 # hidden column J
    assert html.count("<table") == 1
    con_calculo = excel_html.hoja_a_html(xlsx, "Presupuesto", calculado=True)
    assert "sin-calculo" not in con_calculo and "=H14" not in con_calculo
    assert "Origen de trafico (MF): MF X" in excel_html.hoja_a_html(xlsx, "Ruta proyectada", True)


@pytest.mark.parametrize("valor, formato, texto", [
    (1234.5, MONEDA, "C$ 1,234.50"), (-12, MONEDA, "-C$ 12.00"), (0, MONEDA, "C$ -"),
    (0.15, "0%", "15%"), (1731, "General", "1731"), (2.5, "0.00", "2.50"), (1234567, "#,##0", "1,234,567"),
])
def test_number_formats(valor, formato, texto):
    from visor_rutas import excel_html
    assert excel_html.formatear_numero(valor, formato) == texto


def test_summary_lists_lines_with_quantity_and_template_totals(entorno, tmp_path):
    _, pv, _, _ = entorno
    xlsx = _libro_como_plantilla(tmp_path / "p.xlsx")
    resumen = pv.resumen_presupuesto(xlsx, calculado=False)
    mano, materiales = resumen["secciones"]
    assert mano["nombre"] == "Mano de obra" and mano["total"] == 1200.0
    assert [(l["descripcion"], l["cantidad"], l["total"]) for l in mano["lineas"]] == [
        ("Instalacion ADSS", 100, 1000.0), ("Empalme", 2, None)]      # 0 left out; formula total unknown
    assert [l["descripcion"] for l in materiales["lineas"]] == ["ODF 2 Hilos"]
    assert materiales["por_formula"] == 1                                # G35 = G12: needs LibreOffice
    assert [(t["clave"], t["etiqueta"]) for t in resumen["totales"]] == [
        ("alcaldia", "ALCALDIA"), ("subtotal", "SUBTOTAL"), ("costo total", "Costo Total el 7,29 %"), ("total", "TOTAL")]
    assert resumen["totales"][-1]["valor"] is None                       # =I58+I60 without LibreOffice
