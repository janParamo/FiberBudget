"""Metraje final y total de una medida del mapa (visor_rutas/medicion.py): las mismas reglas y la misma plantilla del presupuesto."""
import math
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from test_visor_presupuesto import entorno  # noqa: F401  (fixture)

# ~330 m en dos tramos (hay un giro); con reservas queda bajo el umbral de ADSS
PUNTOS = [[12.1300, -86.2500], [12.1320, -86.2500], [12.1320, -86.2490]]


def _reglas(main):
    return main.load_rules(main.settings.rules_path)


def test_el_metraje_final_es_el_de_las_reglas_de_negocio(entorno):
    main, _, _, _ = entorno
    from visor_rutas import medicion
    r = medicion.calcular(PUNTOS)
    distancia = medicion.longitud_m([tuple(p) for p in PUNTOS])
    assert r["distancia_m"] == pytest.approx(distancia, abs=0.01)
    assert r["metraje_final_m"] == pytest.approx(main.calcular_metraje_final(distancia, _reglas(main)), abs=0.01)
    assert r["metraje_final_m"] > r["distancia_m"]
    assert r["reserva_m"] == pytest.approx(r["metraje_final_m"] - r["distancia_m"], abs=0.01)
    assert r["tipo_cable"] == "Flat"
    # el poste del fixture está en la ruta: 1 existente, el resto nuevos; requeridos = suma
    assert r["postes"]["existentes"] == 1
    assert r["postes"]["requeridos"] == r["postes"]["existentes"] + r["postes"]["nuevos"]
    assert r["postes"]["giros"] == 1
    assert r["supuestos"]["conexiones_nuevas"] == 1 and r["supuestos"]["mufas_intervenidas"] == 0


def test_una_medida_larga_usa_adss(entorno):
    from visor_rutas import medicion
    largo = [[12.00, -86.25], [12.08, -86.25]]          # ~8.9 km
    r = medicion.calcular(largo)
    assert r["tipo_cable"].startswith("ADSS")
    assert r["avisos"] and "5 km" in r["avisos"][0]


@pytest.mark.parametrize("malo", [None, [], [[12.1, -86.2]], [[12.1, -86.2], ["x", 1]], [[12.1, -86.2], [95, 0]],
                                  [[12.1, -86.2], [float("nan"), 0]], [[12.1, -86.2], [12.1]], "12,-86"])
def test_puntos_no_validos_se_rechazan(entorno, malo):
    from visor_rutas import medicion
    with pytest.raises(medicion.MedicionError):
        medicion.calcular(malo)


def test_sin_base_de_postes_se_avisa_y_todos_son_nuevos(entorno, monkeypatch):
    main, _, _, _ = entorno
    from visor_rutas import medicion
    def falla(*a, **k):
        raise main.db.ExistingPolesQueryError("sin base")
    monkeypatch.setattr(main.db, "find_existing_poles_near_route", falla)
    r = medicion.calcular(PUNTOS)
    assert r["postes"]["existentes"] == 0 and r["postes"]["nuevos"] > 0
    assert any("postes existentes" in a for a in r["avisos"])


def test_el_total_llena_la_plantilla_con_las_cantidades_de_la_medida(entorno):
    main, pv, salida, _ = entorno
    from visor_rutas import medicion
    r = medicion.estimar(PUNTOS, ahora=datetime(2026, 10, 14, 9, 0, 0))
    ligera = medicion.calcular(PUNTOS)
    assert r["medicion"] == ligera                                   # los mismos números en las dos llamadas
    resultado = r["resultado"]
    assert Path(salida / "visor" / resultado["archivo"]).is_file()
    assert resultado["calculado"] is False                           # el fixture no tiene LibreOffice
    assert any("LibreOffice" in a for a in resultado["avisos"])
    assert resultado["metraje_final_m"] == pytest.approx(ligera["metraje_final_m"], abs=0.1)
    # el Excel lleva el metraje final y el ticket de medición
    hoja = load_workbook(Path(salida / "visor" / resultado["archivo"]))
    texto = " ".join(str(c.value) for fila in hoja.worksheets[0].iter_rows() for c in fila if c.value is not None)
    assert "MEDICION-20261014-090000" in texto or any("MEDICION-20261014-090000" in str(c.value)
                                                      for ws in hoja.worksheets for fila in ws.iter_rows() for c in fila)
    assert not (salida / "budget_index.json").exists() or "MEDICION" not in (salida / "budget_index.json").read_text(encoding="utf-8")


def test_la_api_del_visor_expone_las_dos_llamadas(entorno):
    main, _, _, _ = entorno
    from visor_rutas.ventana_web import Api
    api = Api()
    ligera = api.medir_ruta(PUNTOS, False)
    assert ligera["ok"] and ligera["medicion"]["metraje_final_m"] > ligera["medicion"]["distancia_m"]
    mala = api.medir_ruta([[1, 2]], False)
    assert mala == {"ok": False, "error": "Una medida necesita al menos dos puntos"}
    total = api.medir_ruta(PUNTOS, True)
    assert total["ok"] and total["resultado"]["archivo"].endswith(".xlsx") and total["supuestos"]["cable"] == "Flat"
    assert not math.isnan(total["medicion"]["metraje_final_m"])
