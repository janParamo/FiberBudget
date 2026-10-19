"""Coordenadas masivas con el cálculo real del visor (Google, base de la red y LibreOffice simulados como en test_visor_presupuesto.py): un lote por la API
del visor, su progreso, el resumen, el .zip y las filas que fallan; y lo que se guarda en Mis lugares / como proyecto (base de prueba)."""
import base64
import io
import time
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from redes_base import base_redes, con_redes  # noqa: F401  (fixtures)
from test_visor_presupuesto import entorno  # noqa: F401  (fixture)
from visor_rutas import masivas
from visor_rutas.redes import db, lugares, masivas_red, usuarios
from visor_rutas.redes.api import METODOS


def _libro(filas) -> str:
    libro = Workbook()
    for f in filas:
        libro.active.append(f)
    b = io.BytesIO()
    libro.save(b)
    return base64.b64encode(b.getvalue()).decode("ascii")


def _esperar(api, id_, seg=30):
    t0 = time.time()
    while time.time() - t0 < seg:
        lote = api.masivas_estado(id_)["lote"]
        if lote["estado"] != "corriendo":
            return lote
        time.sleep(0.05)
    raise AssertionError("el lote no terminó")


def test_a_file_becomes_a_batch_of_budgets_with_a_summary_and_a_zip(entorno):
    from visor_rutas.ventana_web import Api
    main, pv, salida, llamadas = entorno
    api = Api()
    r = api.masivas_leer("clientes.xlsx", _libro([["Nombre", "Código", "Latitud", "Longitud"], ["Casa Pérez", "C12345", 12.1364, -86.2514], ["Tienda Ana", None, 12.1400, -86.2600],
                                                  ["Mal", None, "x", "y"], ["Casa Pérez", None, 12.1364, -86.2514]]))
    assert r["ok"] and r["formato"] == "Excel" and [p["estado"] for p in r["puntos"]] == ["ok", "ok", "invalida", "duplicada"]
    assert api.masivas_leer("a.xls", "eA==")["error"].startswith("Los .xls viejos") and api.masivas_leer("a.txt", "")["ok"] is False
    inicio = api.masivas_iniciar("clientes.xlsx", r["puntos"])
    assert inicio["ok"] and inicio["lote"]["total"] == 2 and inicio["lote"]["estado"] == "corriendo"
    lote = _esperar(api, inicio["lote"]["id"])
    assert (lote["estado"], lote["hechos"], lote["total"]) == ("terminado", 2, 2)
    a, b = lote["resultados"]
    assert (a["ok"], a["nombre"], a["codigo"], a["conexion"]) == (True, "Casa Pérez", "C12345", "MF PRUEBA")
    assert a["distancia_m"] == 1850.0 and a["tipo_cable"].startswith("ADSS") and a["postes_existentes"] == 1 and a["ruta"] and a["conexion_lat"] == 12.13
    assert a["calculado"] is False and a["total"] is None                                                   # the test has no LibreOffice: the amounts show when the Excel is opened
    for x in (a, b):
        assert x["archivo"].endswith(".xlsx") and x["archivo"].startswith(lote["carpeta"])                    # in the batch's folder, not in output/visor
    carpeta = salida / "visor" / lote["carpeta"].rsplit("\\", 1)[-1]
    assert carpeta.is_dir() and sorted(p.name for p in carpeta.glob("*.xlsx")) == ["01_Casa_Perez.xlsx", "02_Tienda_Ana.xlsx", "Resumen.xlsx"]
    assert not list((salida / "visor").glob("Estimacion_*.xlsx"))                                         # nothing left behind
    hoja = load_workbook(carpeta / "Resumen.xlsx")["Resumen"]
    filas = [[c.value for c in f] for f in hoja.iter_rows(min_row=4)]
    assert [f[1] for f in filas] == ["Casa Pérez", "Tienda Ana"] and filas[0][5] == "MF PRUEBA" and filas[0][6] == 1850.0
    z = api.masivas_descargar(inicio["lote"]["id"])
    assert z["ok"] and z["nombre"].endswith(".zip") and sorted(zipfile.ZipFile(io.BytesIO(base64.b64decode(z["datos"]))).namelist()) == ["01_Casa_Perez.xlsx", "02_Tienda_Ana.xlsx", "Resumen.xlsx"]
    assert api.masivas_estado("no")["ok"] is False and api.masivas_cancelar("no")["ok"] is False and api.masivas_descargar("no")["ok"] is False
    assert api.masivas_estado(inicio["lote"]["id"], desde=1)["lote"]["resultados"][0]["nombre"] == "Tienda Ana"


def test_a_point_that_fails_does_not_stop_the_rest(entorno, monkeypatch):
    from visor_rutas.ventana_web import Api
    main, pv, salida, llamadas = entorno
    original = main.db.find_nearest_connection_point
    monkeypatch.setattr(main.db, "find_nearest_connection_point", lambda lat, lon, *a, **k: None if lat > 12.2 else original(lat, lon, *a, **k))
    api = Api()
    puntos = masivas.extraer("a.txt", b"Uno 12.1364, -86.2514\nSin red 12.3000, -86.3000\nTres 12.1400, -86.2600\n")["puntos"]
    lote = _esperar(api, api.masivas_iniciar("a.txt", puntos)["lote"]["id"])
    assert [r["ok"] for r in lote["resultados"]] == [True, False, True] and "No se encontró mufa" in lote["resultados"][1]["error"]
    hoja = load_workbook(lote["carpeta"] + "\\Resumen.xlsx")["Resumen"]
    assert "No se encontró mufa" in [c.value for c in list(hoja.iter_rows(min_row=5, max_row=5))[0]][14]


def test_a_batch_can_be_cancelled_and_the_google_limit_comes_from_the_rules(entorno, monkeypatch):
    from visor_rutas.ventana_web import Api
    main, pv, salida, llamadas = entorno
    api = Api()
    capturado = {}
    real = api._lotes.iniciar

    def espia(nombre, puntos, base, por_minuto, maximo, bloqueo, cache):
        capturado.update(por_minuto=por_minuto, maximo=maximo)
        return real(nombre, puntos, base, por_minuto, maximo, bloqueo, cache)
    monkeypatch.setattr(api._lotes, "iniciar", espia)
    puntos = masivas.extraer("a.txt", "".join(f"P{i} 12.{1300 + i}, -86.2514\n" for i in range(8)).encode())["puntos"]
    inicio = api.masivas_iniciar("a.txt", puntos)
    assert capturado == {"por_minuto": 30, "maximo": 200}                                                    # visor_masivo of business_rules.yaml
    assert api.masivas_cancelar(inicio["lote"]["id"])["ok"] is True
    lote = _esperar(api, inicio["lote"]["id"])
    assert lote["estado"] == "cancelado" and lote["hechos"] < 8


# --------------------------------------------------------------------------- guardar el resultado

PUNTOS = [{"n": 1, "nombre": "Casa Pérez", "codigo": "c12345", "lat": 12.1364, "lon": -86.2514, "conexion": "MF PRUEBA", "conexion_lat": 12.13, "conexion_lon": -86.25,
           "ruta": [[12.13, -86.25], [12.135, -86.25], [12.1364, -86.2514]], "tipo_cable": "Flat"},
          {"n": 2, "nombre": "Tienda Ana", "codigo": "", "lat": 12.14, "lon": -86.26, "conexion": "MF PRUEBA", "conexion_lat": 12.13, "conexion_lon": -86.25,
           "ruta": [[12.13, -86.25], [12.14, -86.26]], "tipo_cable": "ADSS 24 hilos"},
          {"n": 3, "nombre": "Sin conexión", "codigo": "", "lat": 12.15, "lon": -86.27}]


@pytest.fixture
def mundo(con_redes):
    return con_redes, usuarios.identificar(con_redes, usuario_windows="jgutierrez"), usuarios.identificar(con_redes, persona_visor="lector")


def test_the_points_go_to_a_folder_of_mis_lugares_with_their_routes(mundo):
    con, admin, lector = mundo
    r = METODOS["lugares_desde_masivas"](con, lector, nombre="  Visita   norte ", puntos=PUNTOS, con_rutas=True)             # any user has their own places
    assert (r["puntos"], r["rutas"]) == (3, 2)
    elementos = [e for e in r["elementos"] if True]
    marcas = sorted((e["nombre"], e["tipo"], (e["estilo"].get("icono") or {}).get("id")) for e in elementos if e["tipo"] == "marca")
    assert marcas == [("Casa Pérez", "marca", "1"), ("Sin conexión", "marca", "3"), ("Tienda Ana", "marca", "2")]
    assert sorted(e["nombre"] for e in elementos if e["tipo"] == "ruta") == ["Ruta Casa Pérez", "Ruta Tienda Ana"]
    assert {c["nombre"] for c in r["carpetas"]} == {"Visita norte", "Rutas"}
    sin = METODOS["lugares_desde_masivas"](con, lector, nombre="", puntos=PUNTOS[:1])
    assert sin["rutas"] == 0 and any(c["nombre"] == "Coordenadas masivas" for c in sin["carpetas"])
    for malo in ([], None, [{"lat": "x", "lon": 1}], [{"lat": 99, "lon": 1}]):
        with pytest.raises(db.DatosInvalidos):
            METODOS["lugares_desde_masivas"](con, lector, nombre="x", puntos=malo)


def test_the_points_become_a_project_with_clients_the_connection_and_the_routes(mundo):
    con, admin, lector = mundo
    r = METODOS["proyecto_desde_masivas"](con, admin, nombre="Visita norte", puntos=PUNTOS)
    assert r["proyecto"]["codigo"] == "MAS-VISITA-NORTE" and r["resumen"] == {"clientes": 3, "mufas": 1, "cables": 2, "sin_conexion": ["Sin conexión"], "ids": 6}
    pid = r["proyecto"]["id"]
    el = {f["nombre"]: f for f in con.execute("SELECT nombre, tipo, codigo, subtipo FROM redes.elementos WHERE proyecto_id = %s AND borrado_en IS NULL", (pid,)).fetchall()}
    assert set(el) == {"Casa Pérez", "Tienda Ana", "Sin conexión", "MF PRUEBA"} and el["Casa Pérez"]["codigo"] == "C12345" and el["MF PRUEBA"]["tipo"] == "mufa"
    cables = {f["nombre"]: f for f in con.execute("SELECT nombre, clase, tipo, hilos, longitud_m, metraje FROM redes.cables WHERE proyecto_id = %s AND borrado_en IS NULL", (pid,)).fetchall()}
    assert set(cables) == {"Acometida Casa Pérez", "Acometida Tienda Ana"} and cables["Acometida Casa Pérez"]["tipo"] == "FLAT" and cables["Acometida Tienda Ana"]["tipo"] == "ADSS"
    assert all(c["clase"] == "acometida" and c["hilos"] == 2 and c["longitud_m"] > 100 and c["metraje"]["metraje_m"] >= c["longitud_m"] for c in cables.values())
    assert METODOS["proyecto_desde_masivas"](con, admin, nombre="Visita norte", puntos=PUNTOS)["proyecto"]["codigo"] == "MAS-VISITA-NORTE-2"
    with pytest.raises(db.SinPermiso):
        METODOS["proyecto_desde_masivas"](con, lector, nombre="X", puntos=PUNTOS)                      # only editors create projects
