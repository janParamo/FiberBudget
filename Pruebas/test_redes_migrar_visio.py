"""Visio → interactive mufas migration, phase 2 (docs/migracion_visio/PLAN.md): `migrar_visio.py` over a
small synthetic `red_fibra` + KMZ in the throw-away database — the simulation writes nothing, the real
run creates the regional base projects with elements (ODF, reserve, unplaced) and cables (fibre count
from the Visio, by continuity or by the KMZ line text, or unknown), a second run is idempotent, the
--zona / --solo-mufas filters, the Excel report, and the editor working with that data."""
import pytest
from openpyxl import load_workbook

from redes_base import base_redes, con_redes  # noqa: F401  (fixtures)
from visor_rutas.redes import (diagrama, migracion_tablero, migrar_visio, migrar_visio_contenido, mufa, mufas_existentes, proyectos, red,
                               usuarios)

REGLAS = migrar_visio.reglas({"migracion_visio": {"proyectos_base": {
    "defecto": "Red existente – Managua", "zonas": {"Red existente – León": ["Red Leon"]}}}})

# id, nombre, tipo, lat, lon, fuente_coord, zona
NODOS = [(1, "Mufa A", "MUFA", 12.10, -86.30, "visio", "Red Managua"),
         (2, "Mufa B", "MUFA", 12.11, -86.30, "visio", "Red Managua"),
         (3, "Mufa C", "MUFA", 12.12, -86.30, "visio", "Red Managua"),
         (4, "ODF Central", "ODF", 12.09, -86.30, "excel", "Red Managua"),
         (5, "Mufa Perdida", "MUFA", None, None, None, "Red Managua"),
         (6, "Mufa A", "MUFA", 12.10, -86.30, "visio", "Red Managua"),          # repeated in another Visio
         (7, "C12345", "CLIENTE", 12.125, -86.30, "visio", "Red Managua"),
         (8, "Mufa Leon", "MUFA", 12.43, -86.88, "visio", "Red Leon"),
         (9, "Reserva 1", "RESERVA", 12.13, -86.30, "visio", "Red Managua")]
# id, a, b, hilos, fuente, texto
ENLACES = [(1, 4, 1, 48, "excel", None),
           (2, 1, 2, 24, "visio", "CABLE 24H 12X2"),
           (3, 2, 3, None, "visio", None),          # continuity: Mufa B only knows the 24 of A–B
           (4, 3, 9, None, "visio", None),          # the KMZ line says FO 12
           (5, 1, 5, 12, "visio", None),            # an unplaced end
           (6, 3, 7, None, "visio", None),          # a client: not migrated
           (7, 8, 3, None, "visio", None)]          # León → Managua: unknown, between regions


@pytest.fixture
def visio(con_redes):
    con = con_redes
    con.execute("DROP SCHEMA IF EXISTS red_fibra CASCADE")
    con.execute("CREATE SCHEMA red_fibra")
    con.execute("""CREATE TABLE red_fibra.nodos (id int PRIMARY KEY, nombre text, nombre_normalizado text, tipo text, propietario text,
                   lat float8, lon float8, fuente_coord text, zona text, notas text, requiere_revision boolean DEFAULT false, motivo_revision text)""")
    con.execute("CREATE TABLE red_fibra.archivos (id int PRIMARY KEY, ruta_relativa text, md5 text)")
    con.execute("CREATE TABLE red_fibra.nodo_apariciones (nodo_id int, archivo_id int, rol text)")
    con.execute("""CREATE TABLE red_fibra.enlaces (id int PRIMARY KEY, nodo_a_id int, nodo_b_id int, hilos_cable int, longitud_m float8,
                   hilos_usados int, enlace_logico boolean, fuente text, confianza text, texto_original text)""")
    con.execute("""CREATE TABLE red_fibra.servicios (id serial PRIMARY KEY, nodo_id int, archivo_id int, tipo text, codigo_cliente text,
                   id_servicio text, descripcion text)""")
    con.execute("CREATE TABLE red_fibra.textos_crudos (id serial PRIMARY KEY, archivo_id int, texto text)")
    con.execute("""CREATE TABLE red_fibra.anillo_tramos (id serial PRIMARY KEY, hoja text, enlace text, orden int, nodo_id int, en_mufa text,
                   de_mufa text, hilo_entrante text, hacia_mufa text, hilo_saliente text)""")
    for n in NODOS:
        con.execute("INSERT INTO red_fibra.nodos (id, nombre, nombre_normalizado, tipo, lat, lon, fuente_coord, zona) "
                    "VALUES (%s, %s, upper(%s), %s, %s, %s, %s, %s)", (n[0], n[1], n[1], *n[2:]))
    con.execute("INSERT INTO red_fibra.archivos VALUES (1, 'Red Managua/Anillo 1.vsdx', 'abc')")
    con.execute("INSERT INTO red_fibra.nodo_apariciones VALUES (1, 1, 'mufa'), (2, 1, 'mufa')")
    for e in ENLACES:
        con.execute("INSERT INTO red_fibra.enlaces (id, nodo_a_id, nodo_b_id, hilos_cable, fuente, texto_original) "
                    "VALUES (%s, %s, %s, %s, %s, %s)", e)
    con.execute("DROP TABLE IF EXISTS public.red_puntos, public.red_lineas")
    con.execute("CREATE TABLE public.red_puntos (id int PRIMARY KEY, nombre text, categoria text, geom geometry(Point, 4326))")
    con.execute("CREATE TABLE public.red_lineas (id int PRIMARY KEY, nombre text, descripcion text, tipo text, geom geometry(LineString, 4326))")
    con.execute("INSERT INTO public.red_puntos VALUES (1, 'MUFA A KMZ', 'mufa', ST_SetSRID(ST_MakePoint(-86.30005, 12.10005), 4326))")
    # Mufa C → Reserva 1, with a small bend; and a client drop that must be ignored
    con.execute("INSERT INTO public.red_lineas VALUES (1, 'FO 12 RUTA NORTE', NULL, NULL, "
                "ST_GeomFromText('LINESTRING(-86.30 12.12, -86.3008 12.125, -86.30 12.13)', 4326))")
    con.execute("INSERT INTO public.red_lineas VALUES (2, 'C12345', 'CLIENTE', NULL, "
                "ST_GeomFromText('LINESTRING(-86.30 12.12, -86.30 12.13)', 4326))")
    admin = usuarios.identificar(con, usuario_windows="jgutierrez")
    yield con, admin
    con.execute("DROP SCHEMA IF EXISTS red_fibra CASCADE")
    con.execute("DROP TABLE IF EXISTS public.red_puntos, public.red_lineas")


def _corrida(con):
    return con.execute("INSERT INTO redes.migracion_visio_corridas (modo) VALUES ('real') RETURNING id").fetchone()["id"]


def _migrar(con, **filtros):
    plan = migrar_visio.planificar(con, con, REGLAS, **filtros)
    return plan, migrar_visio.aplicar(con, plan, None, _corrida(con))


def test_simulation_writes_nothing_and_reports(visio, tmp_path):
    con, _ = visio
    plan = migrar_visio.planificar(con, con, REGLAS)
    assert con.execute("SELECT count(*) AS n FROM redes.elementos").fetchone()["n"] == 0
    assert con.execute("SELECT count(*) AS n FROM redes.proyectos").fetchone()["n"] == 0
    assert sorted(e["nombre"] for e in plan.elementos) == ["Mufa A", "Mufa B", "Mufa C", "Mufa Leon", "Mufa Perdida", "ODF Central", "Reserva 1"]
    motivos = {o["nombre"]: o["motivo"] for o in plan.omitidos}
    assert "repetido" in motivos["Mufa A"]
    assert "sin ubicar" in motivos["Mufa A – Mufa Perdida"]
    assert any("CLIENTE" in m for m in motivos.values())
    ruta = migrar_visio.reporte(plan, tmp_path / "simulacion.xlsx")
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Resumen", "Mufas", "Cables", "No migrado"]
    assert wb["Mufas"].max_row == 1 + 7 and wb["Cables"].max_row == 1 + 5 and wb["No migrado"].max_row == 1 + 3
    assert "SIMULACIÓN" in wb["Resumen"]["A1"].value


def test_real_run_creates_base_projects_elements_cables_and_fibres(visio):
    con, _ = visio
    plan, conteo = _migrar(con)
    assert conteo["proyectos_base"] == 2 and conteo["elementos_creados"] == 7 and conteo["cables_creados"] == 5
    base = {f["nombre"]: f for f in con.execute("SELECT id, nombre, codigo, es_red_base, zona FROM redes.proyectos").fetchall()}
    assert set(base) == {"Red existente – Managua", "Red existente – León"}
    assert all(p["es_red_base"] for p in base.values()) and base["Red existente – León"]["codigo"] == "RED-LEON"
    el = {f["nombre"]: f for f in con.execute(
        "SELECT nombre, tipo, subtipo, ubicacion IS NULL AS sin, ubicacion_fuente, fuente, confianza, ref_externa, ref_visio, proyecto_id "
        "FROM redes.elementos").fetchall()}
    assert (el["ODF Central"]["tipo"], el["ODF Central"]["fuente"], el["ODF Central"]["confianza"]) == ("odf", "excel_anillos", "alta")
    assert (el["Reserva 1"]["tipo"], el["Reserva 1"]["subtipo"]) == ("mufa", "reserva")
    assert el["Mufa Perdida"]["sin"] and el["Mufa Perdida"]["ubicacion_fuente"] == "sin_ubicar"
    assert el["Mufa A"]["ref_externa"]["kmz"]["nombre"] == "MUFA A KMZ"
    assert el["Mufa A"]["ref_visio"]["archivos"][0]["ruta"] == "Red Managua/Anillo 1.vsdx"
    assert el["Mufa Leon"]["proyecto_id"] == base["Red existente – León"]["id"]
    cab = {f["nombre"]: f for f in con.execute(
        "SELECT nombre, hilos, configuracion, configuracion_estimada, clase, fuente, geom_fuente, color, proyecto_id, "
        "jsonb_array_length(vertices) AS nv, ref_visio, (SELECT count(*) FROM redes.hilos h WHERE h.cable_id = c.id) AS nh "
        "FROM redes.cables c").fetchall()}
    odf = cab["ODF Central – Mufa A"]
    assert (odf["hilos"], odf["clase"], odf["fuente"], odf["nh"]) == (48, "troncal", "excel_anillos", 48)
    ab = cab["Mufa A – Mufa B"]                                     # the Visio text says 12X2: not the 6x4 rule
    assert (ab["hilos"], ab["configuracion"], ab["configuracion_estimada"]) == (24, "12x2", False)
    bc = cab["Mufa B – Mufa C"]
    assert (bc["hilos"], bc["configuracion"], bc["configuracion_estimada"], bc["ref_visio"]["hilos_origen"]) == (24, "6x4", True, "continuidad_un_lado")
    cr = cab["Mufa C – Reserva 1"]
    assert (cr["hilos"], cr["ref_visio"]["hilos_origen"], cr["geom_fuente"], cr["fuente"]) == (12, "texto_kmz", "kmz", "estimado")
    assert cr["nv"] >= 3                                            # follows the KMZ line, not a straight line
    lc = cab["Mufa Leon – Mufa C"]
    assert (lc["hilos"], lc["configuracion"], lc["color"], lc["nh"], lc["geom_fuente"]) == (None, None, "#9e9e9e", 0, "recta")
    assert lc["proyecto_id"] == base["Red existente – León"]["id"]
    assert conteo["hilos_creados"] == 48 + 24 + 24 + 12
    log = {(f["entidad"], f["accion"]): f["n"] for f in con.execute(
        "SELECT entidad, accion, count(*) AS n FROM redes.migracion_visio_log GROUP BY 1, 2").fetchall()}
    assert log == {("nodo", "crear"): 7, ("enlace", "crear"): 5, ("nodo", "omitir"): 1, ("cable", "omitir"): 2}


def test_second_run_is_idempotent(visio):
    con, _ = visio
    _migrar(con)
    _, conteo = _migrar(con)
    assert conteo.get("elementos_creados", 0) == 0 and conteo.get("cables_creados", 0) == 0
    assert (conteo["elementos_ya_estaban"], conteo["cables_ya_estaban"], conteo["proyectos_base"]) == (7, 5, 2)
    assert con.execute("SELECT count(*) AS n FROM redes.elementos").fetchone()["n"] == 7
    assert con.execute("SELECT count(*) AS n FROM redes.proyectos").fetchone()["n"] == 2


def test_reimport_updates_untouched_and_never_overwrites_validated(visio, tmp_path):
    """Phase 7: the Visio changed. Untouched data follows it, what a person validated / edited is only
    reported, and what the Visio no longer gives is kept."""
    con, _ = visio
    _migrar(con)
    cab = lambda n: con.execute("SELECT id, hilos, estado_migracion, "
                                "(SELECT count(*) FROM redes.hilos h WHERE h.cable_id = c.id AND h.borrado_en IS NULL) AS nh "
                                "FROM redes.cables c WHERE nombre = %s AND borrado_en IS NULL", (n,)).fetchone()
    con.execute("UPDATE redes.elementos SET estado_migracion = 'validada' WHERE nombre = 'Mufa C'")       # a person validated it
    con.execute("UPDATE red_fibra.enlaces SET hilos_cable = 48, texto_original = NULL WHERE id = 2")      # Mufa A – Mufa B: 24 → 48
    con.execute("UPDATE red_fibra.nodos SET requiere_revision = true, motivo_revision = 'VALIDAR' WHERE id = 2")
    con.execute("DELETE FROM red_fibra.enlaces WHERE id = 7")                                            # León – Mufa C is gone
    plan = migrar_visio.planificar(con, con, REGLAS)
    dif = {(d["entidad"], d["nombre"]): d for d in migrar_visio.comparar(con, plan)}
    assert not dif[("enlace", "Mufa A – Mufa B")]["protegido"] and dif[("enlace", "Mufa B – Mufa C")]["protegido"]
    assert dif[("enlace", "Mufa Leon – Mufa C")]["desaparecido"]
    assert con.execute("SELECT count(*) AS n FROM redes.cables WHERE hilos = 24").fetchone()["n"] == 2     # simulating wrote nothing
    conteo = migrar_visio.aplicar(con, plan, None, _corrida(con))
    assert (conteo["cables_actualizados"], conteo["elementos_actualizados"], conteo["diferencias_sin_aplicar"]) == (1, 1, 1)
    assert conteo["cables_ya_no_en_visio"] == 1 and conteo.get("cables_creados", 0) == 0
    ab, bc = cab("Mufa A – Mufa B"), cab("Mufa B – Mufa C")
    assert (ab["hilos"], ab["nh"], ab["estado_migracion"]) == (48, 48, "automatica")
    assert (bc["hilos"], bc["nh"]) == (24, 24)                                                             # its mufa was validated
    assert cab("Mufa Leon – Mufa C") is not None                                                           # never deleted
    assert con.execute("SELECT estado_migracion FROM redes.elementos WHERE nombre = 'Mufa B'").fetchone()["estado_migracion"] == "en_revision"
    assert con.execute("SELECT estado_migracion FROM redes.elementos WHERE nombre = 'Mufa C'").fetchone()["estado_migracion"] == "validada"
    assert con.execute("SELECT count(*) AS n FROM redes.historial WHERE entidad = 'cable' AND cambio ? 'hilos'").fetchone()["n"] >= 1
    log = [(f["accion"], f["resultado"]) for f in con.execute(
        "SELECT accion, resultado FROM redes.migracion_visio_log WHERE accion IN ('actualizar', 'conflicto')").fetchall()]
    assert sum(a == "actualizar" for a, _ in log) == 2 and sum(a == "conflicto" for a, _ in log) == 2
    assert any("no se toca" in r and "hilos: 24 → 48" in r for a, r in log)
    # a person edits the cable Mufa A – Mufa B: the next change of the Visio is only reported
    con.execute("UPDATE redes.cables SET estado_migracion = 'editada' WHERE nombre = 'Mufa A – Mufa B'")
    con.execute("UPDATE red_fibra.enlaces SET hilos_cable = 96 WHERE id = 2")
    conteo = migrar_visio.aplicar(con, migrar_visio.planificar(con, con, REGLAS), None, _corrida(con))
    assert conteo.get("cables_actualizados", 0) == 0 and cab("Mufa A – Mufa B")["hilos"] == 48
    # the report has the sheet
    plan = migrar_visio.planificar(con, con, REGLAS)
    plan.diferencias = migrar_visio.comparar(con, plan)
    assert "Diferencias" in load_workbook(migrar_visio.reporte(plan, tmp_path / "dif.xlsx")).sheetnames
    # a partial run cannot tell what is gone
    assert not any(d.get("desaparecido") for d in migrar_visio.comparar(con, migrar_visio.planificar(con, con, REGLAS, zona="León")))


def test_zone_and_list_filters_then_the_rest(visio):
    con, _ = visio
    plan = migrar_visio.planificar(con, con, REGLAS, zona="León")
    assert [e["nombre"] for e in plan.elementos] == ["Mufa Leon"] and plan.cables == [] and plan.omitidos == []
    plan, conteo = _migrar(con, solo={"MUFA A", "MUFA B"})
    assert sorted(e["nombre"] for e in plan.elementos) == ["Mufa A", "Mufa B"]
    assert [c["nombre"] for c in plan.cables] == ["Mufa A – Mufa B"] and conteo["cables_creados"] == 1
    _, conteo = _migrar(con)                                        # the full run completes what is missing
    assert (conteo["elementos_creados"], conteo["elementos_ya_estaban"]) == (5, 2)
    assert (conteo["cables_creados"], conteo["cables_ya_estaban"]) == (4, 1)


def test_trunk_mufa_goes_to_its_region_and_existing_data_moves(visio):
    con, admin = visio
    # a mufa of a trunk route folder (no region), 150 m from León's centre, with a cable to Mufa C
    con.execute("INSERT INTO red_fibra.nodos (id, nombre, nombre_normalizado, tipo, lat, lon, fuente_coord, zona) "
                "VALUES (10, 'Mufa Troncal', 'MUFA TRONCAL', 'MUFA', 12.431, -86.881, 'visio', 'Mufa Metrocentro hacia nodo Leon')")
    con.execute("INSERT INTO red_fibra.enlaces (id, nodo_a_id, nodo_b_id, hilos_cable, fuente) VALUES (8, 10, 3, 12, 'visio')")
    sin_regla = {**REGLAS, "troncales_a_region_km": 0}
    plan = migrar_visio.planificar(con, con, sin_regla)
    assert next(e for e in plan.elementos if e["nombre"] == "Mufa Troncal")["proyecto"] == "Red existente – Managua"
    migrar_visio.aplicar(con, plan, None, _corrida(con))
    plan = migrar_visio.planificar(con, con, REGLAS)                 # the rule on: León
    troncal = next(e for e in plan.elementos if e["nombre"] == "Mufa Troncal")
    assert troncal["proyecto"] == "Red existente – León" and "troncal" in troncal["avisos"][0]
    assert next(c for c in plan.cables if c["nombre"] == "Mufa Troncal – Mufa C")["proyecto"] == "Red existente – León"
    # it was migrated in Managua: the run moves it (and its cable) instead of creating it again
    conteo = migrar_visio.aplicar(con, plan, None, _corrida(con))
    assert (conteo["elementos_movidos"], conteo["cables_movidos"], conteo.get("elementos_creados", 0)) == (1, 1, 0)
    lugar = lambda tabla, n: con.execute(f"SELECT p.nombre FROM redes.{tabla} x JOIN redes.proyectos p ON p.id = x.proyecto_id "
                                         "WHERE x.nombre = %s", (n,)).fetchone()["nombre"]
    assert lugar("elementos", "Mufa Troncal") == lugar("cables", "Mufa Troncal – Mufa C") == "Red existente – León"
    assert con.execute("SELECT count(*) AS n FROM redes.elementos WHERE nombre = 'Mufa Troncal'").fetchone()["n"] == 1
    assert con.execute("SELECT count(*) AS n FROM redes.historial WHERE entidad = 'elemento' AND cambio ? 'proyecto'").fetchone()["n"] == 2
    conteo = migrar_visio.aplicar(con, migrar_visio.planificar(con, con, REGLAS), None, _corrida(con))
    assert "elementos_movidos" not in conteo and "cables_movidos" not in conteo
    # what a user edited is not moved back, only reported
    con.execute("UPDATE redes.elementos SET estado_migracion = 'editada' WHERE nombre = 'Mufa Troncal'")
    conteo = migrar_visio.aplicar(con, migrar_visio.planificar(con, con, sin_regla), None, _corrida(con))
    assert conteo["elementos_no_movidos_editados"] == 1 and lugar("elementos", "Mufa Troncal") == "Red existente – León"


CABECERA = "SPL 1X8 OLT ALAMO 0/1/0 -18.5 dBm"


@pytest.fixture
def visio_con_contenido(visio):
    """Phase 3 data: a splitter with its port list in Mufa B, clients, drops in the KMZ and a ring."""
    con, admin = visio
    servicios = [(2, "splitter", None, None, CABECERA), (2, "splitter", None, None, "3- SPL 1X8 MF LEJANA"),
                 (2, "splitter", None, None, "SPL CHINA MALL"), (3, "splitter", None, None, "SPL 1X6 RARO"),
                 (2, "cliente", "C11111", "ID123456", "1- C11111 Ana Perez ID123456"),
                 (2, "cliente", "C22222", "ID654321", "C22222 Luis Gomez/ID654321/-21.5dBm"),
                 (2, "cliente", "C22222", None, "C22222 Luis Gomez"),                         # same client, another Visio
                 (3, "otro", None, None, "DE MF B A LA MF C")]
    for n, tipo, codigo, ids, texto in servicios:
        con.execute("INSERT INTO red_fibra.servicios (nodo_id, archivo_id, tipo, codigo_cliente, id_servicio, descripcion) "
                    "VALUES (%s, 1, %s, %s, %s, %s)", (n, tipo, codigo, ids, texto))
    con.execute("INSERT INTO red_fibra.textos_crudos (archivo_id, texto) VALUES (1, %s), (1, %s)",
                (f"{CABECERA}\n1- C11111 Ana Perez ID123456\n2- LIBRE\n3- SPL 1/16 MF Lejana\n4-",
                 "DE MF B A LA MF C\n29- SANGRIA X\n30- SPL CASA"))                        # a list of fibres, not of ports
    # drops: from Mufa B (12.11, -86.30) to each house
    for i, (codigo, la, lo) in enumerate((("C11111 Ana Perez", 12.112, -86.302), ("C22222 Luis Gomez", 12.109, -86.299)), 10):
        con.execute("INSERT INTO public.red_lineas VALUES (%s, %s, 'CLIENTE AZUL', NULL, "
                    "ST_GeomFromText(%s, 4326))", (i, codigo, f"LINESTRING(-86.30001 12.11001, {lo} {la})"))
    for orden, nodo, de, entra, hacia, sale in ((1, 1, "ODF Central", "HILOS 1 / 2", "Mufa B", "HILOS 3 / 4"),
                                                (2, 2, "Mufa A", "HILOS 3 / 4", "Mufa C", "HILOS 5 / 6"),
                                                (3, 3, "Mufa B", "HILOS 5 / 6", "Reserva 1", "HILOS 30 / 31")):
        con.execute("INSERT INTO red_fibra.anillo_tramos (hoja, enlace, orden, nodo_id, en_mufa, de_mufa, hilo_entrante, hacia_mufa, hilo_saliente) "
                    "VALUES ('ANILLO T', 'A-C', %s, %s, '', %s, %s, %s, %s)", (orden, nodo, de, entra, hacia, sale))
    return con, admin


def _migrar_todo(con):
    plan = migrar_visio.planificar(con, con, REGLAS)
    cont = migrar_visio_contenido.planificar_contenido(con, con, plan)
    corrida = _corrida(con)
    return plan, cont, {**migrar_visio.aplicar(con, plan, None, corrida), **migrar_visio_contenido.aplicar_contenido(con, cont, None, corrida)}


def test_phase3_splitters_ports_clients_and_ring_splices(visio_con_contenido, tmp_path):
    con, admin = visio_con_contenido
    plan, cont, conteo = _migrar_todo(con)
    assert (conteo["splitters_creados"], conteo["puertos_con_cliente"], conteo["clientes_de_mufa"], conteo["clientes_en_mapa"]) == (1, 2, 1, 2)
    # Mufa C – Reserva 1 was deduced as 12 fibres (KMZ text), but the ring splices fibres 30 and 31: it becomes 48
    assert (conteo["empalmes_creados"], conteo.get("empalmes_fuera_de_rango", 0), conteo["cables_hilos_minimo_del_anillo"]) == (6, 0, 1)
    cr = con.execute("SELECT hilos, configuracion, configuracion_estimada, fuente, confianza, ref_visio ->> 'hilos_origen' AS o, "
                     "(SELECT count(*) FROM redes.hilos h WHERE h.cable_id = c.id AND h.borrado_en IS NULL) AS nh "
                     "FROM redes.cables c WHERE nombre = 'Mufa C – Reserva 1'").fetchone()
    assert (cr["hilos"], cr["configuracion"], cr["configuracion_estimada"], cr["fuente"], cr["confianza"], cr["o"], cr["nh"]) == (
        48, "6x8", True, "estimado", "baja", "minimo_del_anillo", 48)
    motivos = " | ".join(o["motivo"] for o in cont.omitidos)
    assert "no dice el tamaño" in motivos and "1x6 no existe" in motivos and "sin un splitter con tamaño" in motivos
    mufa_b = con.execute("SELECT id FROM redes.elementos WHERE nombre = 'Mufa B'").fetchone()["id"]
    s = con.execute("SELECT * FROM redes.splitters WHERE mufa_id = %s", (mufa_b,)).fetchone()
    assert (s["salidas"], s["x"], s["fuente"], s["confianza"]) == (8, 470, "texto_diagrama", "media")
    assert (s["propiedades"]["olt"], s["propiedades"]["puerto_olt"], s["propiedades"]["potencia_dbm"]) == ("ALAMO", "0/1/0", -18.5)
    clientes = {f["codigo"]: f for f in con.execute(
        "SELECT codigo, nombre, ST_Y(ubicacion) AS lat, ST_X(ubicacion) AS lon, fuente FROM redes.elementos WHERE tipo = 'cliente'").fetchall()}
    assert set(clientes) == {"C11111", "C22222"} and clientes["C11111"]["fuente"] == "kmz"
    assert (round(clientes["C11111"]["lat"], 3), round(clientes["C11111"]["lon"], 3)) == (12.112, -86.302)   # the far end of the drop
    # what the mufa editor reads
    vista = mufa.leer(con, admin, mufa_id=mufa_b)
    servicios = vista["splitters"][0]["servicios"]
    assert set(servicios) == {1, 3} and servicios[1]["codigo_cliente"] == "C11111" and servicios[1]["id_servicio"] == "ID123456"
    assert servicios[1]["elemento_id"] is not None and servicios[3] == {**servicios[3], "codigo_cliente": None, "nombre": "SPL 1/16 MF Lejana"}
    assert [(c["codigo_cliente"], c["nombre"], c["id_servicio"]) for c in vista["clientes"]] == [("C22222", "Luis Gomez", "ID654321")]
    anillo = [c for c in vista["conexiones"] if c["fuente"] == "excel_anillos"]
    assert len(anillo) == 2 and all(c["confianza"] == "alta" for c in anillo)
    numeros = lambda cid: {f["numero"] for f in con.execute(
        "SELECT h.numero FROM redes.conexiones x JOIN redes.hilos h ON h.id IN (x.a_hilo_id, x.b_hilo_id) WHERE x.mufa_id = %s", (cid,)).fetchall()}
    assert numeros(mufa_b) == {3, 4, 5, 6}
    # a second run repeats nothing
    _, _, conteo = _migrar_todo(con)
    assert conteo["splitters_ya_estaban"] == 1 and conteo["empalmes_ya_conectados"] == 6
    assert not any(conteo.get(k) for k in ("splitters_creados", "puertos_con_cliente", "clientes_de_mufa", "clientes_en_mapa", "empalmes_creados",
                                           "cables_hilos_minimo_del_anillo"))
    assert con.execute("SELECT count(*) AS n FROM redes.conexiones").fetchone()["n"] == 6
    # phase 7: a re-import does not undo the minimum (the plan deduces 12 for that cable, the base has 48)
    assert not [d for d in migrar_visio.comparar(con, migrar_visio.planificar(con, con, REGLAS)) if d["nombre"] == "Mufa C – Reserva 1"]
    # the report has the phase 3 sheets
    wb = load_workbook(migrar_visio.reporte(plan, tmp_path / "r.xlsx", contenido=cont))
    assert {"Splitters", "Clientes", "Empalmes"} <= set(wb.sheetnames) and wb["Empalmes"].max_row == 1 + 6


def test_ring_fibre_numbers_set_the_size_of_cables_without_a_count(visio_con_contenido):
    con, _ = visio_con_contenido
    plan = migrar_visio.planificar(con, con, REGLAS)
    migrar_visio.aplicar(con, plan, None, _corrida(con))
    # Mufa B – Mufa C: forget its count (the Visio never gave it); the ring splices its fibres 5 and 6 -> at least 6
    con.execute("UPDATE redes.cables SET hilos = NULL WHERE nombre = 'Mufa B – Mufa C'")
    cont = migrar_visio_contenido.planificar_contenido(con, con, plan)
    conteo = migrar_visio_contenido.aplicar_contenido(con, cont, None, _corrida(con))
    bc = con.execute("SELECT hilos, configuracion, configuracion_estimada FROM redes.cables WHERE nombre = 'Mufa B – Mufa C'").fetchone()
    assert (bc["hilos"], bc["configuracion"], bc["configuracion_estimada"]) == (6, "1x6", True)
    assert conteo["cables_hilos_minimo_del_anillo"] == 2 and conteo["empalmes_creados"] == 6      # B–C and C–Reserva 1
    # a person validated a cable's mufa: its size is not touched (a cable of Mufa C, size 12 -> would need 48)
    con.execute("UPDATE redes.cables SET hilos = 12, estado_migracion = 'editada' WHERE nombre = 'Mufa C – Reserva 1'")
    con.execute("DELETE FROM redes.conexiones WHERE mufa_id = (SELECT id FROM redes.elementos WHERE nombre = 'Mufa C')")
    conteo = migrar_visio_contenido.aplicar_contenido(con, cont, None, _corrida(con))
    assert con.execute("SELECT hilos FROM redes.cables WHERE nombre = 'Mufa C – Reserva 1'").fetchone()["hilos"] == 12
    assert conteo.get("empalmes_fuera_de_rango", 0) == 2


def test_splitter_text_parsing():
    leer = migrar_visio_contenido.leer_splitter
    assert leer("SPL 1X8 SECUNDARIO OLT SD PUERTO 0/2/9 -20.42 dBm")["propiedades"] | {"texto_visio": ""} == {
        "olt": "SD", "puerto_olt": "0/2/9", "potencia_dbm": -20.42, "texto_visio": ""}
    assert leer("SPL DEPENDE ¼ ENTRADA CONTEMPO")["salidas"] == 4
    assert leer("SPL DEPENDE ¼ ENTRADA CONTEMPO")["propiedades"]["depende_de"] == "ENTRADA CONTEMPO"
    assert "olt" not in leer("SPL 1X4 PRINCIPAL OLT PUERTO 0/0/3")["propiedades"]
    assert leer("SPL 1X2 70/30")["relacion"] == "70/30" and leer("SPL 1/16 P0/1/8")["salidas"] == 16
    nombre = migrar_visio_contenido.nombre_splitter
    assert nombre("SPL SECUNDARIO 1X8 MUFA LA JOYA OLT ALAMO 0/ 0/2") == "SPL SECUNDARIO 1X8 MUFA LA JOYA"
    assert nombre("SPL 1/8 SECUNDARIO -17 dBm OLT ALAMO 0/ 0/2") == "SPL 1/8 SECUNDARIO"
    assert nombre("SPL 1X8 -17.85 dBm OLT ENTHEOS P0/2/12") == "SPL 1X8" and nombre("SPL 1X4 PRINCIPAL PUERTO 0/0/3 DEL OLT") == "SPL 1X4 PRINCIPAL"
    servicio = migrar_visio_contenido.leer_servicio
    assert servicio("C35896 Roger David Vallecillo Lanzas/ID496760/-24.59dBm") == {
        "codigo_cliente": "C35896", "nombre": "Roger David Vallecillo Lanzas", "id_servicio": "ID496760"}
    assert servicio("LIBRE") is None and servicio("") is None and servicio("LIBRE VALIDAR POTENCIA") is None


def test_dashboard_and_correction_queue(visio, tmp_path):
    import base64
    con, admin = visio
    _migrar(con)
    t = migracion_tablero.tablero(con, admin)
    assert (t["total"]["mufas"], t["total"]["sin_ubicar"], t["total"]["validadas"], t["total"]["avance"]) == (7, 1, 0, 0)
    assert {z["proyecto"]: z["mufas"] for z in t["zonas"]} == {"Red existente – Managua": 6, "Red existente – León": 1}
    cola = migracion_tablero.correcciones(con, admin)
    primera = cola["correcciones"][0]
    assert (primera["mufa"], primera["motivo"], primera["prioridad"]) == ("Mufa Perdida", "sin_ubicar", 1)
    motivos_c = {f["motivo"] for f in cola["correcciones"] if f["mufa"] == "Mufa C"}
    assert motivos_c == {"hilos_desconocidos", "sin_diagramar", "configuracion_estimada"}
    assert not any(f["mufa"] == "Reserva 1" and f["motivo"] == "sin_diagramar" for f in cola["correcciones"])   # a reserve has no splices
    solo = migracion_tablero.correcciones(con, admin, motivo="hilos_desconocidos")
    assert {f["mufa"] for f in solo["correcciones"]} == {"Mufa C", "Mufa Leon"}
    assert migracion_tablero.correcciones(con, admin, texto="leon")["total"] == len([f for f in cola["correcciones"] if f["mufa"] == "Mufa Leon"])
    # a validated mufa leaves the queue and counts as progress
    mufa_c = con.execute("SELECT id FROM redes.elementos WHERE nombre = 'Mufa C'").fetchone()["id"]
    mufas_existentes.validar(con, admin, mufa_id=mufa_c)
    assert not any(f["mufa"] == "Mufa C" for f in migracion_tablero.correcciones(con, admin)["correcciones"])
    t = migracion_tablero.tablero(con, admin)
    assert (t["total"]["validadas"], t["total"]["avance"]) == (1, round(100 / 7, 1))
    x = migracion_tablero.exportar_correcciones(con, admin, motivo="sin_ubicar")
    ruta = tmp_path / x["nombre"]
    ruta.write_bytes(base64.b64decode(x["datos"]))
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["Tablero", "Correcciones"] and wb["Correcciones"].max_row == 2 and x["filas"] == 1


def test_editor_and_permissions_with_migrated_data(visio):
    con, admin = visio
    _migrar(con)
    managua = con.execute("SELECT id FROM redes.proyectos WHERE nombre = 'Red existente – Managua'").fetchone()["id"]
    lector = usuarios.guardar(con, admin, datos={"nombre": "Ana", "rol": "lector", "persona_visor": "ana"})["usuario"]
    editor = usuarios.guardar(con, admin, datos={"nombre": "Luis", "rol": "editor", "persona_visor": "luis"})["usuario"]
    # base projects are visible to everyone (and editable by editors) without being assigned
    assert {p["nombre"] for p in proyectos.listar(con, lector)["proyectos"]} >= {"Red existente – Managua", "Red existente – León"}
    assert proyectos.puede_ver(con, lector, managua) and not proyectos.puede_editar(con, lector, managua)
    assert proyectos.puede_editar(con, editor, managua)
    r = red.red(con, lector, proyecto_id=managua)
    perdida = next(e for e in r["elementos"] if e["nombre"] == "Mufa Perdida")
    assert perdida["lat"] is None and perdida["ubicacion_fuente"] == "sin_ubicar"
    sin_hilos = con.execute("SELECT id FROM redes.cables WHERE hilos IS NULL").fetchone()["id"]
    mufa_c = next(e for e in r["elementos"] if e["nombre"] == "Mufa C")
    vista = mufa.leer(con, lector, mufa_id=mufa_c["id"])
    assert any(c["id"] == sin_hilos and c["hilos"] is None for c in vista["cables"])
    odf = next(e for e in r["elementos"] if e["tipo"] == "odf")
    assert mufa.leer(con, lector, mufa_id=odf["id"])["mufa"]["nombre"] == "ODF Central"
    assert diagrama.color_de(None, diagrama.COLORES_POR_DEFECTO) == diagrama.COLOR_DESCONOCIDO
    d = diagrama.generar(con, lector, proyecto_id=managua, inicio=f"e:{odf['id']}")
    assert {c["hilos"] for c in d["cables"]} >= {48, 24}
    # an editor moves the cable B–C: its unknown neighbour keeps "unknown", the edit marks the trace as the user's
    cable = next(c for c in r["cables"] if c["nombre"] == "Mufa B – Mufa C")
    assert cable["configuracion"] == "6x4" and cable["configuracion_estimada"] and cable["geom_fuente"] == "recta"
    medio = {"lat": 12.115, "lon": -86.301, "elemento_id": None}
    red.guardar(con, editor, proyecto_id=managua, cambios={"cables": [{**cable, "vertices": [cable["vertices"][0], medio, cable["vertices"][-1]]}]})
    f = con.execute("SELECT geom_fuente, configuracion_estimada, hilos FROM redes.cables WHERE nombre = 'Mufa B – Mufa C'").fetchone()
    assert (f["geom_fuente"], f["configuracion_estimada"], f["hilos"]) == ("usuario", True, 24)
    leon = con.execute("SELECT id FROM redes.proyectos WHERE nombre = 'Red existente – León'").fetchone()["id"]
    desconocido = next(c for c in red.red(con, admin, proyecto_id=leon)["cables"] if c["hilos"] is None)
    red.guardar(con, admin, proyecto_id=leon, cambios={"cables": [{**desconocido, "cruces": 1}]})
    assert con.execute("SELECT hilos FROM redes.cables WHERE id = %s", (desconocido["id"],)).fetchone()["hilos"] is None
    red.guardar(con, admin, proyecto_id=leon, cambios={"cables": [{**desconocido, "hilos": 12, "version": desconocido["version"] + 1}]})
    f = con.execute("SELECT hilos, configuracion_estimada, (SELECT count(*) FROM redes.hilos WHERE cable_id = c.id) AS nh "
                    "FROM redes.cables c WHERE id = %s", (desconocido["id"],)).fetchone()
    assert (f["hilos"], f["configuracion_estimada"], f["nh"]) == (12, False, 12)
