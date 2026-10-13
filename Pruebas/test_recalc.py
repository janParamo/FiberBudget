"""Formula-value caching through LibreOffice (skipped when it is not installed)."""
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

import recalc

SOFFICE = recalc.find_soffice()


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice not installed")
def test_formula_results_are_cached_for_non_recalculating_readers(tmp_path):
    path = tmp_path / "book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Presupuesto"
    ws["A1"], ws["A2"], ws["A3"] = 2, 3, "=A1*A2"
    ws["B1"] = '=CONCATENATE("PRESUPUESTO DE OBRA"," - ","Cliente")'
    ws["B2"] = '=IF(A3>5,A3*0.1,"0")'
    wb.save(path)
    assert load_workbook(path, data_only=True)["Presupuesto"]["A3"].value is None   # nothing cached yet

    assert recalc.cache_formula_values(path, ["Presupuesto"], SOFFICE) == 3

    cached = load_workbook(path, data_only=True)["Presupuesto"]
    assert cached["A3"].value == 6
    assert cached["B1"].value == "PRESUPUESTO DE OBRA - Cliente"
    assert cached["B2"].value == pytest.approx(0.6)
    formulas = load_workbook(path)["Presupuesto"]
    assert formulas["A3"].value == "=A1*A2"   # formulas stay live for Excel
    with zipfile.ZipFile(path) as package:
        assert package.testzip() is None


@pytest.mark.skipif(SOFFICE is None, reason="LibreOffice not installed")
def test_stale_cached_values_from_the_template_are_recomputed(tmp_path):
    """The real template ships with cached zeros; inputs written later must be reflected."""
    from excel_generator import _write_cells_preserving_package

    template = tmp_path / "template.xlsx"
    wb = Workbook()
    wb.active.title = "Presupuesto"
    wb.active["A1"], wb.active["A2"] = 0, "=A1*2"
    wb.save(template)
    recalc.cache_formula_values(template, ["Presupuesto"], SOFFICE)   # caches A2 = 0, like the template
    output = tmp_path / "budget.xlsx"
    _write_cells_preserving_package(template, output, {"xl/worksheets/sheet1.xml": {"A1": 245}})

    recalc.cache_formula_values(output, ["Presupuesto"], SOFFICE)

    assert load_workbook(output, data_only=True)["Presupuesto"]["A2"].value == 490


def test_missing_configured_soffice_falls_back_to_autodetection(tmp_path, monkeypatch):
    missing = str(tmp_path / "no-soffice.exe")
    assert recalc.find_soffice(missing, autodetect_fallback=False) is None
    monkeypatch.setattr(recalc.shutil, "which", lambda name: None)
    monkeypatch.setattr(recalc, "_COMMON_SOFFICE_PATHS", (str(tmp_path / "also-missing.exe"),))
    assert recalc.find_soffice(missing) is None
    found = tmp_path / "soffice.exe"
    found.write_bytes(b"")
    monkeypatch.setattr(recalc, "_COMMON_SOFFICE_PATHS", (str(found),))
    assert recalc.find_soffice(missing) == str(found)   # a stale LIBREOFFICE_PATH still works
