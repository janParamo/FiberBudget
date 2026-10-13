"""budget_note: "Sujeto a disponibilidad de postes de concreto sobre la ruta" on every printed
sheet of every budget, from the API (n8n) and from the viewer."""
from datetime import datetime

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from test_visor_presupuesto import BODY, HEADERS, _campos, entorno  # noqa: F401  (fixture)

NOTA = "Sujeto a disponibilidad de postes de concreto sobre la ruta"


def _celdas_con_nota(xlsx, hojas):
    wb = load_workbook(xlsx)
    return {hoja: [c.coordinate for fila in wb[hoja].iter_rows() for c in fila if c.value == NOTA] for hoja in hojas}


def test_api_and_viewer_budgets_carry_the_note_on_each_page(entorno):
    main, pv, salida, _ = entorno
    rules = main.load_rules(main.settings.rules_path)
    assert rules["budget_note"]["text"] == NOTA
    hojas = [rules["route_sheet"]["sheet"], "Presupuesto"]

    api = TestClient(main.app).post("/presupuesto", json=BODY, headers=HEADERS).json()
    visor = pv.generar(pv.construir_solicitud(_campos()), ahora=datetime(2026, 10, 8, 9, 0, 0))
    for respuesta in (api, visor):
        celdas = _celdas_con_nota(respuesta["excel_path"], hojas)
        assert all(len(c) == 1 for c in celdas.values()), celdas
        assert not any("Budget note" in a for a in respuesta["warnings"])


def test_note_goes_under_the_last_filled_row_in_the_configured_column(entorno):
    main, _, _, _ = entorno
    from excel_generator import _note_assignments
    rules = main.load_rules(main.settings.rules_path)
    wb = load_workbook(main.settings.template_path)
    ws = wb["Presupuesto"]
    ultima = max([c.row for fila in ws.iter_rows() for c in fila if c.value not in (None, "")]
                 + [m.max_row for m in ws.merged_cells.ranges])
    asignaciones, avisos = {}, []
    _note_assignments(wb, ["Presupuesto"], {"text": NOTA, "columns": {"Presupuesto": "E"}, "rows_below_last": 2},
                      asignaciones, avisos)
    assert asignaciones == {"Presupuesto": {f"E{ultima + 2}": NOTA}} and avisos == []
    sin_nota = {}
    _note_assignments(wb, ["Presupuesto"], None, sin_nota, avisos)
    assert sin_nota == {}
    assert rules["budget_note"]["columns"]["Presupuesto"] == "E"
