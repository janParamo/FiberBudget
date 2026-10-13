#!/usr/bin/env python3
"""
Reporte para Planta Externa: cómo se empareja cada mufa del KMZ con la red de los diagramas Visio.

Recorre todas las mufas del KMZ (red_puntos de tipo mufa, PostGIS) y aplica el mismo emparejamiento
que el presupuesto y el visor (origen.match_connection):
  a) mufa Visio con coordenada a <= match_radius_m          b) a <= fallback_radius_m
  c) por nombre en la red Visio (nodos)                     d) por nombre en un diagrama, sin coordenada
  e) sin emparejar
y genera un Excel con una fila por mufa (con filtros) y una hoja "Resumen" con cuántas hay por caso.
Solo lee: no modifica ninguna base.

Uso (desde la carpeta del proyecto, con el mismo .venv):
    .venv\\Scripts\\python.exe red_fibra\\reporte_emparejamiento.py
    .venv\\Scripts\\python.exe red_fibra\\reporte_emparejamiento.py --sqlite --salida output\\emparejamiento.xlsx
    .venv\\Scripts\\python.exe red_fibra\\reporte_emparejamiento.py --limite 50

Red Visio: la copia PostgreSQL (VISIO_DB_*) si responde, si no red_fibra/red_fibra.db (o --sqlite).
"""
import argparse
import datetime as dt
import os
import sqlite3
import sys
from collections import Counter
from pathlib import Path
from typing import Any

RAIZ = Path(__file__).resolve().parent.parent
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

import psycopg  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

import origen  # noqa: E402
import red_fibra_pg  # noqa: E402
from config import get_settings, load_rules  # noqa: E402

CASOS = {
    "a": "Por coordenada (<= radio de emparejamiento)",
    "b": "Por coordenada (<= radio de respaldo)",
    "c": "Por nombre en la red Visio (validar)",
    "d": "Por nombre en diagrama, sin coordenada (validar)",
    "e": "Sin emparejar",
}
DIAGNOSTICOS = {
    "a_mayor_distancia": "Hay mufa Visio con nombre parecido a mayor distancia",
    "sin_coordenada_en_visio": "El nombre aparece en un diagrama sin coordenada",
    "marcada_validar": "El diagrama dice VALIDAR",
    "no_diagramada": "No aparece en ningún diagrama",
    "es_reserva": "Es reserva (se instala mufa nueva)",
}
COLUMNAS = [("ID KMZ", 9), ("Nombre KMZ", 34), ("Tipo KMZ", 12), ("Latitud", 11), ("Longitud", 12),
            ("Resultado", 10), ("Descripción del resultado", 36), ("Nodo Visio emparejado", 34),
            ("Distancia (m)", 12), ("Similitud", 10), ("Archivo del diagrama", 44), ("Inicio del recorrido", 30),
            ("Diagnóstico", 24), ("Casos del diagnóstico", 34), ("Detalle", 80), ("Otros candidatos", 60)]
_KMZ_SQL = """SELECT id, nombre, tipo, ST_Y(geom), ST_X(geom) FROM red_puntos
              WHERE geom IS NOT NULL AND tipo ILIKE 'mufa%' ORDER BY id"""


def mufas_kmz(limite: int | None = None) -> list[tuple]:
    """(id, nombre, tipo, lat, lon) of every KMZ mufa (read-only query)."""
    s = get_settings()
    with psycopg.connect(host=s.db_host, port=s.db_port, dbname=s.db_name, user=s.db_user,
                         password=s.db_password, connect_timeout=5) as conn:
        conn.read_only = True
        filas = conn.execute(_KMZ_SQL).fetchall()
    return filas[:limite] if limite else filas


def _archivo_de_nodo(con: Any, nodo_id: int) -> str | None:
    """Diagram where the Visio node is drawn (its own diagram first)."""
    fila = con.execute(
        """SELECT f.ruta_relativa FROM nodo_apariciones a JOIN archivos f ON f.id = a.archivo_id
           WHERE a.nodo_id = ? AND f.estado = 'ok' AND f.formato <> 'xlsx'
           ORDER BY CASE a.rol WHEN 'principal' THEN 0 WHEN 'vecino' THEN 1 ELSE 2 END, f.ruta_relativa""",
        (nodo_id,)).fetchone()
    return fila[0] if fila else None


def fila_reporte(con: Any, punto: tuple, rules: dict[str, Any], indice: Any) -> list[Any]:
    """One report row for a KMZ mufa (id, nombre, tipo, lat, lon)."""
    kmz_id, nombre, tipo, lat, lon = punto
    m = origen.match_connection(con, origen.ChainPoint(nombre, float(lat), float(lon), tipo), rules, indice)
    emparejado = m.name_match or {}
    archivo = emparejado.get("ruta_relativa")
    if archivo is None and m.node_id is not None:
        archivo = _archivo_de_nodo(con, m.node_id)
    distancia = m.distance_m
    if m.method == "d":
        nodo = m.drawn_name
        inicio = m.node_name or m.direct_node
        distancia = emparejado.get("distancia_diagrama_m")   # nearest georeferenced mufa of that diagram
    else:
        nodo, inicio = m.node_name, None
    diagnostico = m.diagnostico or {}
    candidatos = emparejado.get("candidatos") or (diagnostico.get("candidatos_nodos") or []) + (
        diagnostico.get("candidatos_diagramas") or [])
    otros = "; ".join(
        f"{c.get('nombre') or c.get('nombre_diagrama')}"
        + (f" ({c['archivo']})" if c.get("archivo") else "")
        + (f" {c['distancia_m']:.0f} m" if c.get("distancia_m") is not None else "")
        + f" sim {c.get('similitud', 0):.2f}" for c in candidatos)
    detalle = diagnostico.get("detalle")
    if detalle is None and m.method == "b":
        detalle = "mufa Visio fuera del radio de emparejamiento: validar"
    return [kmz_id, nombre, tipo, round(float(lat), 7), round(float(lon), 7), m.method, CASOS[m.method], nodo,
            None if distancia is None else round(distancia, 1), emparejado.get("similitud"), archivo,
            inicio, DIAGNOSTICOS.get(diagnostico.get("caso"), diagnostico.get("caso")),
            ", ".join(diagnostico.get("casos") or []), detalle, otros or None]


def escribir_excel(filas: list[list[Any]], salida: Path, fuente: str) -> Path:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Mufas"
    hoja.append([c for c, _ in COLUMNAS])
    for fila in filas:
        hoja.append(fila)
    negrita, relleno = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")
    for i, (_, ancho) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=i)
        celda.font, celda.fill = negrita, relleno
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.freeze_panes = "C2"
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}{max(len(filas) + 1, 2)}"
    naranja = PatternFill("solid", fgColor="FFE3C2")
    rojo = PatternFill("solid", fgColor="FFD6D6")
    for r, fila in enumerate(filas, start=2):
        if fila[5] in ("c", "d"):
            hoja.cell(row=r, column=6).fill = naranja
        elif fila[5] == "e":
            hoja.cell(row=r, column=6).fill = rojo

    resumen = libro.create_sheet("Resumen", 0)
    resumen.append(["Emparejamiento de mufas KMZ con la red Visio"])
    resumen["A1"].font = Font(bold=True, size=14)
    resumen.append([f"Generado {dt.datetime.now():%Y-%m-%d %H:%M} · red Visio: {fuente} · {len(filas)} mufas del KMZ"])
    resumen.append([])
    resumen.append(["Resultado", "Descripción", "Mufas", "%"])
    por_caso = Counter(f[5] for f in filas)
    total = max(len(filas), 1)
    for caso, descripcion in CASOS.items():
        resumen.append([caso, descripcion, por_caso.get(caso, 0), round(100 * por_caso.get(caso, 0) / total, 1)])
    resumen.append(["", "Total", len(filas), 100.0 if filas else 0])
    resumen.append([])
    resumen.append(["Diagnóstico (c, d, e)", "Descripción", "Mufas", "%"])
    por_diagnostico = Counter(caso for f in filas for caso in (f[13] or "").split(", ") if caso)
    for caso, descripcion in DIAGNOSTICOS.items():
        resumen.append([caso, descripcion, por_diagnostico.get(caso, 0),
                        round(100 * por_diagnostico.get(caso, 0) / total, 1)])
    resumen.append([])
    resumen.append(["Una mufa puede tener más de un caso de diagnóstico (p. ej. sin coordenada y marcada VALIDAR)."])
    for fila in (4, 4 + len(CASOS) + 3):
        for col in range(1, 5):
            resumen.cell(row=fila, column=col).font, resumen.cell(row=fila, column=col).fill = negrita, relleno
    for col, ancho in zip("ABCD", (26, 52, 10, 8)):
        resumen.column_dimensions[col].width = ancho
    salida.parent.mkdir(parents=True, exist_ok=True)
    libro.save(salida)
    return salida


def generar(salida: Path, usar_sqlite: bool = False, limite: int | None = None) -> Path:
    settings = get_settings()
    rules = load_rules(settings.rules_path)
    puntos = mufas_kmz(limite)
    pg = None if usar_sqlite else red_fibra_pg.connect_or_none()
    if pg is not None:
        con, fuente = origen._PgAdapter(pg), "PostgreSQL (red_fibra)"
    else:
        archivo = Path(settings.red_fibra_db)
        if not archivo.is_file():
            raise SystemExit(f"No se encontró la base de diagramas: {archivo}")
        con = sqlite3.connect(f"{archivo.resolve().as_uri()}?mode=ro", uri=True)
        fuente = f"SQLite ({archivo.name})"
    print(f"{len(puntos)} mufas del KMZ; red Visio: {fuente}")
    try:
        indice = origen._NameIndex(con)
        filas = []
        for i, punto in enumerate(puntos, start=1):
            filas.append(fila_reporte(con, punto, rules, indice))
            if i % 250 == 0:
                print(f"  {i}/{len(puntos)}")
    finally:
        con.close()
        if pg is not None:
            pg.close()
    escribir_excel(filas, salida, fuente)
    resumen = Counter(f[5] for f in filas)
    print("  " + "  ".join(f"{caso}: {resumen.get(caso, 0)}" for caso in CASOS))
    return salida


def main() -> None:
    ap = argparse.ArgumentParser(description="Reporte de emparejamiento de mufas KMZ con la red Visio (solo lectura)")
    ap.add_argument("--salida", default=None,
                    help="Excel de salida (por defecto red_fibra/reporte_emparejamiento_AAAAMMDD_HHMM.xlsx)")
    ap.add_argument("--sqlite", action="store_true", help="Usar red_fibra.db aunque PostgreSQL responda")
    ap.add_argument("--limite", type=int, default=None, help="Solo las primeras N mufas (prueba rápida)")
    args = ap.parse_args()
    os.chdir(RAIZ)                      # .env and business_rules.yaml are found from the project folder
    salida = Path(args.salida) if args.salida else (
        RAIZ / "red_fibra" / f"reporte_emparejamiento_{dt.datetime.now():%Y%m%d_%H%M}.xlsx")
    print(f"Reporte guardado en {generar(salida, args.sqlite, args.limite)}")


if __name__ == "__main__":
    main()
