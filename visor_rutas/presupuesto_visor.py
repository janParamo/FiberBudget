"""Presupuestos desde el visor de rutas: la MISMA generación que POST /presupuesto.

Todo pasa por main._generar_presupuesto_con (punto de conexión, ruta de Google, cantidades,
mapa, plantilla, recalc con LibreOffice), con dos diferencias que no cambian el presupuesto:
  - output_dir = output/visor: los archivos del visor nunca van a output/ ni a budget_index.json
    (ese índice lo escribe solo el endpoint de la API) y no se envía nada;
  - get_route = CacheRutas: la ruta de Google ya pedida para esa coordenada en esta sesión del
    visor se reutiliza (Directions sigue en modo driving).

El formulario se valida con main.TicketRequest, el mismo modelo que valida el body de n8n.
"""
import logging
import math
import re
import threading
import time
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Any

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
from pydantic import ValidationError

import main
import red_fibra_pg
import routing
from excel_generator import _normalized, _safe_file_stem
from presupuesto import seleccionar_tecnologia
from visor_rutas.calculo import CoordenadaInvalida, parsear_coordenadas

log = logging.getLogger(__name__)

CARPETA = "visor"                       # output/visor
SERVICIOS = ("Residencial", "Empresarial", "Corporativo")
MEDIOS = ("FO",)
TECNOLOGIAS = {"": None, "auto": None, "automatico": None, "gpon": "gpon", "metro": "metro"}
HOJAS = ("Presupuesto", "Ruta proyectada")
ESTIMACION_POR_DEFECTO = {"bandwidth": "10 Mbps", "medium": "FO", "servicio_hasta_umbral": "Residencial",
                          "servicio_sobre_umbral": "Empresarial", "cliente": "Estimación visor"}
CAMPO_POSTES_MARCADOS = "postes_existentes_m"   # form key: new poles marked as existing on the map
MAX_POSTES_MARCADOS = 1000


class PresupuestoError(Exception):
    """No se pudo generar el presupuesto (mismo motivo que la API respondería como error)."""


class ErrorFormulario(Exception):
    """Campos no válidos: {campo: mensaje}."""

    def __init__(self, errores: dict[str, str]):
        super().__init__("; ".join(f"{k}: {v}" for k, v in errores.items()))
        self.errores = errores


# --------------------------------------------------------------------------- caché de rutas

class CacheRutas:
    """routing.get_route con memoria por coordenada durante la sesión del visor.

    Clave: origen y destino redondeados a 6 decimales (~10 cm). Los errores no se guardan, así
    que se reintenta. Se llama igual que routing.get_route(origin=..., destination=...).
    """

    def __init__(self, vigencia_s: float = 1800, maximo: int = 300):
        self._vigencia_s = vigencia_s
        self._maximo = maximo
        self._rutas: OrderedDict[tuple, tuple[float, routing.Route]] = OrderedDict()
        self._lock = threading.Lock()
        self.llamadas_google = 0

    @staticmethod
    def _clave(origin: tuple[float, float], destination: tuple[float, float]) -> tuple:
        return tuple(round(float(v), 6) for v in (*origin, *destination))

    def __call__(self, origin: tuple[float, float], destination: tuple[float, float]) -> routing.Route:
        clave = self._clave(origin, destination)
        ahora = time.monotonic()
        with self._lock:
            guardada = self._rutas.get(clave)
            if guardada is not None and ahora - guardada[0] <= self._vigencia_s:
                self._rutas.move_to_end(clave)
                return guardada[1]
        ruta = routing.get_route(origin=origin, destination=destination)   # resolved at call time
        with self._lock:
            self.llamadas_google += 1
            self._rutas[clave] = (ahora, ruta)
            self._rutas.move_to_end(clave)
            while len(self._rutas) > self._maximo:
                self._rutas.popitem(last=False)
        return ruta


# --------------------------------------------------------------------------- archivos

def carpeta_visor() -> Path:
    return main.settings.output_dir / CARPETA


def limpiar_archivos_viejos(carpeta: Path | None = None) -> int:
    """Borra los archivos de output/visor (y su carpeta mapas) más viejos que OUTPUT_RETENTION_DAYS.
    No toca output/ ni budget_index.json."""
    carpeta = carpeta or carpeta_visor()
    limite = time.time() - main.settings.output_retention_days * 24 * 60 * 60
    borrados = 0
    for sub in (carpeta, carpeta / "mapas"):
        if not sub.is_dir():
            continue
        for archivo in sub.iterdir():
            try:
                if archivo.is_file() and archivo.stat().st_mtime < limite:
                    archivo.unlink()
                    borrados += 1
            except OSError as exc:
                log.warning("Visor: could not remove expired file %s (%s)", archivo.name, type(exc).__name__)
    return borrados


def archivo_del_visor(ruta: str | Path, sufijos: tuple[str, ...] = (".xlsx", ".pdf")) -> Path:
    """La ruta, solo si es un archivo existente dentro de output/visor (lo que pide la página)."""
    archivo = Path(str(ruta or "")).resolve()
    carpeta = carpeta_visor().resolve()
    if carpeta not in archivo.parents or archivo.suffix.lower() not in sufijos or not archivo.is_file():
        raise FileNotFoundError("El archivo ya no está en output/visor")
    return archivo


def _nombre_archivo(base: str, lat: float, lon: float, ahora: datetime) -> str:
    return f"{_safe_file_stem(base)}_{lat:.6f}_{lon:.6f}_{ahora:%Y%m%d-%H%M%S}.xlsx"


_NOMBRE_VISOR = re.compile(r"^(?P<base>.+)_(?P<lat>-?\d{1,3}\.\d{6})_(?P<lon>-?\d{1,3}\.\d{6})_(?P<fecha>\d{8}-\d{6})$")


def historial(carpeta: Path | None = None, limite: int = 500) -> list[dict[str, Any]]:
    """Budgets saved by the viewer in output/visor, newest first (menu "Presupuestos" > Historial).

    Client, coordinate and date come from the file name (<cliente>_<lat>_<lon>_<AAAAMMDD-HHMMSS>.xlsx);
    the totals are read only when one is opened (abrir_guardado)."""
    carpeta = carpeta or carpeta_visor()
    if not carpeta.is_dir():
        return []
    archivos = sorted((p for p in carpeta.glob("*.xlsx") if p.is_file() and not p.name.startswith("~$")),
                      key=lambda p: p.stat().st_mtime, reverse=True)[:limite]
    salida = []
    for p in archivos:
        m = _NOMBRE_VISOR.match(p.stem)
        fecha = (datetime.strptime(m["fecha"], "%Y%m%d-%H%M%S") if m
                 else datetime.fromtimestamp(p.stat().st_mtime)).isoformat(timespec="seconds")
        base = (m["base"] if m else p.stem).replace("_", " ").strip()
        salida.append({"archivo": str(p), "nombre": p.name, "cliente": base, "fecha": fecha,
                       "coordenada": f"{m['lat']}, {m['lon']}" if m else None,
                       "estimacion": base.casefold() == "estimacion", "kb": round(p.stat().st_size / 1024),
                       "tiene_pdf": p.with_suffix(".pdf").is_file()})
    return salida


def abrir_guardado(archivo: str | Path) -> dict[str, Any]:
    """A saved viewer budget, with the same shape as resultado_para_visor (for the result area).
    The values cached by LibreOffice are used when the file has them."""
    xlsx = archivo_del_visor(archivo, (".xlsx",))
    resumen = resumen_presupuesto(xlsx, True)
    calculado = any(t.get("valor") is not None for t in resumen["totales"])
    if not calculado:
        resumen = resumen_presupuesto(xlsx, False)
    m = _NOMBRE_VISOR.match(xlsx.stem)
    avisos = [] if calculado else ["Este Excel no tiene los totales calculados: se ven al abrirlo en Excel"]
    return {"archivo": str(xlsx), "nombre": xlsx.name, "carpeta": str(xlsx.parent), "ticket_id": None, "budget_id": None,
            "calculado": calculado, "distancia_km": None, "conexion": None, "ruta": None, "tecnologia": None,
            "tipo_cable": None, "metraje_final_m": None, "postes_marcados": 0,
            "origen_trafico": {"node": None, "confidence": None}, "avisos": avisos,
            "pendientes": {"campos": [], "items": [], "postes": 0}, "resumen": resumen, "guardado": True,
            "coordenada": f"{m['lat']}, {m['lon']}" if m else None}


# --------------------------------------------------------------------------- formulario

def _texto(valor: Any) -> str:
    return "" if valor is None else str(valor).strip()


def construir_solicitud(campos: dict[str, Any], ahora: datetime | None = None) -> main.TicketRequest:
    """TicketRequest desde el formulario del visor, con las mismas validaciones que la API.

    ticket_id vacío -> "VISOR-AAAAMMDD-HHMMSS"; contact_name vacío -> client_name; coordenada en
    un solo campo ("12.15, -86.26", como la barra del visor); tecnología Automático/GPON/Metro;
    ODF de rack Ninguno/12/24.
    """
    ahora = ahora or datetime.now()
    errores: dict[str, str] = {}
    lat = lon = None
    try:
        lat, lon = parsear_coordenadas(_texto(campos.get("coordenada")))
    except CoordenadaInvalida as exc:
        errores["coordenada"] = str(exc)
    tecnologia = _normalized(campos.get("technology"))
    if tecnologia not in TECNOLOGIAS:
        errores["technology"] = "Elija Automático, GPON o Metro"
    odf = _normalized(campos.get("odf_rack_ports"))
    if odf not in ("", "ninguno", "12", "24"):
        errores["odf_rack_ports"] = "Elija Ninguno, 12 o 24"
    cliente = _texto(campos.get("client_name"))
    datos = {
        "ticket_id": _texto(campos.get("ticket_id")) or f"VISOR-{ahora:%Y%m%d-%H%M%S}",
        "client_name": cliente,
        "client_code": _texto(campos.get("client_code")),
        "sales_exec": _texto(campos.get("sales_exec")),
        "contact_name": _texto(campos.get("contact_name")) or cliente,
        "contact_phone": _texto(campos.get("contact_phone")),
        "raw_comment": str(campos.get("raw_comment") or ""),
        "technology": TECNOLOGIAS.get(tecnologia),
        "odf_rack_ports": int(odf) if odf in ("12", "24") else None,
        "latitude": lat if lat is not None else 0.0,
        "longitude": lon if lon is not None else 0.0,
        "service_type": _texto(campos.get("service_type")),
        "medium": _texto(campos.get("medium")) or "FO",
        "bandwidth": _texto(campos.get("bandwidth")),
    }
    try:
        solicitud = main.TicketRequest(**datos)
    except ValidationError as exc:
        for error in exc.errors():
            campo = str(error["loc"][0]) if error.get("loc") else "coordenada"
            campo = "coordenada" if campo in ("latitude", "longitude") else campo
            mensaje = "Campo obligatorio" if error.get("type") in ("string_too_short", "missing") else error["msg"]
            errores.setdefault(campo, mensaje)
    if errores:
        raise ErrorFormulario(errores)
    return solicitud


def postes_a_mano(valor: Any) -> int:
    """How many poles the user added by hand on the map (0 when the value is not a sensible number)."""
    try:
        n = int(valor)
    except (TypeError, ValueError):
        return 0
    return max(0, min(n, MAX_POSTES_MARCADOS))


def mufas_tocadas(valor: Any) -> int | None:
    """How many mufas the user's fusion route intervenes (None: nobody edited it, the budget keeps the computed count)."""
    if valor is None or valor == "":
        return None
    try:
        return max(0, min(int(valor), 100))
    except (TypeError, ValueError):
        return None


def postes_marcados(valor: Any) -> list[float] | None:
    """Metros de ruta de los postes nuevos marcados como existentes en el mapa (None: ninguno).

    Se descartan los valores que no son números finitos >= 0; como máximo MAX_POSTES_MARCADOS.
    """
    if not isinstance(valor, (list, tuple)):
        return None
    metros: list[float] = []
    for v in valor[:MAX_POSTES_MARCADOS]:
        try:
            m = float(v)
        except (TypeError, ValueError):
            continue
        if math.isfinite(m) and m >= 0:
            metros.append(m)
    return metros or None


# --------------------------------------------------------------------------- generación

def generar(solicitud: main.TicketRequest, get_route: Any = None, nombre_base: str | None = None,
            ahora: datetime | None = None, postes_existentes_m: list[float] | None = None,
            postes_extra: int = 0, mufas_tocadas: int | None = None) -> dict[str, Any]:
    """Genera el presupuesto con main._generar_presupuesto_con en output/visor.

    Devuelve la misma respuesta que POST /presupuesto, con excel_path apuntando al archivo
    renombrado <cliente>_<lat>_<lon>_<AAAAMMDD-HHMMSS>.xlsx. No escribe budget_index.json.
    postes_existentes_m: postes nuevos marcados como existentes en el mapa (metros de la ruta);
    se cobran como existentes y no como instalación de poste nuevo.
    """
    ahora = ahora or datetime.now()
    carpeta = carpeta_visor()
    carpeta.mkdir(parents=True, exist_ok=True)
    limpiar_archivos_viejos(carpeta)
    visio_con = red_fibra_pg.connect_or_none()
    opcionales = {"new_poles_marked_existing_m": postes_existentes_m} if postes_existentes_m else {}
    if postes_extra:
        opcionales["new_poles_extra"] = int(postes_extra)
    if mufas_tocadas is not None:
        opcionales["mufa_count_override"] = int(mufas_tocadas)
    try:
        respuesta = main._generar_presupuesto_con(solicitud, visio_con, output_dir=carpeta, get_route=get_route,
                                                  **opcionales)
    except HTTPException as exc:
        raise PresupuestoError(str(exc.detail)) from exc
    finally:
        if visio_con is not None:
            visio_con.close()
    generado = Path(respuesta["excel_path"])
    destino = carpeta / _nombre_archivo(nombre_base or solicitud.client_name, solicitud.latitude,
                                         solicitud.longitude, ahora)
    if destino.exists():
        destino.unlink()
    generado.replace(destino)
    respuesta["excel_path"] = str(destino)
    log.info("Visor: budget %s saved as %s", respuesta["budget_id"], destino.name)
    return respuesta


def supuestos_estimacion(metraje_final_m: float, rules: dict[str, Any]) -> dict[str, Any]:
    """Valores por defecto del presupuesto aproximado (business_rules.yaml, visor_estimacion).

    Servicio: Residencial hasta cable.adss_threshold_m de metraje final y Empresarial por encima
    (el mismo umbral que elige Flat o ADSS). Tecnología: la regla actual con ese ancho de banda.
    """
    config = {**ESTIMACION_POR_DEFECTO, **(rules.get("visor_estimacion") or {})}
    umbral = float(rules["cable"]["adss_threshold_m"])
    hasta = float(metraje_final_m) <= umbral
    tecnologia, _ = seleccionar_tecnologia(config["bandwidth"], None, "", rules)
    return {
        "service_type": config["servicio_hasta_umbral"] if hasta else config["servicio_sobre_umbral"],
        "bandwidth": config["bandwidth"], "medium": config["medium"], "technology": tecnologia,
        "cable": "Flat" if hasta else "ADSS", "metraje_final_m": round(float(metraje_final_m), 1),
        "umbral_m": umbral, "cliente": config["cliente"],
    }


def campos_estimacion(lat: float, lon: float, supuestos: dict[str, Any], ahora: datetime | None = None) -> dict[str, Any]:
    """Formulario equivalente a la estimación (para generarla y para "Ajustar datos")."""
    ahora = ahora or datetime.now()
    cliente = supuestos["cliente"]
    return {"ticket_id": f"VISOR-{ahora:%Y%m%d-%H%M%S}", "client_name": cliente, "client_code": cliente,
            "sales_exec": cliente, "contact_name": cliente, "contact_phone": cliente,
            "coordenada": f"{lat:.6f}, {lon:.6f}", "service_type": supuestos["service_type"],
            "medium": supuestos["medium"], "bandwidth": supuestos["bandwidth"], "technology": "auto",
            "odf_rack_ports": "ninguno", "raw_comment": ""}


def estimar(lat: float, lon: float, metraje_final_m: float, get_route: Any = None,
            postes_existentes_m: list[float] | None = None, postes_extra: int = 0, mufas_tocadas: int | None = None) -> dict[str, Any]:
    """Presupuesto aproximado de la vista principal: solo la coordenada, supuestos por defecto.

    metraje_final_m: el de la ruta que el visor acaba de dibujar (calculo.calcular_ruta), que
    decide Residencial / Empresarial. Con el caché de rutas, la generación no vuelve a llamar a Google.
    postes_existentes_m: postes nuevos marcados como existentes en el mapa (ver generar).
    """
    rules = main.load_rules(main.settings.rules_path)
    supuestos = supuestos_estimacion(metraje_final_m, rules)
    campos = campos_estimacion(lat, lon, supuestos)
    respuesta = generar(construir_solicitud(campos), get_route, nombre_base="Estimacion",
                        postes_existentes_m=postes_existentes_m, postes_extra=postes_extra, mufas_tocadas=mufas_tocadas)
    supuestos["postes_marcados"] = len((respuesta.get("budget") or {}).get("new_poles_marked_existing_m") or [])
    supuestos["technology"] = respuesta["equipment"]["technology"]
    supuestos["cable"] = respuesta["budget"]["cable_type"]
    return {"supuestos": supuestos, "campos": campos, "respuesta": respuesta}


# --------------------------------------------------------------------------- lectura del Excel

def _es_formula(valor: Any) -> bool:
    return isinstance(valor, ArrayFormula) or (isinstance(valor, str) and valor.startswith("="))


def _numero(valor: Any) -> float | None:
    if isinstance(valor, bool) or valor is None:
        return None
    if isinstance(valor, (int, float)):
        return float(valor) if math.isfinite(valor) else None
    try:
        return float(str(valor).replace(",", ""))
    except ValueError:
        return None


# The template's total rows, found by their label (not by row number); shown with the template's text.
_TOTALES = ("alcaldia", "subtotal", "costo total", "total")


def resumen_presupuesto(xlsx: Path, calculado: bool) -> dict[str, Any]:
    """Desglose del Excel generado: líneas con cantidad > 0 de MANO DE OBRA y MATERIALES
    (descripción E, unidad F, cantidad G, precio H, total I), totales de cada sección y los
    totales de la plantilla (Alcaldía, Subtotal, Costo total, TOTAL). Las secciones y totales
    se ubican por sus etiquetas, no por número de fila.

    Los montos son los valores que dejó recalc.py; sin LibreOffice (calculado False) las celdas
    con fórmula quedan en None, porque el valor guardado sería el de la plantilla vacía, y las
    cantidades que calcula la plantilla se cuentan en "por_formula".
    """
    valores = load_workbook(xlsx, data_only=True)["Presupuesto"]
    formulas = load_workbook(xlsx, data_only=False)["Presupuesto"]

    def leer(fila: int, columna: int) -> Any:
        if _es_formula(formulas.cell(fila, columna).value) and not calculado:
            return None
        return valores.cell(fila, columna).value

    secciones: list[dict[str, Any]] = []
    totales: list[dict[str, Any]] = []
    actual = None
    for fila in range(1, formulas.max_row + 1):
        titulo = _normalized(formulas.cell(fila, 3).value)
        if titulo in ("mano de obra", "materiales"):
            actual = {"nombre": "Mano de obra" if titulo == "mano de obra" else "Materiales",
                      "total": _numero(leer(fila, 9)), "lineas": [], "por_formula": 0}
            secciones.append(actual)
            continue
        textos = [(c, _normalized(formulas.cell(fila, c).value)) for c in range(3, 9)]
        encontrado = next(((clave, c) for clave in _TOTALES for c, e in textos
                           if e == clave or (clave != "total" and e.startswith(clave))), None)
        if encontrado is not None:
            clave, columna = encontrado
            actual = None
            etiqueta = " ".join(str(formulas.cell(fila, columna).value).split())
            totales.append({"clave": clave, "etiqueta": etiqueta, "valor": _numero(leer(fila, 9))})
            if clave == "total":
                break
            continue
        if actual is None or not formulas.cell(fila, 5).value:
            continue
        cantidad = _numero(leer(fila, 7))
        if cantidad is None and _es_formula(formulas.cell(fila, 7).value):
            actual["por_formula"] += 1         # quantity computed by the template: needs LibreOffice / Excel
            continue
        if not cantidad:
            continue
        actual["lineas"].append({"descripcion": str(formulas.cell(fila, 5).value).strip(),
                                 "unidad": leer(fila, 6), "cantidad": cantidad,
                                 "precio": _numero(leer(fila, 8)), "total": _numero(leer(fila, 9))})
    return {"calculado": calculado, "secciones": secciones, "totales": totales}


def resultado_para_visor(respuesta: dict[str, Any]) -> dict[str, Any]:
    """Lo que la página necesita de una respuesta de presupuesto (sin el HTML de las hojas)."""
    archivo = Path(respuesta["excel_path"])
    calculado = bool(respuesta.get("values_cached"))
    avisos = list(respuesta.get("warnings") or [])
    if not calculado:
        avisos.insert(0, "LibreOffice no disponible: los montos y totales se ven al abrir el Excel")
    return {
        "archivo": str(archivo), "nombre": archivo.name, "carpeta": str(archivo.parent),
        "ticket_id": respuesta["ticket_id"], "budget_id": respuesta["budget_id"],
        "calculado": calculado, "distancia_km": respuesta.get("distance_km"),
        "conexion": respuesta.get("connection_point"), "ruta": respuesta.get("route"),
        "tecnologia": (respuesta.get("equipment") or {}).get("technology"),
        "tipo_cable": (respuesta.get("budget") or {}).get("cable_type"),
        "metraje_final_m": (respuesta.get("budget") or {}).get("final_distance_m"),
        "postes_marcados": len((respuesta.get("budget") or {}).get("new_poles_marked_existing_m") or []),
        "origen_trafico": {k: (respuesta.get("traffic_origin") or {}).get(k) for k in ("node", "confidence")},
        "avisos": avisos,
        "pendientes": {
            "campos": respuesta.get("pending_validation_fields") or [],
            "items": respuesta.get("pending_validation_items") or [],
            "postes": len(respuesta.get("pending_validation_poles") or []),
        },
        "resumen": resumen_presupuesto(archivo, calculado),
    }
