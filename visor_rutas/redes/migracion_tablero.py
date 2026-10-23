"""Migración Visio, fase 6: cola de corrección y tablero (docs/migracion_visio/PLAN.md §7).

- `tablero`: por proyecto base (zona), cuántas mufas migradas hay, cuántas con fusiones conocidas, con
  supuestos (estimados), validadas y editadas, y el avance (validadas / migradas).
- `correcciones`: la lista ordenada de lo que hay que corregir, una fila por mufa y motivo, sin las
  mufas ya validadas; filtros por zona, motivo y texto. Cada fila trae el id de la mufa para saltar a
  ella (o a su proyecto, si está sin ubicar).
- `exportar_correcciones`: lo mismo en Excel (hojas Tablero y Correcciones) para Planta Externa.
Solo lee.
"""
import base64
import io
from datetime import datetime
from typing import Any

from visor_rutas.redes.db import SinPermiso

# (prioridad, clave, título): the order of the queue
MOTIVOS = [
    (1, "sin_ubicar", "Sin ubicar"),
    (2, "en_revision", "El Visio pedía revisarla"),
    (3, "excede_capacidad", "Excede la capacidad del modelo"),
    (4, "hilos_desconocidos", "Cables sin cantidad de hilos"),
    (5, "sin_diagramar", "Sin diagramar (sin fusiones)"),
    (6, "ubicacion_estimada", "Ubicación estimada"),
    (7, "configuracion_estimada", "Configuración de cable estimada"),
    (8, "modelo_estimado", "Modelo de mufa estimado"),
]
TITULO = {k: t for _, k, t in MOTIVOS}
PRIORIDAD = {k: p for p, k, _ in MOTIVOS}

_CONSULTA = """
WITH m AS (
  SELECT e.id, e.nombre, e.tipo, e.subtipo, e.proyecto_id, p.nombre AS proyecto, e.estado_migracion, e.ubicacion_fuente,
         e.ref_visio ->> 'motivo_revision' AS motivo_revision
  FROM redes.elementos e JOIN redes.proyectos p ON p.id = e.proyecto_id
  WHERE p.es_red_base AND p.borrado_en IS NULL AND e.borrado_en IS NULL AND e.tipo IN ('mufa', 'nodo', 'odf') AND e.ref_visio ? 'clave'
    AND (%(proyecto)s::bigint IS NULL OR e.proyecto_id = %(proyecto)s)),
extremos AS (
  SELECT origen_id AS e_id, hilos, configuracion_estimada FROM redes.cables WHERE borrado_en IS NULL
  UNION ALL SELECT destino_id, hilos, configuracion_estimada FROM redes.cables WHERE borrado_en IS NULL),
cab AS MATERIALIZED (SELECT e_id, count(*) AS n, count(*) FILTER (WHERE hilos IS NULL) AS desconocidos,
               count(*) FILTER (WHERE hilos IS NOT NULL AND configuracion_estimada) AS config_estimada
        FROM extremos GROUP BY 1),
con AS MATERIALIZED (SELECT mufa_id, count(*) AS n FROM redes.conexiones WHERE borrado_en IS NULL GROUP BY 1),
modelo AS MATERIALIZED (SELECT mufa_id, modelo_estimado, avisos FROM redes.mufa_modelos WHERE borrado_en IS NULL)
SELECT m.*, coalesce(cab.n, 0) AS cables, coalesce(cab.desconocidos, 0) AS desconocidos, coalesce(cab.config_estimada, 0) AS config_estimada,
       coalesce(con.n, 0) AS conexiones, modelo.modelo_estimado, coalesce(modelo.avisos, '[]'::jsonb) AS avisos_modelo
FROM m LEFT JOIN cab ON cab.e_id = m.id LEFT JOIN con ON con.mufa_id = m.id LEFT JOIN modelo ON modelo.mufa_id = m.id
ORDER BY m.proyecto, m.nombre"""


def _mufas(con: Any, proyecto_id: int | None = None) -> list[dict[str, Any]]:
    return [dict(f) for f in con.execute(_CONSULTA, {"proyecto": proyecto_id}).fetchall()]


def motivos_de(m: dict[str, Any]) -> list[tuple[str, str]]:
    """(motivo, detalle) of one migrated mufa; none if a person validated it."""
    if m["estado_migracion"] == "validada":
        return []
    salida = []
    if m["ubicacion_fuente"] == "sin_ubicar":
        salida.append(("sin_ubicar", "no tiene coordenada: colóquela en el mapa de su proyecto"))
    if m["estado_migracion"] == "en_revision":
        salida.append(("en_revision", m["motivo_revision"] or "marcada para revisar en el Visio"))
    if m["avisos_modelo"]:
        salida.append(("excede_capacidad", "; ".join(m["avisos_modelo"])))
    if m["desconocidos"]:
        salida.append(("hilos_desconocidos", f"{m['desconocidos']} de {m['cables']} cables sin cantidad de hilos"))
    if m["cables"] and not m["conexiones"] and not (m["tipo"] == "mufa" and m["subtipo"] == "reserva"):
        salida.append(("sin_diagramar", "ninguna fusión conocida"))
    if m["ubicacion_fuente"] == "estimada":
        salida.append(("ubicacion_estimada", "coordenada estimada por sus vecinas"))
    if m["config_estimada"]:
        salida.append(("configuracion_estimada", f"{m['config_estimada']} cable(s) con la configuración de buffers por regla"))
    if m["modelo_estimado"]:
        salida.append(("modelo_estimado", "48 / 96 y bandejas por regla"))
    return salida


def _visible(actor: dict[str, Any]) -> None:
    """Base projects are visible to every active user of the module (proyectos.puede_ver)."""
    if not actor or actor.get("borrado_en") is not None:
        raise SinPermiso("Usuario desactivado")


def tablero(con: Any, actor: dict[str, Any]) -> dict[str, Any]:
    """Progress of the migrated network, per base project and in total."""
    _visible(actor)
    zonas: dict[int, dict[str, Any]] = {}
    motivos = {k: 0 for _, k, _ in MOTIVOS}
    for m in _mufas(con):
        z = zonas.setdefault(m["proyecto_id"], {"proyecto_id": m["proyecto_id"], "proyecto": m["proyecto"], "mufas": 0, "con_fusiones": 0,
                                                "con_supuestos": 0, "validadas": 0, "editadas": 0, "sin_ubicar": 0, "sin_diagramar": 0,
                                                "por_corregir": 0})
        z["mufas"] += 1
        z["con_fusiones"] += m["conexiones"] > 0
        z["validadas"] += m["estado_migracion"] == "validada"
        z["editadas"] += m["estado_migracion"] == "editada"
        lista = [k for k, _ in motivos_de(m)]
        z["con_supuestos"] += any(k in ("hilos_desconocidos", "ubicacion_estimada", "configuracion_estimada", "modelo_estimado") for k in lista)
        z["sin_ubicar"] += "sin_ubicar" in lista
        z["sin_diagramar"] += "sin_diagramar" in lista
        z["por_corregir"] += bool(lista)
        for k in lista:
            motivos[k] += 1
    filas = sorted(zonas.values(), key=lambda z: -z["mufas"])
    for z in filas:
        z["avance"] = round(100 * z["validadas"] / z["mufas"], 1) if z["mufas"] else 0
    total = {k: sum(z[k] for z in filas) for k in ("mufas", "con_fusiones", "con_supuestos", "validadas", "editadas", "sin_ubicar",
                                                   "sin_diagramar", "por_corregir")}
    total["avance"] = round(100 * total["validadas"] / total["mufas"], 1) if total["mufas"] else 0
    return {"zonas": filas, "total": total, "motivos": [{"clave": k, "titulo": t, "mufas": motivos[k]} for _, k, t in MOTIVOS]}


def sugerencias_de_ubicacion(con: Any, ids: list[int]) -> dict[int, dict[str, Any]]:
    """Where an element without coordinate probably is, only as a hint to centre the map (never saved):
      - `vecinas`: the median position of the located mufas its cables join it to in the Visio (`red_fibra.enlaces`; estimated coordinates and
        clients do not count);
      - `diagrama`: failing that, the median of the located mufas that are drawn in the same Visio diagram (the smallest one it appears in);
      - `carpeta`: failing that, the median of the located elements of the same Visio folder (its diagrams are of one stretch of network).
    {id: {lat, lon, fuente, zoom, vecinas, nombres}}; an element with none of them has no suggestion."""
    if not ids:
        return {}
    salida: dict[int, dict[str, Any]] = {}
    try:
        for f in con.execute("""
            WITH sin AS (SELECT id, (ref_visio ->> 'nodo_id')::int AS nodo FROM redes.elementos WHERE id = ANY(%s) AND ref_visio ? 'nodo_id'),
            vec AS (SELECT s.id, v.nombre, v.lat, v.lon FROM sin s
                    JOIN red_fibra.enlaces e ON e.nodo_a_id = s.nodo OR e.nodo_b_id = s.nodo
                    JOIN red_fibra.nodos v ON v.id = CASE WHEN e.nodo_a_id = s.nodo THEN e.nodo_b_id ELSE e.nodo_a_id END
                    WHERE v.id <> s.nodo AND v.lat IS NOT NULL AND v.lon IS NOT NULL AND coalesce(v.fuente_coord, '') <> 'estimada_vecinos'
                      AND coalesce(v.tipo, '') <> 'CLIENTE')
            SELECT id, percentile_cont(0.5) WITHIN GROUP (ORDER BY lat) AS lat, percentile_cont(0.5) WITHIN GROUP (ORDER BY lon) AS lon,
                   count(DISTINCT nombre) AS vecinas, (array_agg(DISTINCT nombre ORDER BY nombre))[1:3] AS nombres FROM vec GROUP BY id""", (list(ids),)).fetchall():
            salida[f["id"]] = {"lat": round(float(f["lat"]), 6), "lon": round(float(f["lon"]), 6), "fuente": "vecinas", "zoom": 17, "vecinas": int(f["vecinas"]),
                               "nombres": list(f["nombres"] or [])}
        faltan = [i for i in ids if i not in salida]
        if faltan:
            for f in con.execute("""
                WITH sin AS (SELECT id, (ref_visio ->> 'nodo_id')::int AS nodo FROM redes.elementos WHERE id = ANY(%s) AND ref_visio ? 'nodo_id'),
                cerca AS (SELECT s.id, a.archivo_id, v.nombre, v.lat, v.lon FROM sin s
                          JOIN red_fibra.nodo_apariciones a ON a.nodo_id = s.nodo
                          JOIN red_fibra.nodo_apariciones b ON b.archivo_id = a.archivo_id AND b.nodo_id <> s.nodo
                          JOIN red_fibra.nodos v ON v.id = b.nodo_id
                          WHERE v.lat IS NOT NULL AND v.lon IS NOT NULL AND coalesce(v.fuente_coord, '') <> 'estimada_vecinos' AND coalesce(v.tipo, '') <> 'CLIENTE'),
                tam AS (SELECT id, archivo_id, count(DISTINCT nombre) AS c FROM cerca GROUP BY 1, 2),
                mejor AS (SELECT DISTINCT ON (id) id, archivo_id FROM tam ORDER BY id, c, archivo_id)
                SELECT c.id, percentile_cont(0.5) WITHIN GROUP (ORDER BY c.lat) AS lat, percentile_cont(0.5) WITHIN GROUP (ORDER BY c.lon) AS lon,
                       count(DISTINCT c.nombre) AS vecinas, (array_agg(DISTINCT c.nombre ORDER BY c.nombre))[1:3] AS nombres
                FROM cerca c JOIN mejor m ON m.id = c.id AND m.archivo_id = c.archivo_id GROUP BY c.id""", (faltan,)).fetchall():
                salida[f["id"]] = {"lat": round(float(f["lat"]), 6), "lon": round(float(f["lon"]), 6), "fuente": "diagrama", "zoom": 16, "vecinas": int(f["vecinas"]),
                                   "nombres": list(f["nombres"] or [])}
    except Exception:                                                  # noqa: BLE001 - no red_fibra schema (a test, or Visio not imported): no hint
        return {}
    faltan = [i for i in ids if i not in salida]
    if faltan:
        for f in con.execute("""
            SELECT s.id, percentile_cont(0.5) WITHIN GROUP (ORDER BY ST_Y(o.ubicacion)) AS lat, percentile_cont(0.5) WITHIN GROUP (ORDER BY ST_X(o.ubicacion)) AS lon,
                   count(*) AS n
            FROM redes.elementos s JOIN redes.elementos o ON o.proyecto_id = s.proyecto_id AND o.id <> s.id AND o.borrado_en IS NULL AND o.ubicacion IS NOT NULL
                 AND o.ref_visio ->> 'zona' = s.ref_visio ->> 'zona'
            WHERE s.id = ANY(%s) AND s.ref_visio ? 'zona' GROUP BY s.id""", (faltan,)).fetchall():
            salida[f["id"]] = {"lat": round(float(f["lat"]), 6), "lon": round(float(f["lon"]), 6), "fuente": "carpeta", "zoom": 14, "vecinas": int(f["n"]), "nombres": []}
    return salida


def _filas(con: Any, proyecto_id: int | None, motivo: str | None, texto: str | None) -> list[dict[str, Any]]:
    t = (texto or "").strip().upper()
    filas = []
    for m in _mufas(con, int(proyecto_id) if proyecto_id else None):
        if t and t not in m["nombre"].upper():
            continue
        for k, detalle in motivos_de(m):
            if motivo and k != motivo:
                continue
            filas.append({"mufa_id": m["id"], "mufa": m["nombre"], "tipo": m["tipo"], "proyecto_id": m["proyecto_id"], "proyecto": m["proyecto"],
                          "motivo": k, "titulo": TITULO[k], "detalle": detalle, "prioridad": PRIORIDAD[k], "estado": m["estado_migracion"],
                          "sugerencia": None})
    pistas = sugerencias_de_ubicacion(con, [f["mufa_id"] for f in filas if f["motivo"] == "sin_ubicar"])
    for f in filas:
        if f["motivo"] == "sin_ubicar":
            sug = pistas.get(f["mufa_id"])
            f["sugerencia"] = sug
            if sug and sug["fuente"] == "vecinas":
                resto = sug["vecinas"] - len(sug["nombres"])
                f["detalle"] = f"no tiene coordenada: se une a {', '.join(sug['nombres'])}" + (f" y {resto} más" if resto > 0 else "") + " (el mapa se centra ahí; colóquela con un clic)"
            elif sug and sug["fuente"] == "diagrama":
                f["detalle"] = (f"no tiene coordenada: en su diagrama aparecen {', '.join(sug['nombres'])}" + (f" y {sug['vecinas'] - len(sug['nombres'])} más" if sug["vecinas"] > len(sug["nombres"]) else "")
                                + " (el mapa se centra entre ellas; colóquela con un clic)")
            elif sug:
                f["detalle"] = f"no tiene coordenada: el mapa se centra en el resto de su carpeta del Visio ({sug['vecinas']} elementos ubicados); colóquela con un clic"
            else:
                f["detalle"] = "no tiene coordenada ni vecinas ubicadas: busque su diagrama y colóquela en el mapa"
    filas.sort(key=lambda f: (f["prioridad"], f["proyecto"], f["mufa"]))
    return filas


def correcciones(con: Any, actor: dict[str, Any], *, proyecto_id: int | None = None, motivo: str | None = None, texto: str | None = None,
                 limite: int = 300) -> dict[str, Any]:
    _visible(actor)
    filas = _filas(con, proyecto_id, motivo, texto)
    limite = max(1, min(int(limite or 300), 2000))
    return {"total": len(filas), "correcciones": filas[:limite], "limite": limite}


def exportar_correcciones(con: Any, actor: dict[str, Any], *, proyecto_id: int | None = None, motivo: str | None = None,
                          texto: str | None = None) -> dict[str, Any]:
    """Excel for Planta Externa: the dashboard and the whole queue (with the filters)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    _visible(actor)
    t = tablero(con, actor)
    filas = _filas(con, proyecto_id, motivo, texto)
    wb = Workbook()
    negrita, cabeza = Font(bold=True), PatternFill("solid", fgColor="DCE6F1")

    def hoja(ws, columnas, datos, anchos):
        ws.append(columnas)
        for c in ws[ws.max_row]:
            c.font, c.fill = negrita, cabeza
        for d in datos:
            ws.append(d)
        for i, w in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(row=ws.max_row - len(datos) + 1, column=1).coordinate if datos else None

    ws = wb.active
    ws.title = "Tablero"
    claves = ["mufas", "con_fusiones", "con_supuestos", "validadas", "editadas", "sin_ubicar", "sin_diagramar", "por_corregir", "avance"]
    hoja(ws, ["Zona (proyecto base)", "Mufas migradas", "Con fusiones", "Con supuestos", "Validadas", "Editadas", "Sin ubicar",
              "Sin diagramar", "Por corregir", "Avance %"],
         [[z["proyecto"], *[z[k] for k in claves]] for z in t["zonas"]] + [["TOTAL", *[t["total"][k] for k in claves]]],
         [30, 12, 12, 12, 10, 10, 10, 12, 12, 10])
    for c in ws[ws.max_row]:
        c.font = negrita
    ws2 = wb.create_sheet("Correcciones")
    hoja(ws2, ["Prioridad", "Motivo", "Mufa", "Zona", "Detalle", "Estado", "Ubicación sugerida (lat, lon)"],
         [[f["prioridad"], f["titulo"], f["mufa"], f["proyecto"], f["detalle"], f["estado"],
           f"{f['sugerencia']['lat']:.6f}, {f['sugerencia']['lon']:.6f}" if f.get("sugerencia") else ""] for f in filas], [9, 30, 40, 28, 70, 12, 30])
    if filas:
        ws2.auto_filter.ref = ws2.dimensions
    salida = io.BytesIO()
    wb.save(salida)
    return {"nombre": f"correcciones_red_existente_{datetime.now():%Y-%m-%d}.xlsx", "tipo": "xlsx",
            "datos": base64.b64encode(salida.getvalue()).decode("ascii"), "filas": len(filas)}
