"""Coordenadas masivas (visor v2, fase 6): leer muchas coordenadas de un archivo y calcular el presupuesto de cada una en segundo plano.

Lectura (`extraer`): Excel (.xlsx), CSV, Word (.docx), texto (.txt) o texto pegado.
  - En Excel y CSV se buscan las columnas por su encabezado (latitud / longitud, una sola «coordenada», nombre o cliente, código); sin
    encabezado se revisan las celdas de cada fila.
  - En Word y texto se busca en cada línea, con la misma gramática del buscador (visor_rutas/coordenadas.py: decimal, grados-minutos-segundos
    y grados-minutos decimales). El resto de la línea es el nombre.
  Cada punto sale `ok`, `invertida` (latitud y longitud al revés: trae la sugerencia), `fuera` (de Nicaragua), `invalida` o `duplicada`
  (igual a una fila anterior); solo los `ok` se incluyen por defecto y la página deja corregir los demás.

Cálculo (`GestorLotes` / `Lote`): un hilo por lote, un punto tras otro con el mismo cálculo del presupuesto aproximado del visor (`calculo.calcular_ruta`
+ `presupuesto_visor.estimar`, supuestos de `visor_estimacion`). Las rutas de Google se piden con un límite por minuto (`visor_masivo` en
business_rules.yaml) y el lote se puede cancelar. Cada presupuesto queda como Excel en output/visor/masivo_<fecha>_<nombre>/ junto a Resumen.xlsx.
"""
import csv
import io
import logging
import re
import shutil
import threading
import time
import unicodedata
import zipfile
from collections import deque
from datetime import datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from visor_rutas import almacen_archivos, coordenadas

log = logging.getLogger(__name__)

MAX_ARCHIVO = 15 * 1024 * 1024
MAX_PUNTOS_LEIDOS = 2000
MAX_LOTES_GUARDADOS = 20
OK, INVERTIDA, FUERA, INVALIDA, DUPLICADA = "ok", "invertida", "fuera", "invalida", "duplicada"
_NUM = re.compile(r"^[-+]?\d+(?:[.,]\d+)?$")
_CODIGO = re.compile(r"^C\s?\d{3,8}$", re.I)
_DECIMAL_PAR = re.compile(r"[-+]?\d{1,3}\.\d{3,}\s*[,;]?\s*[-+]?\d{1,3}\.\d{3,}")
_DECIMAL_SUELTO = re.compile(r"[-+]?\d{1,3}\.\d+[\s,;]+[-+]?\d{1,3}\.\d+")          # 12.1 -86.2: accepted only when it falls in Nicaragua
_GMS = re.compile(r"[NSEWO]?\s*\d{1,3}\s*[°º˚]\s*\d{1,2}(?:\s*['′’´]\s*\d{1,2}(?:\.\d+)?\s*[\"″”]?)?\s*[NSEWO]?|"
                  r"[NSEWO]\s*\d{1,3}\s+\d{1,2}(?:\.\d+)?(?:\s+\d{1,2}(?:\.\d+)?)?|"
                  r"\d{1,3}\s+\d{1,2}(?:\.\d+)?(?:\s+\d{1,2}(?:\.\d+)?)?\s*[NSEWO]\b")
ENCABEZADOS_LAT = {"lat", "latitud", "latitude", "y"}
ENCABEZADOS_LON = {"lon", "lng", "long", "longitud", "longitude", "x"}
ENCABEZADOS_COORD = {"coordenada", "coordenadas", "coord", "ubicacion", "gps", "posicion", "punto", "localizacion"}


def _normal(texto: Any) -> str:
    t = "".join(c for c in unicodedata.normalize("NFD", str(texto if texto is not None else "")) if unicodedata.category(c) != "Mn")
    return " ".join(t.lower().split())


# --------------------------------------------------------------------------- lectura

def _punto(n: int, texto: str, nombre: str = "", codigo: str = "") -> dict[str, Any]:
    r = coordenadas.interpretar(texto)
    p: dict[str, Any] = {"n": n, "nombre": nombre.strip()[:160], "codigo": codigo.strip()[:40], "texto": texto.strip()[:200], "lat": None, "lon": None,
                         "estado": INVALIDA, "motivo": "", "sugerencia": None, "incluir": False}
    if r.estado == coordenadas.OK:
        p.update(lat=r.lat, lon=r.lon, estado=OK, incluir=True)
    elif r.estado == coordenadas.INVERTIDA:
        p.update(lat=r.lat, lon=r.lon, estado=INVERTIDA, motivo=r.mensaje, sugerencia=list(r.sugerencia))
    elif r.estado == coordenadas.FUERA:
        p.update(lat=r.lat, lon=r.lon, estado=FUERA, motivo=r.mensaje)
    else:
        p["motivo"] = r.mensaje or "No se reconoce la coordenada"
    return p


def _marcar_duplicadas(puntos: list[dict[str, Any]]) -> None:
    vistos: dict[tuple[float, float], int] = {}
    for p in puntos:
        if p["estado"] != OK:
            continue
        clave = (round(p["lat"], 5), round(p["lon"], 5))
        if clave in vistos:
            p.update(estado=DUPLICADA, incluir=False, motivo=f"Igual a la fila {vistos[clave]}")
        else:
            vistos[clave] = p["n"]


def _decodificar(crudo: bytes) -> str:
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return crudo.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return crudo.decode("utf-8", errors="replace")


def _texto_docx(crudo: bytes) -> list[str]:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    w = "{%s}" % ns["w"]
    try:
        with zipfile.ZipFile(io.BytesIO(crudo)) as z:
            raiz = ET.fromstring(z.read("word/document.xml"))
    except (zipfile.BadZipFile, KeyError, ET.ParseError):
        raise ValueError("El archivo Word no se pudo leer (¿es un .docx?)") from None

    def parrafo(p: ET.Element) -> str:
        partes = []
        for e in p.iter():
            if e.tag == w + "t":
                partes.append(e.text or "")
            elif e.tag in (w + "tab", w + "br"):
                partes.append("\t" if e.tag == w + "tab" else " ")
        return "".join(partes)

    lineas: list[str] = []
    cuerpo = raiz.find("w:body", ns)
    if cuerpo is not None:
        for hijo in cuerpo:
            if hijo.tag == w + "p":
                lineas.append(parrafo(hijo))
            elif hijo.tag == w + "tbl":
                for fila in hijo.iter(w + "tr"):
                    lineas.append("\t".join(" ".join(parrafo(p) for p in c.iter(w + "p")) for c in fila.findall("w:tc", ns)))
    return lineas


def _candidatos(linea: str) -> list[str]:
    """The coordinates written in a line of free text: pairs of decimals, or two groups of degrees."""
    salida = [m.group(0) for m in _DECIMAL_PAR.finditer(linea)]
    if salida:
        return salida
    grupos = [m for m in _GMS.finditer(linea) if re.search(r"\d", m.group(0))]
    return [linea[a.start():b.end()] for a, b in zip(grupos[::2], grupos[1::2])] if len(grupos) >= 2 else []


def _desde_texto(lineas: list[str]) -> list[dict[str, Any]]:
    puntos: list[dict[str, Any]] = []
    for linea in lineas:
        linea = linea.strip()
        if not linea:
            continue
        candidatas = _candidatos(linea)
        if not candidatas:
            candidatas = [m.group(0) for m in _DECIMAL_SUELTO.finditer(linea)
                          if coordenadas.interpretar(m.group(0)).estado in (coordenadas.OK, coordenadas.INVERTIDA)]
        for candidata in candidatas:
            resto = linea.replace(candidata, " ", 1)
            nombre = re.sub(r"\b(lat(itud)?|lon(g(itud)?)?|coordenadas?)\b\s*[:=]?", " ", resto, flags=re.I)
            nombre = re.sub(r"^[\s\-:,;|()\[\]\t]+|[\s\-:,;|()\[\]\t]+$", "", nombre)
            nombre = re.sub(r"\s{2,}|\t", " ", nombre)
            puntos.append(_punto(len(puntos) + 1, candidata, "" if _NUM.match(nombre) else nombre))
    return puntos


def _filas_xlsx(crudo: bytes) -> list[list[Any]]:
    try:
        libro = load_workbook(io.BytesIO(crudo), read_only=True, data_only=True)
    except Exception:
        raise ValueError("El archivo Excel no se pudo leer (¿es un .xlsx?)") from None
    hoja = libro.active
    filas = [list(f) for f in hoja.iter_rows(values_only=True, max_row=MAX_PUNTOS_LEIDOS + 20)]
    libro.close()
    return filas


def _filas_csv(crudo: bytes) -> list[list[Any]]:
    texto = _decodificar(crudo)
    try:
        dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t|")
    except csv.Error:
        dialecto = csv.excel
    return [list(f) for f in csv.reader(io.StringIO(texto), dialecto)]


def _celda(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return repr(v) if v != int(v) else str(int(v))
    return str(v).strip()


def _decimal(c: str) -> str:
    """«12,1364» (decimal comma of a Spanish-language Excel / CSV) -> «12.1364»; anything else stays."""
    return c.replace(",", ".") if re.fullmatch(r"[-+]?\d+,\d+", c) else c


def _columnas(fila: list[Any]) -> dict[str, int]:
    cols: dict[str, int] = {}
    for i, c in enumerate(fila):
        t = _normal(c)
        if not t:
            continue
        if t in ENCABEZADOS_LAT or t.startswith("latitud"):
            cols.setdefault("lat", i)
        elif t in ENCABEZADOS_LON or t.startswith("longitud"):
            cols.setdefault("lon", i)
        elif t in ENCABEZADOS_COORD or t.startswith("coordenada"):
            cols.setdefault("coord", i)
        elif "codigo" in t or t in ("cod", "code", "id", "id cliente", "c#"):
            cols.setdefault("codigo", i)
        elif any(k in t for k in ("nombre", "cliente", "sitio", "ticket", "name", "descripcion", "direccion")):
            cols.setdefault("nombre", i)
    return cols


def _desde_tabla(filas: list[list[Any]]) -> list[dict[str, Any]]:
    filas = [f for f in filas if any(_celda(c) for c in f)]
    inicio, cols = 0, {}
    for i, f in enumerate(filas[:15]):
        c = _columnas(f)
        if ("lat" in c and "lon" in c) or "coord" in c:
            inicio, cols = i + 1, c
            break
    puntos: list[dict[str, Any]] = []
    for fila in filas[inicio:]:
        celdas = [_celda(c) for c in fila]
        if cols:
            def en(k: str) -> str:
                return celdas[cols[k]] if k in cols and cols[k] < len(celdas) else ""
            texto = f"{_decimal(en('lat'))}, {_decimal(en('lon'))}" if "lat" in cols and "lon" in cols else en("coord")
            if not texto.strip(", "):
                if any(celdas):
                    puntos.append({**_punto(len(puntos) + 1, "", en("nombre"), en("codigo")), "motivo": "La fila no tiene coordenada"})
                continue
            puntos.append(_punto(len(puntos) + 1, texto, en("nombre"), en("codigo")))
            continue
        # no header: one cell with a coordinate, or two adjacent numeric cells
        encontrada, usadas = None, set()
        for j, c in enumerate(celdas):
            if not c:
                continue
            if _NUM.match(c) and j + 1 < len(celdas) and _NUM.match(celdas[j + 1]):
                r = coordenadas.interpretar(f"{c}, {celdas[j + 1]}")
                if r.estado != coordenadas.NO_ES and r.estado != coordenadas.ERROR:
                    encontrada, usadas = f"{c}, {celdas[j + 1]}", {j, j + 1}
                    break
            if not _NUM.match(c) and coordenadas.interpretar(c).estado in (coordenadas.OK, coordenadas.INVERTIDA, coordenadas.FUERA):
                encontrada, usadas = c, {j}
                break
        if encontrada is None:
            continue
        resto = [c for j, c in enumerate(celdas) if j not in usadas and c]
        codigo = next((c for c in resto if _CODIGO.match(c)), "")
        nombre = next((c for c in resto if c != codigo and not _NUM.match(c)), "")
        puntos.append(_punto(len(puntos) + 1, encontrada, nombre, codigo))
    return puntos


def extraer(nombre: str, crudo: bytes) -> dict[str, Any]:
    """The coordinates of a file: {"formato", "puntos": [...], "avisos": [...]}. ValueError with a clear message when it cannot be read."""
    if not crudo:
        raise ValueError("El archivo está vacío")
    if len(crudo) > MAX_ARCHIVO:
        raise ValueError(f"El archivo debe pesar hasta {MAX_ARCHIVO // (1024 * 1024)} MB")
    sufijo = Path(str(nombre or "")).suffix.lower()
    avisos: list[str] = []
    if sufijo in (".xlsx", ".xlsm"):
        formato, puntos = "Excel", _desde_tabla(_filas_xlsx(crudo))
    elif sufijo == ".xls":
        raise ValueError("Los .xls viejos no se pueden leer: ábralo en Excel y guárdelo como .xlsx (o como .csv)")
    elif sufijo == ".csv":
        formato, puntos = "CSV", _desde_tabla(_filas_csv(crudo))
    elif sufijo == ".docx":
        formato, puntos = "Word", _desde_texto(_texto_docx(crudo))
    elif sufijo in (".txt", ".text", ".md", ".log", ""):
        formato, puntos = "Texto", _desde_texto(_decodificar(crudo).splitlines())
    else:
        raise ValueError("Formato no admitido: use Excel (.xlsx), CSV, Word (.docx) o texto (.txt)")
    if not puntos:
        raise ValueError("No se encontró ninguna coordenada en el archivo")
    if len(puntos) > MAX_PUNTOS_LEIDOS:
        puntos = puntos[:MAX_PUNTOS_LEIDOS]
        avisos.append(f"Se leyeron solo las primeras {MAX_PUNTOS_LEIDOS} coordenadas")
    _marcar_duplicadas(puntos)
    return {"formato": formato, "puntos": puntos, "avisos": avisos}


def revisar(puntos: Any) -> list[dict[str, Any]]:
    """The points of the page (possibly edited): validates the text of each one again, keeps the choice to include it."""
    if not isinstance(puntos, list) or not puntos:
        raise ValueError("No hay coordenadas")
    salida = []
    for i, p in enumerate(puntos[:MAX_PUNTOS_LEIDOS], 1):
        p = p if isinstance(p, dict) else {}
        nuevo = _punto(i, str(p.get("texto") or ""), str(p.get("nombre") or ""), str(p.get("codigo") or ""))
        if not p.get("incluir", nuevo["incluir"]):
            nuevo["incluir"] = False
        salida.append(nuevo)
    _marcar_duplicadas(salida)
    return salida


# --------------------------------------------------------------------------- límite de Google y lote

class CacheLimitada:
    """CacheRutas with a limit of Google routes per minute: waits (and says so) before asking for one that is not in the cache."""

    def __init__(self, cache: Any, por_minuto: int, lote: "Lote"):
        self._cache = cache
        self._por_minuto = max(1, int(por_minuto))
        self._lote = lote
        self._llamadas: deque[float] = deque()

    def __call__(self, origin: tuple[float, float], destination: tuple[float, float]) -> Any:
        if self._cache._clave(origin, destination) not in self._cache._rutas:
            while True:
                ahora = time.monotonic()
                while self._llamadas and ahora - self._llamadas[0] >= 60:
                    self._llamadas.popleft()
                if len(self._llamadas) < self._por_minuto:
                    break
                espera = 60 - (ahora - self._llamadas[0])
                self._lote.esperando_s = round(espera)
                if self._lote.cancelar.wait(timeout=min(1.0, max(0.05, espera))):
                    self._lote.esperando_s = 0
                    raise InterruptedError("lote cancelado")
            self._lote.esperando_s = 0
            self._llamadas.append(time.monotonic())
        return self._cache(origin, destination)


def _seguro(texto: str) -> str:
    t = "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")
    return re.sub(r"[^A-Za-z0-9_-]+", "_", t).strip("_")[:50] or "lote"


class Lote:
    def __init__(self, id_: str, nombre: str, puntos: list[dict[str, Any]], carpeta: Path):
        self.id = id_
        self.nombre = nombre
        self.puntos = puntos
        self.carpeta = carpeta
        self.estado = "corriendo"
        self.resultados: list[dict[str, Any]] = []
        self.actual: str | None = None
        self.mensaje = ""
        self.esperando_s = 0
        self.cancelar = threading.Event()
        self.inicio = time.time()
        self.fin: float | None = None
        self.hilo: threading.Thread | None = None

    def publico(self, desde: int = 0) -> dict[str, Any]:
        hechos = len(self.resultados)
        return {"id": self.id, "nombre": self.nombre, "estado": self.estado, "total": len(self.puntos), "hechos": hechos, "actual": self.actual,
                "mensaje": self.mensaje, "esperando_s": self.esperando_s, "carpeta": str(self.carpeta), "segundos": round((self.fin or time.time()) - self.inicio),
                "resultados": self.resultados[max(0, int(desde or 0)):]}


def _total(resumen: dict[str, Any] | None) -> float | None:
    for t in (resumen or {}).get("totales") or []:
        if t.get("clave") == "total":
            return t.get("valor")
    return None


def procesar_punto(p: dict[str, Any], carpeta: Path, cache: Any, bloqueo: Any, ahora: datetime | None = None) -> dict[str, Any]:
    """Route + approximate budget of one point (the same calculation as the viewer's «Presupuesto aproximado»). The Excel goes to `carpeta`."""
    from visor_rutas import calculo, presupuesto_visor
    base = {"n": p["n"], "nombre": p["nombre"], "codigo": p["codigo"], "lat": p["lat"], "lon": p["lon"]}
    datos = calculo.calcular_ruta(p["lat"], p["lon"], get_route=cache)
    ruta = datos["ruta"]
    with bloqueo:
        est = presupuesto_visor.estimar(p["lat"], p["lon"], ruta["metraje_final_m"], cache)
    resultado = presupuesto_visor.resultado_para_visor(est["respuesta"])
    origen = Path(resultado["archivo"])
    nombre = f"{p['n']:02d}_{_seguro(p['nombre'] or p['codigo'] or 'punto')}.xlsx"
    destino = carpeta / nombre
    shutil.move(str(origen), destino)
    almacen_archivos.borrar("visor", origen.name)                 # generar() stored it at the top level; the batch keeps it in its folder
    avisos = list(dict.fromkeys(list(datos.get("avisos") or []) + list(resultado.get("avisos") or [])))
    return {**base, "ok": True, "conexion": datos["conexion"]["nombre"], "conexion_lat": datos["conexion"]["lat"], "conexion_lon": datos["conexion"]["lon"],
            "distancia_m": ruta["distancia_m"], "metraje_final_m": ruta["metraje_final_m"], "tipo_cable": ruta["tipo_cable"],
            "aproximada": ruta["aproximada"], "ruta": [list(x) for x in ruta["puntos"]], "postes_existentes": len(datos["postes_existentes"]),
            "postes_nuevos": datos["postes_nuevos"]["cantidad"], "mufas": len(datos["mufas_intervenidas"]), "total": _total(resultado.get("resumen")),
            "calculado": resultado["calculado"], "archivo": str(destino), "nombre_archivo": nombre, "avisos": avisos}


COLUMNAS_RESUMEN = [("#", "n", 5), ("Nombre", "nombre", 28), ("Código", "codigo", 12), ("Latitud", "lat", 12), ("Longitud", "lon", 12),
                    ("Punto de conexión", "conexion", 26), ("Distancia por ruta (m)", "distancia_m", 14), ("Metraje final (m)", "metraje_final_m", 14),
                    ("Cable", "tipo_cable", 14), ("Postes existentes", "postes_existentes", 10), ("Postes nuevos", "postes_nuevos", 10),
                    ("Mufas intervenidas", "mufas", 10), ("Total", "total", 14), ("Presupuesto", "nombre_archivo", 30), ("Avisos / error", "avisos", 60)]


def escribir_resumen(carpeta: Path, nombre: str, resultados: list[dict[str, Any]]) -> Path:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Resumen"
    hoja.append([f"Presupuestos aproximados — {nombre}"])
    hoja["A1"].font = Font(bold=True, size=13)
    hoja.append([])
    hoja.append([c[0] for c in COLUMNAS_RESUMEN])
    for celda in hoja[3]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F6FEB")
        celda.alignment = Alignment(wrap_text=True, vertical="center")
    for r in resultados:
        fila = []
        for _, clave, _ in COLUMNAS_RESUMEN:
            v = r.get(clave) if r.get("ok") or clave in ("n", "nombre", "codigo", "lat", "lon", "avisos", "error") else None
            if clave == "avisos":
                v = "; ".join(r.get("avisos") or []) if r.get("ok") else (r.get("error") or "")
            fila.append(v)
        hoja.append(fila)
    for i, (_, _, ancho) in enumerate(COLUMNAS_RESUMEN, 1):
        hoja.column_dimensions[hoja.cell(3, i).column_letter].width = ancho
    hoja.freeze_panes = "A4"
    ruta = carpeta / "Resumen.xlsx"
    libro.save(ruta)
    return ruta


class GestorLotes:
    """The batches of this process (memory): a thread each, the newest MAX_LOTES_GUARDADOS kept."""

    def __init__(self) -> None:
        self._lotes: dict[str, Lote] = {}
        self._lock = threading.Lock()
        self._activo: Lote | None = None

    def iniciar(self, nombre: str, puntos: list[dict[str, Any]], carpeta_base: Path, por_minuto: int, maximo: int, bloqueo: Any, cache_factory: Any) -> Lote:
        puntos = [p for p in puntos if p.get("incluir") and p.get("estado") in (OK,)]
        if not puntos:
            raise ValueError("No hay coordenadas válidas marcadas para calcular")
        if len(puntos) > maximo:
            raise ValueError(f"Un lote puede tener hasta {maximo} coordenadas (hay {len(puntos)}): divida el archivo")
        with self._lock:
            if self._activo is not None and self._activo.estado == "corriendo":
                raise ValueError(f"Ya hay un lote en curso ({self._activo.nombre}): espere a que termine o cancélelo")
            ahora = datetime.now()
            id_ = f"{ahora:%Y%m%d-%H%M%S}"
            carpeta = carpeta_base / f"masivo_{id_}_{_seguro(Path(nombre).stem)}"
            carpeta.mkdir(parents=True, exist_ok=True)
            lote = Lote(id_, Path(nombre).stem or "lote", puntos, carpeta)
            self._lotes[id_] = lote
            self._activo = lote
            for viejo in sorted(self._lotes)[:-MAX_LOTES_GUARDADOS]:
                del self._lotes[viejo]
        lote.hilo = threading.Thread(target=self._correr, args=(lote, por_minuto, bloqueo, cache_factory()), name=f"lote-{id_}", daemon=True)
        lote.hilo.start()
        return lote

    def _correr(self, lote: Lote, por_minuto: int, bloqueo: Any, cache: Any) -> None:
        limitada = CacheLimitada(cache, por_minuto, lote)
        final = "terminado"
        try:
            for p in lote.puntos:
                if lote.cancelar.is_set():
                    break
                lote.actual = p["nombre"] or p["texto"]
                try:
                    lote.resultados.append(procesar_punto(p, lote.carpeta, limitada, bloqueo))
                except InterruptedError:
                    break
                except Exception as exc:                                    # one point failing never stops the batch
                    log.warning("Bulk: point %s failed (%s: %s)", p["n"], type(exc).__name__, exc)
                    error = str(exc) if exc.__class__.__name__ in ("VisorError", "SinPuntoConexion", "PresupuestoError", "CoordenadaInvalida") else \
                        f"No se pudo calcular ({type(exc).__name__}): {exc}"
                    lote.resultados.append({"n": p["n"], "nombre": p["nombre"], "codigo": p["codigo"], "lat": p["lat"], "lon": p["lon"], "ok": False, "error": error})
            final = "cancelado" if lote.cancelar.is_set() else "terminado"
        except Exception as exc:
            log.exception("Bulk: the batch failed")
            final, lote.mensaje = "error", f"{type(exc).__name__}: {exc}"
        finally:
            lote.actual = None
            lote.esperando_s = 0
            try:
                escribir_resumen(lote.carpeta, lote.nombre, lote.resultados)
            except Exception:
                log.exception("Bulk: could not write Resumen.xlsx")
            try:
                almacen_archivos.subir_carpeta("visor", lote.carpeta.parent, patron=f"{lote.carpeta.name}/*", excluir=("mapas",))
            except Exception:
                log.exception("Bulk: could not store the batch in the database")
            lote.fin = time.time()
            lote.estado = final                                             # only now: the summary is already written when the page sees it finished
            log.info("Bulk %s: %s, %d of %d", lote.id, lote.estado, len(lote.resultados), len(lote.puntos))

    def obtener(self, id_: str) -> Lote:
        with self._lock:
            lote = self._lotes.get(str(id_))
        if lote is None:
            raise KeyError("El lote ya no existe (el visor o el servicio se reinició)")
        return lote

    def comprimir(self, id_: str) -> tuple[str, bytes]:
        lote = self.obtener(id_)
        if lote.estado == "corriendo":
            raise ValueError("El lote todavía está calculando")
        salida = io.BytesIO()
        with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as z:
            for archivo in sorted(lote.carpeta.glob("*.xlsx")):
                z.write(archivo, archivo.name)
        return f"{lote.carpeta.name}.zip", salida.getvalue()
